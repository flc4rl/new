#!/usr/bin/env python3
"""
Build the two master files from one or more result folders:

  master_network   one row per event x organisation x source: which
                   organisations were found for which event, and where.
                   Sheets:
                     event_organisations  the network (long format: a source
                                          tied to two events gives two rows): by
                                          default only rows that are evidence of
                                          taking part in the event (see Network
                                          scope below)
                     context_organisations rows that are not (linked reports and
                                          studies, other pages of a site several
                                          events link to, sources about another
                                          year, rows tied to no event, copies of
                                          a document), with context_reason
                     events               each event of the network once, with
                                          its organisations
                     all_events           every event found (agenda entries,
                                          calendar, event pages), with how well
                                          it was covered and whether it is in
                                          the network
                     menus_and_footers    names found only in menus, headers,
                                          footers or sidebars (site furniture),
                                          kept apart so they do not inflate the
                                          network. Names in a passage repeated in
                                          the main content of many pages stay in
                                          the network (repeated_on_site=True)
                     ner_only_candidates  rows whose organisation was found only by
                                          NER (spaCy) with no other evidence in the
                                          corpus (see OrgEvidence): kept apart so
                                          nothing is lost, but not in the network
                                          unless --keep-ner-only is given; also
                                          rows found only by NER of a one-word
                                          name read as an ordinary word, when
                                          other rows keep the name
                     rejected_names       names and mentions the run's vetting
                                          dropped (rejected_names.csv: ordinary
                                          words, weekdays, places, people), with
                                          the reason, for review (newer folders)
                     sources              every source read (web page, PDF,
                                          agenda entry, calendar cell) with its
                                          theme hits
                     excluded_sources     sources whose organisations were left
                                          out by the theme filter, with the reason
                     about                how the file was made
  master_themes    one row per theme term found and event (long format): the
                   event, the source link, and the passage (two sentences
                   before and after). Terms in menus, sidebars, footers, site
                   furniture and short labels are in a separate sheet
                   (menus_and_sidebars) unless --keep-menu-themes is given.

Theme filter (--theme-filter, default "source"): a theme hit is the inclusion
criterion for master_network. Organisations and events are kept only from
sources with at least one hit of a focus theme in their content (not only in
menus, headers, footers or sidebars; themes marked "#: descriptive" in the
themes file do not count): "source" = the web page, the PDF, or for seed agendas
the agenda entry itself; "direct" = as "source", but agenda entries, calendar
cells, the event's own page and its landing page also pass when any source of
their event has a hit; "site" = any page of the same website (an agenda entry
is still judged on its own, as the agenda PDF counts as one site); "event" = any
source tied to the same event; "none" keeps everything. master_themes is not
filtered. The run log and the about sheet say how much was left out.

Network scope (--network-scope, default "participation"): every network row
says what kind of document its source is (doc_kind: listing, page, programme,
report; see classify_document), how the source was tied to the event
(attribution: event, edition, site_level, not_tied, other_year), whether it is
a copy or translation of another document, and whether it is evidence of
taking part in the event (participation, context_reason). Rows with
participation=no go to the context_organisations sheet; "all" keeps every row
in event_organisations.

Country columns (event_country, country_iso, country_source, location_text)
come from events.csv of each result folder, matched on event_key. For folders
made by older versions (no country columns) they are worked out with
gazetteer.py, as analytics.py does.

Names that join two organisations ("Exxon & Shell", "Greenpeace and Civil
Rights Defenders") are split into one row per organisation when each part is an
organisation the corpus or the run's watchlist corroborates; found_by then says
"split from ...".

Written as .xlsx (one sheet per table, filters, frozen header) and as .csv.
site_actors.py calls this at the end of every run; run it yourself to merge
several runs or to change the filter:

    python make_master.py results/movementhub results/pdf_agendas -o results/master
    python make_master.py results/movementhub -o results/movementhub --theme-filter none
    python make_master.py results/movementhub -o results/movementhub --network-scope all
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import logging
import re
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Sequence, Set, Tuple
from urllib.parse import unquote, urlparse

log = logging.getLogger(__name__)
csv.field_size_limit(10_000_000)

NO_EVENT = "(not tied to an event)"
CONFIDENCE_RANK = {"low": 0, "medium": 1, "high": 2}
THEME_FILTERS = ["direct", "source", "site", "event", "none"]
NETWORK_SCOPES = ["participation", "all"]
# Layers that document the event itself: under the 'direct' theme filter they count when
# any source of their event has a theme hit.
DIRECT_LAYERS = {"listing_seed", "listing_calendar", "main_event_page", "landing"}
LISTING_LAYERS = {"listing_seed", "listing_calendar"}
FURNITURE_REGIONS = {"nav", "header", "footer", "aside"}
COUNTRY_FIELDS = ["event_country", "country_iso", "country_source", "location_text"]

NETWORK_FIELDS = ["event", "event_date", "event_link", "event_key", *COUNTRY_FIELDS,
                  "organisation", "org_key", "org_type", "confidence", "mentions_on_source",
                  "source_link", "source_unit", "source_title", "site", "site_role", "layer", "is_pdf",
                  "page_date", "fetched_at", "content_vs_event", "years_after_event",
                  "recurring_event_editions", "site_linked_from_events", "found_by",
                  "named_in_event_title", "only_in_menus_or_footers", "repeated_on_site", "source_has_theme",
                  "doc_kind", "doc_kind_reason", "attribution", "source_year", "duplicate_document_of",
                  "translation_of", "participation", "context_reason",
                  "example_context"]
EVENT_FIELDS = ["event", "event_date", "event_link", "event_key", *COUNTRY_FIELDS,
                "n_organisations", "organisations", "n_sources", "sources"]
ALL_EVENT_FIELDS = ["event", "date", "event_link", "event_key", *COUNTRY_FIELDS, "listed_on",
                    "calendar_months", "entry_words", "link_followed", "link_status", "pages_analysed",
                    "words_analysed", "org_mentions", "sources_with_theme_hit", "in_network",
                    "network_organisations", "context_organisations", "ner_only_candidates",
                    "left_out_by_theme_filter", "location_structured",
                    "merged_from"]
THEME_FIELDS = ["theme", "term", "event", "event_date", "event_link", "event_key", *COUNTRY_FIELDS,
                "n_events_for_passage", "layer", "content_vs_event", "years_after_event",
                "recurring_event_editions", "site_linked_from_events",
                "source_link", "source_unit", "page_title", "page_date", "fetched_at",
                "context_before", "sentence", "context_after", "organisations_in_paragraph",
                "organisations_in_paragraph_keys",
                "site", "site_role", "is_pdf", "pdf_found_on", "region", "boilerplate",
                "menu_or_sidebar", "duplicate_of", "focus_theme", "doc_kind", "attribution",
                "source_year", "duplicate_document_of", "translation_of", "passage_id", "doc_theme_rows"]
SOURCE_FIELDS = ["source_unit", "source_link", "page_title", "site", "site_role", "site_status", "kind",
                 "is_pdf", "page_date", "fetched_at", "words", "event_keys", "events", "org_mentions",
                 "has_theme_hit", "theme_hits", "menu_theme_hits", "in_network"]
EXCLUDED_FIELDS = ["source_unit", "source_link", "site", "site_role", "events", "n_organisations",
                   "organisations", "reason"]


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------
def _read(path: Path) -> List[dict]:
    return list(_iter(path))


def _iter(path: Path) -> Iterator[dict]:
    """Rows of a CSV file, one at a time (the mentions of a large run do not fit
    comfortably in memory twice)."""
    if not path.exists() or path.stat().st_size == 0:
        return
    with path.open(encoding="utf-8-sig", newline="") as fh:
        yield from csv.DictReader(fh)


def _true(v) -> bool:
    return str(v).strip().lower() in ("true", "yes", "1")


def _split(value: str) -> List[str]:
    return [v.strip() for v in (value or "").split(" | ")]


def _events_of(row: dict) -> List[dict]:
    """The events a mention or theme row is tied to, one dict each (key, name, date,
    link, and where known layer, content_vs_event, years_after_event, in_title).

    New result folders store them as JSON (events_json); older ones joined the
    names, dates and links with " | ", which is split again here."""
    js = row.get("events_json")
    if js:
        try:
            evs = [e for e in json.loads(js) if isinstance(e, dict)]
            if evs:
                return evs
            return [{"name": NO_EVENT, "date": "", "link": "", "key": ""}]
        except ValueError:
            pass
    names, dates, links = _split(row.get("event", "")), _split(row.get("event_date", "")), _split(row.get("event_link", ""))
    keys = _split(row.get("event_key", ""))
    if not any(names) and not any(links):
        return [{"name": NO_EVENT, "date": "", "link": "", "key": ""}]
    if len(names) > len(links) == len(dates) >= 1 and len(links) == 1:
        # One event whose name itself contained " | " ("National Climate March | Netherlands").
        names = [" | ".join(names)]
    n = max(len(names), len(dates), len(links))
    pad = lambda xs: xs + [""] * (n - len(xs))  # noqa: E731
    keys = keys if len(keys) == n else [""] * n
    return [{"name": a, "date": b, "link": c, "key": k} for a, b, c, k in zip(pad(names), pad(dates), pad(links), keys)]


def _event_id(ev: dict) -> str:
    return ev.get("key") or f"{ev.get('name', '')}\x1f{ev.get('link', '')}"


def _layer_fallback(row: dict, ev: dict) -> str:
    """Layer for rows of older result folders (no events_json)."""
    url = row.get("url") or row.get("source_link") or ""
    if row.get("site_role") == "seed" or url.startswith(("file:", "pdf:")):
        return "listing_seed"
    if row.get("calendar_month") or "calendar_month=" in url:
        return "listing_calendar"
    placeholder = str(ev.get("name") or NO_EVENT).startswith("(")
    if row.get("site_role") == "main":
        return "main_page" if placeholder else "main_event_page"
    norm = lambda u: re.sub(r"^https?://(?:www\.)?", "", (u or "").split("#")[0]).rstrip("/")  # noqa: E731
    if ev.get("link") and norm(ev["link"]) == norm(url):
        return "landing"
    pdf = _true(row.get("is_pdf")) or bool(re.search(r"\.pdf(?:\?|$)", url, re.I))
    return "linked_pdf" if pdf else "linked_page"


# --------------------------------------------------------------------------
# What kind of document a source is, and how it was tied to its event
# --------------------------------------------------------------------------
# Words in a PDF's title or file name that make it part of the event itself (a programme,
# flyer, call, invitation, press release) or a publication about something else (a
# report, study, guide, newsletter). Matched anywhere in a word ("Programmheft").
_PROGRAMME_WORDS = re.compile(
    r"(?i)prog(?:ramm|rama|ramme|ram)?\b|prog[_\-. ]|programm|programa|programme|program[_\-. ]|flyer"
    r"|faltblatt|folder|aufruf|appel|call[_\-. ]?(?:for|to)|einladung|invitation|invito|welcome|willkommen"
    r"|bienvenue|accueil|livret|carnet|leaflet|tract|plakat|affiche|poster|ablaufplan|zeitplan|timetable"
    r"|schedule|horaire|fahrplan|pressemitteilung|press[_\- ]?release|communiqu|planprog")
_PROGRAMME_CASED = re.compile(r"\bPM[_\-]")
_REPORT_WORDS = re.compile(
    r"(?i)report|rapport|bericht|studie|study|étude|etude|estudio|informe|relat[oó]rio|zpr[aá]va|dossier"
    r"|working[_\- ]?paper|paper|booklet|brosch|brochure|guide|handbook|handbuch|leitfaden|assessment"
    r"|review|publikation|publication|analys|annual|jahresbericht|newsletter|bulletin|echoes|zeitung"
    r"|beilage|journal|magazin|zine|factsheet|fact[_\- ]sheet|position|memorandum|satzung|statut"
    r"|estatuto|beleidsplan|aanbeveling|recommendation|forderung|faq|strategie|strategy"
    r"|legal[_\- ]?opinion")
# Pages that name many organisations as the event's own partners, not as a report would.
# Coalition and member lists ("over ons", "wer wir sind", "Bündnis", "Trägerkreis") name
# the event's own organisers too.
_PARTNER_WORDS = re.compile(
    r"(?i)partner|partenaires|unterst[uü]tzer|signataires|signatories|who we are|[uü]ber uns"
    r"|qui sommes|line-?up|agenda|over ons|wie (?:we|wij) zijn|wer wir sind|chi siamo"
    r"|qui[eé]nes somos|quem somos|om oss|om os\b|members|mitglieder|membres|miembros|membri|leden"
    r"|b[uü]ndnis|coalit|alliance|alliantie|allianz|tr[aä]gerkreis|organisers|organizers"
    r"|organisatoren|organisateurs")
LONG_PDF_WORDS = 3000      # a PDF without either kind of word is a report from this length
MANY_ORGS_ON_PAGE = 100    # a web page naming more organisations is a list or report


def classify_document(url: str, title: str, words: int, n_orgs: int, is_pdf: bool,
                      layer: str) -> Tuple[str, str]:
    """The kind of document a source is, and why: 'listing' (an agenda entry or calendar
    cell), 'programme' (a PDF that is part of the event: programme, flyer, call,
    invitation, press release, or a short PDF), 'report' (a linked publication: report,
    study, guide, newsletter, or a long PDF; or a web page naming more than
    MANY_ORGS_ON_PAGE organisations that is not a programme or partner list) or 'page'.
    Landing pages and the event's own page are always 'page'."""
    if layer in LISTING_LAYERS:
        return "listing", "agenda entry or calendar cell"
    if layer in ("landing", "main_event_page"):
        return "page", "the event's own page" if layer == "main_event_page" else "the page the event links to"
    try:
        name = unquote(urlparse(url or "").path.rstrip("/").rsplit("/", 1)[-1])
    except ValueError:
        name = ""
    text = f"{title or ''} {name}"
    prog = _PROGRAMME_WORDS.search(text) or _PROGRAMME_CASED.search(text)
    if is_pdf:
        if prog:
            return "programme", f"programme word '{prog.group(0).strip()}' in the title or file name"
        rep = _REPORT_WORDS.search(text)
        if rep:
            return "report", f"report word '{rep.group(0).strip()}' in the title or file name"
        if (words or 0) >= LONG_PDF_WORDS:
            return "report", f"PDF of {words} words (at least {LONG_PDF_WORDS})"
        return "programme", f"PDF of {words or 0} words (fewer than {LONG_PDF_WORDS})"
    if layer in ("linked_page", "main_page") and (n_orgs or 0) > MANY_ORGS_ON_PAGE:
        if prog:
            return "page", f"names {n_orgs} organisations, programme word '{prog.group(0).strip()}'"
        try:   # the whole address counts here (/qui-sommes-nous/notre-manifeste/)
            path = re.sub(r"[-_/.+]+", " ", unquote(urlparse(url or "").path))
        except ValueError:
            path = ""
        partner = _PARTNER_WORDS.search(f"{title or ''} {path}")
        if partner:
            return "page", f"names {n_orgs} organisations, partner word '{partner.group(0).strip()}'"
        return "report", f"web page naming {n_orgs} organisations (more than {MANY_ORGS_ON_PAGE})"
    return "page", ""


ATTRIBUTIONS = ("event", "edition", "site_level", "not_tied", "other_year")
_CONTEXT_REASONS = OrderedDict([("report", "linked report or study"),
                                ("other_year", "about another year"),
                                ("site_level", "other pages of a site several events link to"),
                                ("not_tied", "not tied to an event"),
                                ("copy", "copy of")])


def _attribution(ev: dict) -> str:
    """How a row's source was tied to its event; derived for folders made before
    events_json had it ('(other pages of ...' -> site_level, no event -> not_tied)."""
    a = ev.get("attribution") or ""
    if a in ATTRIBUTIONS:
        return a
    name = str(ev.get("name") or NO_EVENT)
    if name == NO_EVENT or not name.strip():
        return "not_tied"
    return "site_level" if name.startswith("(") else "event"


def _parse_editions(value: str) -> List[Tuple[str, str]]:
    """recurring_event_editions ("Name (2023-07-12 to 2023-07-14); Other (2021-08-01)")
    as (name, date) pairs."""
    out = []
    for part in (value or "").split("; "):
        if not part.strip():
            continue
        m = re.match(r"(.*) \((\d{4}-\d\d(?:-\d\d)?[^()]*)\)$", part)
        out.append((m.group(1), m.group(2)) if m else (part, ""))
    return out


_title_key_cache: Dict[str, str] = {}


def _in_title(org_key: str, event_name: str) -> bool:
    """Organisation named in the event title as whole words (for older folders)."""
    from extract import _KEY_SHAPE, canonical_key, squash_runs
    k = org_key or ""
    # org_key is already a key: keying it again is not idempotent for a glued first word
    # ("A-camp" has the key "a camp", which keyed again becomes "camp" and would match any
    # title with "camp" in it). Only re-key values that do not look like a key.
    if not (_KEY_SHAPE.fullmatch(k) and k == k.casefold()):
        k = canonical_key(k)
    if not k:
        return False
    t = _title_key_cache.get(event_name)
    if t is None:
        # also keyed with its leading article kept ("A Camp, ..." for the key "a camp")
        t = _title_key_cache[event_name] = (f" {canonical_key(event_name or '')} "
                                            f" {canonical_key('zzqx ' + (event_name or ''))[5:]} ")
    if f" {k} " in t:
        return True
    # Spacing variants on word boundaries: "Geef Tegengas" in "Geef Tegen Gas ...".
    sq = re.sub(r"[\W_]+", "", k)   # not squash_key(k), which would key k again
    if len(sq) >= 4 and sq in (squash_runs(event_name or "") | squash_runs("zzqx " + (event_name or ""))):
        return True
    # "#ShellMustFall": a one-word name at the start of a hashtag.
    return " " not in k and len(k) >= 4 and any(
        canonical_key(h).startswith(k) for h in re.findall(r"#(\w+)", event_name or ""))


_country_note: Set[str] = set()


def _seed_entry_texts(res: Path) -> Dict[str, str]:
    """Agenda entry text by event name (casefolded), from seed_links.json."""
    p = res / "seed_links.json"
    out: Dict[str, str] = {}
    try:
        data = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
        for pdf in (data.get("pdfs") or {}).values():
            for e in pdf.get("entries") or []:
                out.setdefault((e.get("name") or "").casefold().strip(), " ".join(e.get("lines") or []))
    except (OSError, ValueError, AttributeError):
        pass
    return out


_OLD_FOLDERS = ("\x00older folder",)   # marker key in by_name: work out missing countries


def _gazetteer_countries(res: Path, rows: List[dict], quiet: bool = False) -> Dict[int, dict]:
    """Country columns for events.csv rows of an older folder (no country columns),
    worked out with gazetteer.py as analytics.py does. {} when it is not available."""
    try:
        import gazetteer
        from events import country_columns
    except ImportError as exc:
        if not quiet and str(res) not in _country_note:
            _country_note.add(str(res))
            log.warning("events.csv in %s has no country columns (older version) and gazetteer.py "
                        "could not be loaded (%s): the country columns stay empty. Run site_actors.py "
                        "--no-crawl on this folder to add them", res, exc)
        return {}
    entries = _seed_entry_texts(res)
    out: Dict[int, dict] = {}
    for i, e in enumerate(rows):
        link = e.get("event_link") or ""
        link = link if link.startswith(("http://", "https://")) else ""
        try:
            got = gazetteer.event_country(e.get("event") or "",
                                          entries.get((e.get("event") or "").casefold().strip(), ""),
                                          link=link, structured=e.get("location_structured") or "")
            if got.get("countries") and not got.get("names"):
                got = {**got, "names": [gazetteer.country_name(c) for c in got["countries"]]}
            if isinstance(got.get("matched"), (list, tuple)):
                got = {**got, "matched": "; ".join(str(x) for x in got["matched"])}
            out[i] = country_columns(got)
        except Exception as exc:  # noqa: BLE001 - one odd event name must not stop the master files
            log.debug("No country for %r: %s", e.get("event"), exc)
    if not quiet and str(res) not in _country_note:
        _country_note.add(str(res))
        log.info("events.csv in %s has no country columns (older version): countries worked out "
                 "with the gazetteer for %d of %d events (run site_actors.py --no-crawl on this "
                 "folder to store them)", res, sum(1 for c in out.values() if c.get("country_iso")),
                 len(rows))
    return out


_countries_memo: Dict[tuple, Tuple[Dict[str, dict], Dict[tuple, dict]]] = {}


def _countries(results: Sequence[Path]) -> Tuple[Dict[str, dict], Dict[tuple, dict]]:
    memo_key = tuple(str(r) for r in results) + tuple(
        str((r / "events.csv").stat().st_mtime_ns) if (r / "events.csv").exists() else "" for r in results)
    if memo_key not in _countries_memo:
        _countries_memo.clear()
        _countries_memo[memo_key] = _read_countries(results)
    return _countries_memo[memo_key]


def _read_countries(results: Sequence[Path]) -> Tuple[Dict[str, dict], Dict[tuple, dict]]:
    by_key: Dict[str, dict] = {}
    by_name: Dict[tuple, dict] = {}
    for res in results:
        rows = list(_iter(res / "events.csv"))
        derived = (_gazetteer_countries(res, rows)
                   if rows and "country_iso" not in rows[0] else {})
        if derived:
            by_name[_OLD_FOLDERS] = {"derive": "yes"}
        for i, e in enumerate(rows):
            c = derived.get(i) or {f: e.get(f, "") for f in COUNTRY_FIELDS}
            if e.get("event_key"):
                by_key.setdefault(e["event_key"], c)
            by_name.setdefault((e.get("event", ""), e.get("event_link", "")), c)
    return by_key, by_name


def _country_of(ev: dict, by_key: Dict[str, dict], by_name: Dict[tuple, dict]) -> dict:
    c = by_key.get(ev.get("key") or "") or by_name.get((ev.get("name", ""), ev.get("link", "")))
    if c is None and by_name.get(_OLD_FOLDERS) and not str(ev.get("name") or "(").startswith("("):
        # Older folders: the event names of mentions.csv do not always match events.csv.
        k = (ev.get("name", ""), ev.get("link", ""))
        c = by_name[k] = _gazetteer_countries(Path("."), [{"event": k[0], "event_link": k[1]}],
                                              quiet=True).get(0)
    return c or {f: "" for f in COUNTRY_FIELDS}


# --------------------------------------------------------------------------
# Corpus evidence for organisation names (NER-only candidates, coordinated names)
# --------------------------------------------------------------------------
NER_ONLY_METHODS = frozenset({"ner", "pattern_after", "link_text", "website_elsewhere"})  # fallback


def _weak_methods() -> frozenset:
    """Methods that are not evidence on their own (extract.NER_LIKE_METHODS, else the
    fallback above). 'website' and 'link' of older folders stay strong."""
    try:
        from extract import NER_LIKE_METHODS
        return frozenset(NER_LIKE_METHODS) or NER_ONLY_METHODS
    except ImportError:
        return NER_ONLY_METHODS


# Weekday names are never organisations, however often they appear ("Freitag", "Samedi").
WEEKDAYS = frozenset("""
monday tuesday wednesday thursday friday saturday sunday
montag dienstag mittwoch donnerstag freitag samstag sonnabend sonntag
lundi mardi mercredi jeudi vendredi samedi dimanche
lunes martes miércoles miercoles jueves viernes sábado sabado domingo
lunedì lunedi martedì martedi mercoledì mercoledi giovedì giovedi venerdì venerdi sabato domenica
maandag dinsdag woensdag donderdag vrijdag zaterdag zondag
segunda terça terca quarta quinta sexta
""".split())
CANDIDATE_FIELDS = NETWORK_FIELDS + ["corpus_sources", "corpus_sites", "why_candidate"]
REJECTED_FIELDS = ["organisation", "org_key", "scope", "reason", "reason_text", "mentions_dropped",
                   "mentions_kept", "methods_dropped", "pages", "sites", "example_url", "example_context",
                   "result_folder"]
try:   # weekday names as organisation keys ("Tuesday" -> "tusday"), see extract.WEEKDAY_KEYS
    from extract import WEEKDAY_KEYS as _WEEKDAY_KEYS
except ImportError:
    _WEEKDAY_KEYS = WEEKDAYS
_LEGAL_FORM = re.compile(r"\s(?:AG|SE|plc|PLC|Ltd\.?|Limited|GmbH|LLC|Inc\.?|S\.?A\.?|SpA|S\.p\.A\.|N\.?V\.?|"
                         r"B\.?V\.?|e\.\s?V\.|Oy|Oyj|AB|ASA|A/S|SAS|SARL|Srl|KG|Corp\.?|Corporation)$")
_CAMEL = re.compile(r"[A-ZÀ-Ý][a-zà-ÿ]+(?:[A-ZÀ-Ý][a-zà-ÿ]+)+")
_ACRONYM = re.compile(r"[A-Z0-9ÀÂÄÇÉÈÊËÎÏÔÖÙÛÜ][A-Z0-9ÀÂÄÇÉÈÊËÎÏÔÖÙÛÜ&\-]{1,5}")


def _method_names(value: str) -> Set[str]:
    """"ner (2); pattern (1)" (organisations.csv), "ner+pattern" (mentions.csv) or
    "ner; website; split from ..." (found_by of master_network)."""
    return {re.sub(r"\s*\(\d+\)$", "", m).strip() for m in re.split(r";\s*|\+", value or "")
            if not m.strip().startswith("split from")} - {""}


def _weak_for(weak: frozenset, surface: str) -> frozenset:
    """`weak` without 'website_elsewhere' for a name of several words, an acronym or a
    camelCase name (extract.weak_methods_for): only for a one-word name is a website
    given in another paragraph weak evidence."""
    try:
        from extract import weak_methods_for
        if "website_elsewhere" in weak_methods_for(surface):
            return weak
    except ImportError:
        if not (" " in surface.strip() or _ACRONYM.fullmatch(surface) or _CAMEL.fullmatch(surface)):
            return weak
    return frozenset(weak - {"website_elsewhere"})


class OrgEvidence:
    """What the whole corpus (every mention in every result folder, theme hit or
    not) says about each organisation key: the rules that found it, how many
    sources and websites name it, and whether a watchlist of the run lists it.

    A name found only by NER (spaCy) is "corroborated" when it has other evidence:
    - another rule found it somewhere (a descriptor or organisational noun, a
      website or link, an agenda title, the watchlist);
    - its own words include an organisational noun (org_type is set);
    - an acronym was spelt out for it in the text ("Climate Action Network (CAN)");
    - it ends in a legal form ("Pan American Energy LLC", "RWE AG");
    - it is an acronym of 2 to 6 capitals named in 2 or more sources, not a place;
    - or, not being a place, it is named in the content (not menus or footers) of 3 or
      more sources on 2 or more websites, or of 3 or more agenda entries.
    Other NER-only names go to the ner_only_candidates sheet instead of the network."""

    def __init__(self, results: Sequence[Path]):
        self.methods: Dict[str, Set[str]] = defaultdict(set)
        self.sources: Dict[str, Set[str]] = defaultdict(set)
        self.sites: Dict[str, Set[str]] = defaultdict(set)
        self.content: Dict[str, Set[Tuple[str, str, bool]]] = defaultdict(set)   # (source, site, seed)
        self.surfaces: Dict[str, Counter] = defaultdict(Counter)
        self.typed: Set[str] = set()
        self.spelt: Set[str] = set()
        self.watch: Set[str] = set()
        self.by_name: Dict[str, str] = {}     # canonical key of a display name -> org_key
        self.weak = _weak_methods()
        # Part-of-speech votes of one-word names (organisations.csv pos_votes, pos_common):
        # key -> (votes, common). Folders without the columns keep the older rule.
        self.pos: Dict[str, Tuple[int, int]] = {}
        self.no_pos_folders: List[str] = []
        for res in results:
            for m in _iter(res / "mentions.csv"):
                k = m.get("org_key") or ""
                if not k:
                    continue
                self.methods[k] |= _method_names(m.get("methods", ""))
                src = m.get("source_unit") or m.get("url", "")
                self.sources[k].add(src)
                self.sites[k].add(m.get("site", ""))
                if not _is_menu_mention(m):
                    self.content[k].add((src, m.get("site", ""), m.get("site_role") == "seed"))
                self.surfaces[k][m.get("organisation", "")] += 1
                if m.get("org_type"):
                    self.typed.add(k)
            has_pos = None
            for o in _iter(res / "organisations.csv"):
                k = o.get("org_key") or ""
                if has_pos is None:
                    has_pos = "pos_votes" in o
                if k and o.get("acronyms"):
                    self.spelt.add(k)
                if k and o.get("methods"):
                    self.methods[k] |= _method_names(o["methods"])
                if k and has_pos and str(o.get("pos_votes") or "").strip():
                    try:
                        v, c = int(float(o["pos_votes"])), int(float(o.get("pos_common") or 0))
                    except ValueError:
                        continue
                    pv, pc = self.pos.get(k, (0, 0))
                    self.pos[k] = (pv + v, pc + c)
            if has_pos is False and (res / "mentions.csv").exists():
                self.no_pos_folders.append(str(res))
            self.watch |= _watchlist_keys(res)
        try:
            from extract import canonical_key, tidy_surface
            for k, forms in self.surfaces.items():
                for f in forms:
                    self.by_name.setdefault(canonical_key(tidy_surface(f)), k)
                self.by_name.setdefault(k, k)
        except ImportError:   # extract.py missing: names are compared by their keys only
            self.by_name = {k: k for k in self.surfaces}

    def display(self, key: str) -> str:
        forms = self.surfaces.get(key)
        return forms.most_common(1)[0][0] if forms else key

    def why_not(self, key: str) -> str:
        """"" when the name is corroborated, else why it is only a candidate."""
        methods = self.methods.get(key, set())
        surface = self.display(key)
        if not methods or methods - _weak_for(self.weak, surface):
            return ""   # unknown here (keep it) or found by another rule as well
        if key in self.typed or key in self.spelt or key in self.watch:
            return ""
        n_src = len(self.sources.get(key, set()) - {""})
        n_sites = len(self.sites.get(key, set()) - {""})
        if _LEGAL_FORM.search(surface):
            return ""
        if _ACRONYM.fullmatch(surface) and n_src >= 2 and not _is_place(surface):
            return ""
        if _CAMEL.fullmatch(surface) and not _is_place(surface):
            return ""   # a coined one-word name with an inner capital ("GroenFront", "BlackRock")
        content = self.content.get(key, set())
        seed_entries = {src for src, _, seed in content if seed}
        broad = ((len({src for src, _, _ in content}) >= 3 and len({site for _, site, _ in content}) >= 2)
                 or len(seed_entries) >= 3) and not _is_place(surface)
        common = self.ordinary_word(key, surface)
        if broad and not common:
            return ""
        how = ("NER (spaCy)" if methods <= {"ner", "pattern_after"}
               else "NER (spaCy), link text or a website given elsewhere")
        why = (f"found only by {how}, in {n_src} source{'s' if n_src != 1 else ''} on "
               f"{n_sites} website{'s' if n_sites != 1 else ''}, with no other evidence")
        return f"{why}; {common}" if broad else why

    def ordinary_word(self, key: str, surface: str) -> str:
        """Why a one-word name may not use the breadth rule (named in 3 sources on 2
        websites, or in 3 agenda entries), or "": a weekday never; another one-word name
        that is not an acronym or camelCase when the part-of-speech tagger mostly read
        it as an ordinary word (pos_votes >= 2, at least half of them pos_common).
        Names without votes (never tagged, or folders made before pos_votes existed)
        keep the older rule; so do older folders' pos_votes of 0, which meant the same."""
        if " " in key.strip():
            return ""
        if (surface.casefold() in WEEKDAYS
                or (not surface and (key.casefold() in WEEKDAYS or key.casefold() in _WEEKDAY_KEYS))):
            return "a weekday"
        if _ACRONYM.fullmatch(surface) or _CAMEL.fullmatch(surface):
            return ""
        votes, common = self.pos.get(key, (0, 0))
        if votes < 2 or common < 0.5 * votes:
            return ""
        return f"one word read as an ordinary word in {common} of {votes} contexts"

    def why_not_row(self, row: dict) -> str:
        """why_not() of the row's organisation; else, when the row itself was found only
        by weak methods (its found_by) and the name is an ordinary word (ordinary_word:
        a weekday, or mostly read as a noun, verb ...), why the row is only a candidate:
        "die Welt retten" next to a name kept elsewhere by "Die Welt (2019) www.welt.de"."""
        key = row.get("org_key") or ""
        why = self.why_not(key)
        if why or not key:
            return why
        surface = row.get("organisation") or self.display(key)
        ms = _method_names(row.get("found_by", ""))
        common = self.ordinary_word(key, surface)
        weak = _weak_for(self.weak, surface)
        if common == "a weekday":
            weak = weak | {"appositive"}   # "forum FREITAG, SAMSTAG"
        if not ms or ms - weak:
            return ""
        if not common:
            return ""
        return (f"this row found only by {', '.join(sorted(ms))}; the name is {common}, "
                "kept elsewhere only where a rule, website or link supports it")

    def attested(self, name: str) -> str:
        """org_key of `name` when the corpus or a watchlist knows it as a corroborated
        organisation, else ""."""
        try:
            from extract import canonical_key, tidy_surface
            ck = canonical_key(tidy_surface(name))
        except ImportError:
            ck = name.casefold()
        if ck in self.watch:
            return self.by_name.get(ck, ck)
        k = self.by_name.get(ck, "")
        return k if k and not self.why_not(k) else ""


def _is_place(surface: str) -> bool:
    try:
        from extract import _is_place_name
        return _is_place_name(surface)
    except ImportError:
        return False


def _stoplist_for(res: Path):
    """The stoplist a run used (run_info.json; the bundled default_stoplist.txt when the
    folder does not say), as a function name -> True when stoplisted. Older result
    folders were made before some stoplist entries existed (FAQ, Google Calendar ...):
    re-reading them applies the current list."""
    try:
        from extract import canonical_key, load_stoplist, tidy_surface
    except ImportError:
        return lambda name: False
    path = Path(__file__).with_name("default_stoplist.txt")
    try:
        info = json.loads((res / "run_info.json").read_text(encoding="utf-8"))
        rec = ((info.get("input_files") or {}).get("stoplist") or {}).get("path")
        if rec and Path(rec).exists():
            path = Path(rec)
    except (OSError, ValueError, AttributeError):
        pass
    exact, pats = load_stoplist(path)

    def stopped(name: str) -> bool:
        surface = tidy_surface(name)
        return canonical_key(surface) in exact or any(p.search(surface) for p in pats)
    return stopped


def _watchlist_keys(res: Path) -> Set[str]:
    """Canonical keys of the watchlist a run used (from run_info.json)."""
    p = res / "run_info.json"
    try:
        info = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
        path = ((info.get("input_files") or {}).get("watchlist") or {}).get("path")
        if not path or not Path(path).exists():
            return set()
        from extract import Watchlist
        wl = Watchlist.load(Path(path))
        return set(wl.canonical) | set(wl.canonical.values())
    except (OSError, ValueError, ImportError, AttributeError):
        return set()


def split_coordinated(rows: List[dict], ev: OrgEvidence) -> int:
    """Split organisation names such as "Greenpeace and Civil Rights Defenders" or
    "Exxon & Shell" into their parts when every part is itself an organisation the
    corpus or a watchlist corroborates (see OrgEvidence). Names like "Perth and
    Kinross Council" or "Bund für Umwelt und Naturschutz" stay whole because their
    parts are not organisations of their own. Works in place; returns the number of
    rows split."""
    try:
        from extract import coordinated_parts
    except ImportError:
        return 0
    eid = lambda r: r.get("event_key") or f"{r['event']}\x1f{r['event_link']}"  # noqa: E731
    keep: List[dict] = []
    index: Dict[tuple, dict] = {}
    split: List[Tuple[dict, List[str]]] = []
    for r in rows:
        parts = coordinated_parts(r["organisation"])
        keys = [ev.attested(p) for p in parts] if len(parts) >= 2 else []
        if keys and all(keys) and len(set(keys)) == len(keys):
            split.append((r, keys))
            continue
        keep.append(r)
        index.setdefault((eid(r), r.get("org_key") or r["organisation"], r["source_link"]), r)
    for r, keys in split:
        for k in keys:
            have = index.get((eid(r), k, r["source_link"]))
            if have is not None:
                have["mentions_on_source"] += r["mentions_on_source"]
                continue
            q = {**r, "organisation": ev.display(k), "org_key": k,
                 "found_by": "; ".join(x for x in (r.get("found_by", ""), "split from " + r["organisation"]) if x)}
            index[(eid(q), k, q["source_link"])] = q
            keep.append(q)
    rows[:] = keep
    return len(split)


def _is_menu_mention(m: dict) -> bool:
    """A mention in site furniture: a menu, header, footer or sidebar. A passage repeated
    in the main content of many pages (boilerplate=True, region main) is not furniture:
    it names, say, the organisers on every event page. Older mentions.csv files without
    in_menu or region fall back to the boilerplate flag."""
    if "in_menu" not in m and "region" not in m:
        return _true(m.get("boilerplate"))
    return (_true(m.get("in_menu")) or _true(m.get("menu_or_sidebar"))
            or m.get("region", "") in FURNITURE_REGIONS)


def _is_menu_theme(r: dict) -> bool:
    if "menu_or_sidebar" in r and r["menu_or_sidebar"] != "":
        return _true(r["menu_or_sidebar"])
    return r.get("region", "") in FURNITURE_REGIONS or _true(r.get("boilerplate"))


# --------------------------------------------------------------------------
# Theme filter (inclusion criterion)
# --------------------------------------------------------------------------
class ThemeFilter:
    """Which sources, websites and events have a theme hit in their content."""

    def __init__(self, mode: str, results: Sequence[Path]):
        if mode not in THEME_FILTERS:
            raise ValueError(f"--theme-filter must be one of {', '.join(THEME_FILTERS)}")
        self.mode = mode
        self.units: Set[str] = set()
        self.sites: Set[str] = set()
        self.events: Set[str] = set()
        self.menu_only_units: Set[str] = set()
        self.old_folders: List[str] = []
        self.no_theme_file: List[str] = []   # folders without theme_hits.csv or page theme counts
        self.from_mentions: List[str] = []   # folders whose theme hits came from mentions.csv
        for res in results:
            src = res / "sources.csv"
            if src.exists():
                for r in _iter(src):
                    u = r.get("source_unit") or r.get("source_link", "")
                    if _true(r.get("has_theme_hit")):
                        self.units.add(u)
                        self.sites.add(r.get("site", ""))
                        self.events.update(k for k in (r.get("event_keys") or "").split("; ") if k)
                    elif int(r.get("menu_theme_hits") or 0):
                        self.menu_only_units.add(u)
                # Events tied to themed passages (also names-only ids for completeness).
                for r in _iter(res / "theme_hits.csv"):
                    if (r.get("focus_theme") or "yes") == "no":
                        continue
                    if not _is_menu_theme(r) and (r.get("source_unit") or r.get("source_link")) in self.units:
                        self.events.update(_event_id(e) for e in _events_of(r))
            elif not (res / "theme_hits.csv").exists():
                # The four files of an old run copied on their own (no theme_hits.csv):
                # the page-level theme counts of mentions.csv say which pages had a hit.
                self.old_folders.append(str(res))
                found_col = False
                for m in _iter(res / "mentions.csv"):
                    found_col = found_col or "page_themes" in m
                    if (m.get("page_themes") or "").strip():
                        u = self.unit_of(m)
                        self.units.add(u)
                        self.sites.add(m.get("site", ""))
                        self.events.update(_event_id(e) for e in _events_of(m))
                if found_col:
                    self.from_mentions.append(str(res))
                else:
                    self.no_theme_file.append(str(res))
            else:
                # Older folder: a source is a page or a whole PDF.
                self.old_folders.append(str(res))
                for r in _iter(res / "theme_hits.csv"):
                    u = r.get("source_link", "")
                    if (r.get("focus_theme") or "yes") == "no":
                        continue    # a descriptive theme does not count for the filter
                    if _is_menu_theme(r):
                        self.menu_only_units.add(u)
                        continue
                    self.units.add(u)
                    self.sites.add(r.get("site", ""))
                    self.events.update(_event_id(e) for e in _events_of(r))
        self.menu_only_units -= self.units

    def unit_of(self, m: dict) -> str:
        return m.get("source_unit") or m.get("url") or m.get("source_link", "")

    def keeps(self, m: dict, ev: dict) -> bool:
        if self.mode == "none":
            return True
        if self.mode == "direct":
            # The listing, the event's own page and its landing page pass when any source
            # of their event has a theme hit; every other source needs its own hit.
            if self.unit_of(m) in self.units:
                return True
            return (not str(ev.get("name") or NO_EVENT).startswith("(")
                    and (ev.get("layer") or _layer_fallback(m, ev)) in DIRECT_LAYERS
                    and _event_id(ev) in self.events)
        if self.mode == "site":
            if m.get("site_role") == "seed":
                # The "site" of an agenda entry is the whole agenda PDF: one hit anywhere in
                # it would keep every entry, so each entry is judged on its own.
                return self.unit_of(m) in self.units
            return m.get("site", "") in self.sites
        if self.mode == "event" and not str(ev.get("name", "")).startswith("("):
            return _event_id(ev) in self.events
        return self.unit_of(m) in self.units

    def reason(self, m: dict, ev: Optional[dict] = None) -> str:
        if (self.mode == "direct" and ev is not None
                and (ev.get("layer") or _layer_fallback(m, ev)) in DIRECT_LAYERS):
            return "event has no theme hit"
        if self.mode == "site" and m.get("site_role") != "seed":
            return "no theme term on any page of this website"
        if self.mode == "event":
            return "no theme term in any source tied to this event, nor in this source"
        if self.unit_of(m) in self.menu_only_units:
            return "theme terms only in menus, sidebars, headers, footers or short labels"
        return "no theme term in this source"


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------
class DocInfo:
    """What the result folders say about each source, for classify_document and the
    copy rule: words (sources.csv, else pages.csv), distinct organisations named in its
    content, and the events each source link is tied to."""

    def __init__(self, results: Sequence[Path]):
        self.words: Dict[str, int] = {}
        self.orgs: Dict[str, Set[str]] = defaultdict(set)
        self.link_events: Dict[str, Set[str]] = defaultdict(set)
        self.event_links: Dict[tuple, str] = {}           # (event name, date) -> link
        self.links_by_name: Dict[str, Set[str]] = defaultdict(set)
        for res in results:
            src = res / "sources.csv"
            rows = _iter(src) if src.exists() else (
                {"source_unit": p.get("url", ""), "source_link": p.get("url", ""), "words": p.get("words"),
                 "event_keys": p.get("event_keys", "")} for p in _iter(res / "pages.csv"))
            for r in rows:
                u = r.get("source_unit") or r.get("source_link") or ""
                try:
                    self.words[u] = int(float(r.get("words") or 0))
                except ValueError:
                    pass
                self.link_events[r.get("source_link") or u].update(
                    k for k in (r.get("event_keys") or "").split("; ") if k)
            for e in _iter(res / "events.csv"):
                self.event_links.setdefault((e.get("event", ""), e.get("date", "")), e.get("event_link", ""))
                self.links_by_name[e.get("event", "")].add(e.get("event_link", ""))

    def n_orgs(self, unit: str) -> int:
        return len(self.orgs.get(unit, ()))

    def classify(self, row: dict) -> Tuple[str, str]:
        unit = row.get("source_unit") or row.get("source_link") or ""
        return classify_document(row.get("source_link", ""), row.get("source_title") or row.get("page_title", ""),
                                 self.words.get(unit, 0), self.n_orgs(unit), _true(row.get("is_pdf")),
                                 row.get("layer", ""))

    def link_of(self, name: str, date: str) -> str:
        link = self.event_links.get((name, date))
        if link is None:
            links = sorted(self.links_by_name.get(name, ()))
            link = links[0] if links else ""
        return link


def _derived_attribution(m: dict, ev: dict, layer: str, docs: Optional[DocInfo]) -> Tuple[str, str]:
    """(attribution, source_year) of a row. New folders store them in events_json. For
    older folders they are worked out here: placeholders as in _attribution; a linked
    page or PDF tied to the nearest of several editions whose events are not one series
    (events.same_series) is site_level; one whose address or title names a year that is
    not its event's is other_year."""
    if "attribution" in ev:
        return _attribution(ev), str(ev.get("source_year") or "")
    a = _attribution(ev)
    if a != "event" or layer not in ("linked_page", "linked_pdf"):
        return a, ""
    try:
        from events import event_years, same_series, source_year
    except ImportError:
        return a, ""
    eds = _parse_editions(m.get("recurring_event_editions", ""))
    if eds and docs is not None:
        name, link = ev.get("name", ""), ev.get("link", "")
        if not any(same_series(name, link, n, docs.link_of(n, d)) for n, d in eds):
            return "site_level", ""
    years = event_years(ev.get("date", ""))
    if years:
        sy = source_year(m.get("url") or m.get("source_link", ""), m.get("page_title", ""), years)
        if sy is not None and sy not in years:
            return "other_year", str(sy)
    return a, ""


def participation(rows: List[dict], docs: DocInfo) -> Counter:
    """Fill doc_kind, doc_kind_reason, participation and context_reason of network rows
    (in place). A row is evidence of taking part in its event (participation=yes)
    unless its source is a linked report or study, it is about another year, it comes
    from other pages of a site several events link to, it is tied to no event, or its
    source is a copy of a document already tied to the same event. Returns the rows
    per reason."""
    reasons: Counter = Counter()
    for r in rows:
        kind, why = docs.classify(r)
        r["doc_kind"], r["doc_kind_reason"] = kind, why
        out = []
        if kind == "report":
            out.append(("report", _CONTEXT_REASONS["report"]))
        att = r.get("attribution") or "event"
        if att == "other_year":
            out.append(("other_year", f"about another year ({r.get('source_year') or '?'})"))
        if att == "site_level" or (r["event"].startswith("(") and r["event"] != NO_EVENT and att != "not_tied"):
            out.append(("site_level", _CONTEXT_REASONS["site_level"]))
        if att == "not_tied" or r["event"] == NO_EVENT:
            out.append(("not_tied", _CONTEXT_REASONS["not_tied"]))
        orig = r.get("duplicate_document_of") or ""
        if orig and _event_id({"key": r.get("event_key"), "name": r["event"], "link": r["event_link"]}) \
                in docs.link_events.get(orig, ()):
            out.append(("copy", f"copy of {orig}"))
        r["participation"] = "no" if out else "yes"
        r["context_reason"] = "; ".join(t for _, t in out)
        for k, _ in out:
            reasons[k] += 1
    return reasons


def network_rows(results: Sequence[Path], min_confidence: str = "low",
                 theme_filter: Optional[ThemeFilter] = None,
                 excluded: Optional[Dict[str, dict]] = None,
                 before: Optional[Dict[tuple, list]] = None,
                 judge: Optional["ThemeFilter"] = None,
                 docs: Optional[DocInfo] = None) -> List[dict]:
    """master_network rows (one per event x organisation x source). Mentions the
    theme filter leaves out are summarised per source in `excluded`; `before`
    receives every row key before the filter: [only in furniture, organisation,
    event id, placeholder event]. `docs` (if given) receives the organisations named
    in each source's content and the events of each source link, before the filter."""
    rows: "OrderedDict[tuple, dict]" = OrderedDict()
    by_key, by_name = _countries(results)
    judge = judge or theme_filter
    n_stopped = [0]
    for res in results:
        conf = {o["org_key"]: o for o in _iter(res / "organisations.csv")}
        # Older mentions.csv files have no source_has_theme column: record the judgement the
        # theme filter used (from sources.csv, theme_hits.csv or the page theme counts).
        can_judge = judge is not None and str(res) not in judge.no_theme_file
        stopped = _stoplist_for(res)
        stop_cache: Dict[str, bool] = {}
        for m in _iter(res / "mentions.csv"):
            org = conf.get(m["org_key"], {})
            confidence = org.get("confidence", "")
            if CONFIDENCE_RANK.get(confidence, 0) < CONFIDENCE_RANK[min_confidence]:
                continue
            nm = m.get("organisation", "")
            if nm not in stop_cache:
                stop_cache[nm] = stopped(nm)
            if stop_cache[nm]:
                n_stopped[0] += 1
                continue
            is_bp = _is_menu_mention(m)
            is_rep = _true(m.get("boilerplate"))
            m_evs = _events_of(m)
            if docs is not None:
                if not is_bp:
                    docs.orgs[m.get("source_unit") or m["url"]].add(m.get("org_key") or nm)
                docs.link_events[m["url"]].update(_event_id(e) for e in m_evs)
            for ev in m_evs:
                if before is not None:
                    b = before.setdefault((_event_id(ev), m["organisation"], m["url"]),
                                          [True, m["organisation"], _event_id(ev),
                                           str(ev.get("name") or NO_EVENT).startswith("(")])
                    b[0] = b[0] and is_bp
                if theme_filter is not None and not theme_filter.keeps(m, ev):
                    if excluded is not None:
                        u = theme_filter.unit_of(m)
                        x = excluded.setdefault(u, {
                            "source_unit": u, "source_link": m.get("url", ""), "site": m.get("site", ""),
                            "site_role": m.get("site_role", ""), "_events": OrderedDict(),
                            "_orgs": OrderedDict(), "reason": theme_filter.reason(m, ev)})
                        x["_events"][ev.get("name", "")] = None
                        x["_orgs"][m["organisation"]] = None
                    continue
                name = ev.get("name") or NO_EVENT
                key = (_event_id(ev), m["organisation"], m["url"])
                r = rows.get(key)
                if r is None:
                    layer = ev.get("layer") or _layer_fallback(m, ev)
                    attribution, s_year = _derived_attribution(m, ev, layer, docs)
                    r = rows[key] = {
                        "event": name, "event_date": ev.get("date", ""), "event_link": ev.get("link", ""),
                        "event_key": ev.get("key", ""), **_country_of(ev, by_key, by_name),
                        "organisation": m["organisation"], "org_key": m.get("org_key", ""),
                        "org_type": m.get("org_type") or org.get("org_type", ""),
                        "confidence": confidence, "mentions_on_source": 0,
                        "source_link": m["url"], "source_unit": m.get("source_unit") or m["url"],
                        "source_title": m.get("page_title", ""),
                        "site": m.get("site", ""), "site_role": m.get("site_role", ""),
                        "layer": layer,
                        "is_pdf": _true(m.get("is_pdf")),
                        "page_date": m.get("page_date", ""), "fetched_at": m.get("fetched_at", ""),
                        "content_vs_event": ev.get("content_vs_event", ""),
                        "years_after_event": ev.get("years_after_event", ""),
                        "recurring_event_editions": m.get("recurring_event_editions", ""),
                        "site_linked_from_events": m.get("site_linked_from_events", ""),
                        "named_in_event_title": (bool(ev["in_title"]) if "in_title" in ev
                                                 else _in_title(m.get("org_key", ""), name)),
                        "source_has_theme": (_true(m["source_has_theme"]) if m.get("source_has_theme")
                                             else judge.unit_of(m) in judge.units if can_judge else ""),
                        "attribution": attribution, "source_year": s_year,
                        "duplicate_document_of": m.get("duplicate_document_of", ""),
                        "translation_of": m.get("translation_of", ""),
                        "doc_kind": "", "doc_kind_reason": "", "participation": "yes", "context_reason": "",
                        "_methods": Counter(), "_bp": True, "_rep": True, "example_context": "", "_ex_bp": True,
                    }
                r["mentions_on_source"] += 1
                r["_methods"].update(m.get("methods", "").split("+"))
                # Menus, headers, footers and sidebars; passages repeated across a site's pages.
                r["_bp"] = r["_bp"] and is_bp
                r["_rep"] = r["_rep"] and is_rep
                if not r["example_context"] or (r["_ex_bp"] and not is_bp):  # prefer real content
                    r["example_context"] = (m.get("context") or "")[:500]
                    r["_ex_bp"] = is_bp
    out = []
    for r in rows.values():
        r["found_by"] = "; ".join(k for k, _ in r.pop("_methods").most_common() if k)
        r["only_in_menus_or_footers"] = r.pop("_bp")
        r["repeated_on_site"] = r.pop("_rep") and not r["only_in_menus_or_footers"]
        r.pop("_ex_bp")
        out.append(r)
    out.sort(key=_network_order)
    if n_stopped[0]:
        log.info("%d mentions of names on the stoplist (the run's --stoplist, else default_stoplist.txt) "
                 "were left out; they come from a run made before those entries existed", n_stopped[0])
    return out


def _network_order(r: dict) -> tuple:
    return (r["event"].startswith("("), r["event_date"] or "", r["event"].casefold(),
            r["organisation"].casefold())


def event_rows(network: List[dict]) -> List[dict]:
    by_event: Dict[tuple, dict] = OrderedDict()
    for r in network:
        key = (r.get("event_key") or "", r["event"], r["event_link"])
        e = by_event.setdefault(key, {"event": r["event"], "event_date": r["event_date"],
                                      "event_link": r["event_link"], "event_key": r.get("event_key", ""),
                                      **{f: r.get(f, "") for f in COUNTRY_FIELDS},
                                      "_orgs": OrderedDict(), "_src": OrderedDict()})
        e["_orgs"][r["organisation"]] = None
        e["_src"][r["source_link"]] = None
    out = []
    for e in by_event.values():
        orgs, srcs = list(e.pop("_orgs")), list(e.pop("_src"))
        out.append({**e, "n_organisations": len(orgs), "organisations": "; ".join(orgs),
                    "n_sources": len(srcs), "sources": " ".join(srcs)})
    return out


def all_event_rows(results: Sequence[Path], network: List[dict], before: Set[str],
                   context: Sequence[dict] = (), candidates: Sequence[dict] = ()) -> List[dict]:
    """Every event of events.csv, with whether (and with how many organisations) it is
    in the network, so shares can be computed over all events, not only those found.
    context_organisations and ner_only_candidates count the event's organisations in
    those sheets; left_out_by_theme_filter is True only for an event with organisations
    before the theme filter and none in the network or either sheet."""
    def ids(rows) -> Dict[str, Set[str]]:
        out_: Dict[str, Set[str]] = defaultdict(set)
        for r in rows:
            out_[_event_id({"key": r.get("event_key"), "name": r["event"], "link": r["event_link"]})].add(
                r["organisation"])
        return out_
    orgs_of, ctx_of, cand_of = ids(network), ids(context), ids(candidates)
    out, seen = [], set()
    by_key, by_name = _countries(results)
    for res in results:
        for e in _iter(res / "events.csv"):
            if "country_iso" not in e:   # older folder: countries from the gazetteer
                e.update(_country_of({"key": e.get("event_key"), "name": e.get("event", ""),
                                      "link": e.get("event_link", "")}, by_key, by_name))
            eid = _event_id({"key": e.get("event_key"), "name": e.get("event", ""),
                             "link": e.get("event_link", "")})
            if eid in seen:
                continue
            seen.add(eid)
            n = len(orgs_of.get(eid, ()))
            n_ctx, n_cand = len(ctx_of.get(eid, ())), len(cand_of.get(eid, ()))
            out.append({**e, "in_network": n > 0, "network_organisations": n,
                        "context_organisations": n_ctx, "ner_only_candidates": n_cand,
                        "left_out_by_theme_filter": eid in before and n == 0 and n_ctx == 0 and n_cand == 0})
    return out


def passage_id(sentence: str) -> str:
    """The same passage (sentence casefolded, whitespace normalised) has the same id."""
    import hashlib
    return hashlib.sha1(" ".join((sentence or "").casefold().split()).encode("utf-8")).hexdigest()[:12]


def theme_rows(results: Sequence[Path], keep_menu_themes: bool = False,
               docs: Optional[DocInfo] = None) -> Tuple[List[dict], List[dict]]:
    """master_themes rows, one per theme term and event (long format), and the rows
    found in menus, sidebars, site furniture and short labels (kept apart unless
    keep_menu_themes). Every row says whether its theme is a focus theme (descriptive
    themes are recorded with focus_theme=no), the document's kind (doc_kind, when
    `docs` is given), how the source was tied to the event, copies, a passage_id and
    how many rows of the same sheet come from the same source (doc_theme_rows)."""
    by_key, by_name = _countries(results)
    out: List[dict] = []
    menus: List[dict] = []
    kinds: Dict[tuple, str] = {}
    for res in results:
        for r in _iter(res / "theme_hits.csv"):
            evs = _events_of(r)
            menu = _is_menu_theme(r)
            pid = passage_id(r.get("sentence", ""))
            for ev in evs:
                layer = ev.get("layer") or _layer_fallback(r, ev)
                attribution, s_year = _derived_attribution(r, ev, layer, docs)
                row = {**r, "event": ev.get("name") or NO_EVENT, "event_date": ev.get("date", ""),
                       "event_link": ev.get("link", ""), "event_key": ev.get("key", ""),
                       **_country_of(ev, by_key, by_name),
                       "n_events_for_passage": r.get("n_events") or len(evs),
                       "layer": layer,
                       "content_vs_event": ev.get("content_vs_event", ""),
                       "years_after_event": ev.get("years_after_event", ""),
                       "menu_or_sidebar": menu,
                       "focus_theme": r.get("focus_theme") or "yes",
                       "attribution": attribution, "source_year": s_year,
                       "duplicate_document_of": r.get("duplicate_document_of", ""),
                       "translation_of": r.get("translation_of", ""),
                       "passage_id": pid}
                if docs is not None:
                    unit = r.get("source_unit") or r.get("source_link", "")
                    kk = (unit, layer)
                    if kk not in kinds:
                        kinds[kk] = docs.classify({**row, "source_title": r.get("page_title", "")})[0]
                    row["doc_kind"] = kinds[kk]
                row.pop("events_json", None)
                (menus if menu and not keep_menu_themes else out).append(row)
    key = lambda r: (r["event"].startswith("("), r.get("event_date") or "", r["event"].casefold(),  # noqa: E731
                     r.get("source_link", ""), r.get("theme", ""))
    for sheet in (out, menus):
        per_doc = Counter(r.get("source_unit") or r.get("source_link", "") for r in sheet)
        for r in sheet:
            r["doc_theme_rows"] = per_doc[r.get("source_unit") or r.get("source_link", "")]
        sheet.sort(key=key)
    return out, menus


def source_rows(results: Sequence[Path], network: List[dict]) -> List[dict]:
    """Every source read: sources.csv (or pages.csv of older folders) plus the site's crawl status."""
    in_net = {r.get("source_unit") or r["source_link"] for r in network}
    out = []
    for res in results:
        status = {s.get("site", ""): s.get("status", "") for s in _iter(res / "sites.csv")}
        src = res / "sources.csv"
        rows = _iter(src) if src.exists() else (
            {**p, "source_unit": p.get("url", ""), "source_link": p.get("url", ""),
             "page_title": p.get("title", ""), "page_date": p.get("page_date") or p.get("date", ""),
             "has_theme_hit": any(int(v or 0) for k, v in p.items() if k.startswith("theme:"))}
            for p in _iter(res / "pages.csv"))
        for r in rows:
            r["site_status"] = status.get(r.get("site", ""), "")
            r["in_network"] = (r.get("source_unit") or r.get("source_link")) in in_net
            out.append(r)
    return out


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------
def _write_csv(path: Path, rows: List[dict], fields: List[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


EXCEL_MAX_ROWS = 1_048_575    # data rows per sheet (one more for the header)


def _write_xlsx(path: Path, sheets: List[tuple]) -> bool:
    """One sheet per (title, rows, fields), streamed (write-only), with a bold
    frozen header, filters, column widths and wrapped long text. A table longer
    than Excel's row limit continues on sheets <title>_2, <title>_3 ..."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl not installed; master files written as CSV only (pip install openpyxl)")
        return False
    wide = {"context_before": 60, "sentence": 60, "context_after": 60, "example_context": 60,
            "organisations": 60, "organisations_in_paragraph": 40, "event": 40, "source_title": 30,
            "page_title": 30, "sources": 50, "source_link": 40, "event_link": 40, "organisation": 30,
            "events": 40, "reason": 40, "value": 80, "listed_on": 40, "source_unit": 40}
    bold = Font(bold=True)
    a_wrap = Alignment(wrap_text=True, vertical="top")
    a_top = Alignment(vertical="top")
    wb = Workbook(write_only=True)
    tmp = path.with_name(path.name + ".part")
    for title, rows, fields in sheets:
        parts = [rows[i:i + EXCEL_MAX_ROWS] for i in range(0, len(rows), EXCEL_MAX_ROWS)] or [[]]
        if len(parts) > 1:
            log.warning("%s: %d rows are more than one Excel sheet holds; continued on %d sheets "
                        "(%s_2 ...). The CSV file has them all in one table.", title, len(rows),
                        len(parts), title)
        aligns = [a_wrap if wide.get(f, 0) >= 40 else a_top for f in fields]
        for n, part in enumerate(parts, 1):
            ws = wb.create_sheet(title if n == 1 else f"{title}_{n}"[:31])
            # Write-only sheets take their layout before the first row.
            for i, f in enumerate(fields, 1):
                ws.column_dimensions[get_column_letter(i)].width = wide.get(f, 16)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(fields)))}{len(part) + 1}"
            hdr = []
            for f in fields:
                c = WriteOnlyCell(ws, value=f)
                c.font = bold
                hdr.append(c)
            ws.append(hdr)
            for r in part:
                cells = []
                for f, a in zip(fields, aligns):
                    v = _cell(r.get(f, ""))
                    c = WriteOnlyCell(ws, value=v)
                    if isinstance(v, str) and v.startswith("="):
                        c.data_type = "s"   # text such as "=== Antifa camp ===", not a formula
                    c.alignment = a
                    cells.append(c)
                ws.append(cells)
    wb.save(tmp)
    tmp.replace(path)
    return True


try:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as _ILLEGAL
except ImportError:   # openpyxl missing: only CSV files are written anyway
    _ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _cell(v):
    if isinstance(v, bool):
        return "yes" if v else "no"
    if v in ("True", "False"):
        return "yes" if v == "True" else "no"
    if isinstance(v, (int, float)):
        return v
    # Control characters (from PDFs, pasted Word text) are not allowed in Excel cells.
    return _ILLEGAL.sub("", str(v))[:32000]   # Excel's cell limit


def _about_rows(results: Sequence[Path], facts: List[Tuple[str, object]]) -> List[dict]:
    rows = [{"item": k, "value": v if isinstance(v, (int, float)) else str(v)} for k, v in facts]
    for res in results:
        p = res / "run_info.json"
        if not p.exists():
            continue
        try:
            info = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        for k in ("started_at", "finished_at", "status", "command", "spacy_models_used", "ner_size",
                  "pages_fetched_between", "packages"):
            if k in info:
                v = info[k]
                rows.append({"item": f"{res.name}: {k}",
                             "value": " ".join(v) if k == "command" and isinstance(v, list)
                             else json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v)})
    return rows


def build_masters(results: Sequence[Path], out_dir: Path, min_confidence: str = "low",
                  theme_filter: str = "source", keep_menu_themes: bool = False,
                  stats: Optional[dict] = None, keep_ner_only: bool = False,
                  network_scope: str = "participation") -> List[Path]:
    """Write master_network and master_themes (.csv and .xlsx) to out_dir and
    return their paths. `stats` (if given) receives the counts of what the theme
    filter kept and left out. Names found only by NER with no other evidence (see
    OrgEvidence) go to the ner_only_candidates sheet unless keep_ner_only. With
    network_scope 'participation', rows that are not evidence of taking part in the
    event (see participation()) go to the context_organisations sheet and
    master_network_context.csv; with 'all' the network keeps every row."""
    if network_scope not in NETWORK_SCOPES:
        raise ValueError(f"--network-scope must be one of {', '.join(NETWORK_SCOPES)}")
    out_dir.mkdir(parents=True, exist_ok=True)
    results = [Path(r) for r in results]
    tf = ThemeFilter(theme_filter, results)
    if tf.old_folders and theme_filter != "none":
        log.info("No sources.csv in %s (made by an older version): for the theme filter a source "
                 "is a whole page or PDF, not an agenda entry", ", ".join(tf.old_folders))
    if tf.from_mentions and theme_filter != "none":
        log.warning("theme_hits.csv not found in %s: the theme filter uses the page theme counts of "
                    "mentions.csv instead (they may include menus and footers), and master_themes has "
                    "no rows from this folder. Use the whole results folder for full theme rows.",
                    ", ".join(tf.from_mentions))
    excluded: Dict[str, dict] = {}
    before: Dict[tuple, list] = {}
    docs = DocInfo(results)
    kept_net = network_rows(results, min_confidence, tf if theme_filter != "none" else None,
                            excluded, before, judge=tf, docs=docs)
    evidence = OrgEvidence(results)
    if evidence.no_pos_folders and not keep_ner_only:
        log.info("organisations.csv in %s has no pos_votes column (made by an older version): "
                 "one-word names found only by NER keep the older rule (3 sources on 2 websites). "
                 "Run site_actors.py --no-crawl on the folder to add the part-of-speech votes",
                 ", ".join(evidence.no_pos_folders))
    n_split = split_coordinated(kept_net, evidence)
    if n_split:
        kept_net.sort(key=_network_order)
        log.info("Split %d network %s naming several organisations at once (\"Exxon & Shell\") "
                 "into one row per organisation; found_by says 'split from ...'", n_split,
                 "row" if n_split == 1 else "rows")
    participation(kept_net, docs)
    net = [r for r in kept_net if not r["only_in_menus_or_footers"]]
    furniture = [r for r in kept_net if r["only_in_menus_or_footers"]]
    tf_net = list(net)   # what the theme filter kept, before the NER-only candidates go
    candidates: List[dict] = []
    if not keep_ner_only:
        cand_ids = set()
        for r in net:
            why = evidence.why_not_row(r)
            if why:
                k = r.get("org_key") or ""
                cand_ids.add(id(r))
                candidates.append({**r, "corpus_sources": len(evidence.sources.get(k, set()) - {""}),
                                   "corpus_sites": len(evidence.sites.get(k, set()) - {""}),
                                   "why_candidate": why})
        if candidates:
            net = [r for r in net if id(r) not in cand_ids]
            log.info("%d network rows (%d organisations) found only by NER with no other evidence "
                     "are in the ner_only_candidates sheet of master_network, not in the network "
                     "(--keep-ner-only on site_actors.py or make_master.py keeps them; a --watchlist entry "
                     "confirms a name)",
                     len(candidates), len({r["organisation"] for r in candidates}))
    # Participation: what is evidence of taking part in the event stays in the network.
    if network_scope == "participation":
        context = [r for r in net if r["participation"] == "no"]
        net = [r for r in net if r["participation"] == "yes"]
    else:
        context = []
    context_counts: Counter = Counter()    # rows flagged per reason (in either scope)
    for r in net + context:
        if r["participation"] != "no":
            continue
        for k, label in _CONTEXT_REASONS.items():
            if r["context_reason"].startswith(label) or f"; {label}" in r["context_reason"]:
                context_counts[k] += 1
    kept_net = net + furniture
    before_content = [b for b in before.values() if not b[0]]
    del before
    before_ids = {b[2] for b in before_content if not b[3]}
    events = event_rows(net)
    all_events = all_event_rows(results, net, before_ids, context=context, candidates=candidates)
    themes, menu_themes = theme_rows(results, keep_menu_themes, docs)
    sources = source_rows(results, kept_net)
    for x in excluded.values():
        evs, orgs = list(x.pop("_events")), list(x.pop("_orgs"))
        x["events"] = "; ".join(evs[:20])
        x["n_organisations"] = len(orgs)
        x["organisations"] = "; ".join(orgs[:40])
    excl = sorted(excluded.values(), key=lambda x: (-x["n_organisations"], x["source_unit"]))

    # What the theme filter kept, for the log and the about sheet.
    n = {
        "network_rows": (len(tf_net), len(before_content)),
        "organisations": (len({r["organisation"] for r in tf_net}), len({b[1] for b in before_content})),
        "events": (len({_event_id({"key": r.get("event_key"), "name": r["event"], "link": r["event_link"]})
                        for r in tf_net if not r["event"].startswith("(")}), len(before_ids)),
    }
    left_out_sources = len({x["source_unit"] for x in excl})
    # What master_network actually holds (after the NER-only step and the network scope):
    # the counts above are taken before them, so they say what the theme filter alone kept.
    real_events = [e for e in events if not e["event"].startswith("(")]
    in_file = {"network_rows_in_file": len(net),
               "organisations_in_file": len({r["organisation"] for r in net}),
               "events_in_file": len(real_events)}
    if theme_filter != "none":
        log.info("Theme filter '%s': kept %d of %d network rows, %d of %d organisations and %d of %d "
                 "events; %d %s with organisations but no theme hit %s left out (listed in the "
                 "excluded_sources sheet). Use --theme-filter none to keep everything.",
                 theme_filter, *n["network_rows"], *n["organisations"], *n["events"], left_out_sources,
                 "source" if left_out_sources == 1 else "sources",
                 "was" if left_out_sources == 1 else "were")
        if before_content and not tf_net:
            if tf.no_theme_file:
                log.warning("theme_hits.csv not found in %s: the theme hits are unknown, so the theme "
                            "filter left master_network empty. Use the whole results folder, or "
                            "--theme-filter none", ", ".join(tf.no_theme_file))
            else:
                log.warning("The theme filter left master_network empty: no source has a theme hit. "
                            "Check the --themes file.")
    if network_scope == "participation":
        log.info("Network scope 'participation': %d rows are evidence of taking part in their event; "
                 "%d rows are in the context_organisations sheet and master_network_context.csv (%s). "
                 "Use --network-scope all to keep every row in the network.",
                 len(net), len(context),
                 ", ".join(f"{_CONTEXT_REASONS[k]}: {v}" for k, v in context_counts.most_common()) or "none")
    n_rows, n_orgs, n_events = in_file.values()
    n_placeholders = len(events) - len(real_events)
    log.info("master_network: %d row%s, %d organisation%s, %d event%s%s%s",
             n_rows, "" if n_rows == 1 else "s", n_orgs, "" if n_orgs == 1 else "s",
             n_events, "" if n_events == 1 else "s",
             (f" plus {n_placeholders} placeholder event{'s' if n_placeholders != 1 else ''} "
              "(other pages of a site, not tied)") if n_placeholders else "",
             "" if theme_filter == "none" and not candidates and network_scope == "all" else " (after %s)" % (
                 ", ".join(x for x in (
                     f"the theme filter '{theme_filter}'" if theme_filter != "none" else "",
                     f"{len(candidates)} NER-only rows moved to ner_only_candidates" if candidates else "",
                     f"{len(context)} context rows moved to context_organisations"
                     if network_scope == "participation" else "") if x)))
    if menu_themes:
        log.info("%d theme rows from menus, sidebars, site furniture and short labels are in the "
                 "menus_and_sidebars sheet of master_themes (--keep-menu-themes keeps them in the main "
                 "sheet)", len(menu_themes))
    summary = {"theme_filter": theme_filter, "network_scope": network_scope,
               "keep_menu_themes": keep_menu_themes, "min_confidence": min_confidence,
               **{f"{k}_kept": a for k, (a, _) in n.items()}, **{f"{k}_before_filter": b for k, (_, b) in n.items()},
               **in_file,
               "placeholder_events_in_file": n_placeholders,
               "context_rows": len(context),
               **{f"context_rows_{k}": context_counts.get(k, 0) for k in _CONTEXT_REASONS},
               "rows_not_participation": sum(1 for r in net + context if r["participation"] == "no"),
               "sources_left_out": left_out_sources,
               "keep_ner_only": keep_ner_only, "ner_only_candidate_rows": len(candidates),
               "ner_only_candidate_organisations": len({r["organisation"] for r in candidates}),
               "coordinated_names_split": n_split, "theme_rows": len(themes),
               "theme_rows_in_menus_and_sidebars": len(menu_themes), "all_events": len(all_events),
               "sources": len(sources)}
    if stats is not None:
        stats.update(summary)
    about = _about_rows(results, [("made_at", dt.datetime.now().astimezone().isoformat(timespec="seconds")),
                                  ("result_folders", "; ".join(str(r) for r in results))]
                        + list(summary.items()))

    # Names the run's vetting dropped (rejected_names.csv of newer folders), for review.
    rejected = [{**x, "result_folder": res.name} for res in results for x in _iter(res / "rejected_names.csv")]
    paths = []
    _write_csv(out_dir / "master_network.csv", net, NETWORK_FIELDS)
    _write_csv(out_dir / "master_network_context.csv", context, NETWORK_FIELDS)
    _write_csv(out_dir / "master_themes.csv", themes, THEME_FIELDS)
    paths += [out_dir / "master_network.csv", out_dir / "master_network_context.csv",
              out_dir / "master_themes.csv"]
    if _write_xlsx(out_dir / "master_network.xlsx", [("event_organisations", net, NETWORK_FIELDS),
                                                      ("events", events, EVENT_FIELDS),
                                                      ("all_events", all_events, ALL_EVENT_FIELDS),
                                                      ("context_organisations", context, NETWORK_FIELDS),
                                                      ("menus_and_footers", furniture, NETWORK_FIELDS),
                                                      ("ner_only_candidates", candidates, CANDIDATE_FIELDS),
                                                      *([("rejected_names", rejected, REJECTED_FIELDS)]
                                                        if rejected else []),
                                                      ("sources", sources, SOURCE_FIELDS),
                                                      ("excluded_sources", excl, EXCLUDED_FIELDS),
                                                      ("about", about, ["item", "value"])]):
        paths.append(out_dir / "master_network.xlsx")
    if _write_xlsx(out_dir / "master_themes.xlsx", [("theme_instances", themes, THEME_FIELDS),
                                                     ("menus_and_sidebars", menu_themes, THEME_FIELDS),
                                                     ("about", about, ["item", "value"])]):
        paths.append(out_dir / "master_themes.xlsx")
    return paths


def main(argv: Optional[List[str]] = None) -> None:
    ap = argparse.ArgumentParser(description="Merge result folders into the two master files")
    ap.add_argument("results", nargs="+", type=Path, help="Result folders written by site_actors.py")
    ap.add_argument("--output-dir", "-o", type=Path, required=True)
    ap.add_argument("--min-confidence", choices=["low", "medium", "high"], default="low",
                    help="Leave out organisations below this confidence in master_network")
    ap.add_argument("--theme-filter", choices=THEME_FILTERS, default="source",
                    help="Keep organisations and events only from sources with a theme hit: the "
                         "source itself ('source': page, PDF or agenda entry); as 'source', but "
                         "agenda entries, calendar cells, the event's own page and its landing page "
                         "also count when any source of their event has a hit ('direct'); any page "
                         "of its website ('site'; agenda entries are still judged one by one), any "
                         "source of the same event ('event'), or keep everything ('none')")
    ap.add_argument("--network-scope", choices=NETWORK_SCOPES, default="participation",
                    help="'participation': the network keeps the rows that are evidence of taking "
                         "part in the event; linked reports and studies, other pages of a site "
                         "several events link to, sources about another year, rows tied to no event "
                         "and copies of a document go to the context_organisations sheet and "
                         "master_network_context.csv. 'all': every row stays in the network (as in "
                         "earlier versions), with the flag columns")
    ap.add_argument("--keep-ner-only", action="store_true",
                    help="Keep organisations found only by NER (spaCy) with no other evidence in the "
                         "network; by default they are listed in the ner_only_candidates sheet")
    ap.add_argument("--keep-menu-themes", action="store_true",
                    help="Keep theme terms found in menus, sidebars and site furniture in the main "
                         "sheet of master_themes")
    args = ap.parse_args(argv)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    handlers: List[logging.Handler] = [logging.StreamHandler()]
    try:
        fh = logging.FileHandler(args.output_dir / "run.log", mode="a", encoding="utf-8")
        fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s make_master: %(message)s",
                                          "%Y-%m-%d %H:%M:%S"))
        handlers.append(fh)
    except OSError:
        pass
    handlers[0].setFormatter(logging.Formatter("%(message)s"))
    logging.basicConfig(level=logging.INFO, handlers=handlers)
    missing = [str(r) for r in args.results if not (r / "mentions.csv").exists()]
    if missing:
        ap.error(f"no mentions.csv in {', '.join(missing)} (not a site_actors.py result folder?)")
    for p in build_masters(args.results, args.output_dir, args.min_confidence, args.theme_filter,
                           args.keep_menu_themes, keep_ner_only=args.keep_ner_only,
                           network_scope=args.network_scope):
        print(p)


if __name__ == "__main__":
    main()
