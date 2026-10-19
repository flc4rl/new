#!/usr/bin/env python3
"""
Analysis tables, charts and network files for a finished run (read-only).

    python analytics.py RESULTS_DIR -o OUT [options]

RESULTS_DIR is a site_actors.py results folder (or a make_master.py output
folder). It must hold master_network (.xlsx or .csv), master_themes (.xlsx or
.csv) and events.csv. Optional files (seed_links.json, pages.csv, sites.csv,
sources.csv) are read from there and from every --run folder (combined).
Without --run, the folders a make_master.py output was built from (its about
sheet, "result_folders") are used when they still exist. Analyses that need a missing file are skipped and listed in summary.json
("skipped").

Nothing in RESULTS_DIR is changed. OUT receives:

  analytics.xlsx           all tables (README sheet first), one sheet each
  tables/*.csv             the same tables as CSV (UTF-8 with BOM, for Excel)
  org_codes_template.xlsx  organisation authority list to curate, merge
                           candidates and a stratified precision sample
  kwic_for_coding.xlsx     theme passages in keyword-in-context layout with
                           coder columns and a stratified coding sample
  network/                 Gephi files (.gexf, .graphml) and node/edge CSVs
  charts/*.png             six charts (150 dpi)
  summary.json, summary.md provenance, parameters, numbers and caveats

Inclusion criterion (--theme-filter; default: the filter make_master.py
recorded in the masters' about sheet, or "source" when none is recorded):
organisations and events count for the network only when they come from a
source with a hit of a focus theme (colonialism, racism, fascism) in its
content. A source is one web page or one PDF; for the agenda (seed) PDFs each
agenda entry is its own source. "direct" also keeps the listing, the event's own
page and its landing page when any source of their event has a hit; "source"
asks every source for its own hit; "site" keeps every source of a website with a
hit, "event" every source of an event with a hit, "none" keeps all sources. The
sensitivity table always also shows "none", and the number of organisations and
events left out is reported. Theme prevalence is never filtered (an inclusion
criterion on the outcome would make every share 100 %). Descriptive themes
(marked "#: descriptive" in the theme file, or focus_theme=no in the masters:
global_south, authoritarianism) are reported in descriptive_themes only.

Network scope (--network-scope, default "participation"): the default network
uses only rows that are evidence of taking part in the event. Linked reports and
studies (doc_kind report), other pages of a site several events link to, rows
about another year and copies of a document are context (context_organisations
sheet of newer masters; derived here for older ones). "all" uses every row. Tie
levels (--layers, default L1-L3s): L1 the listing, L2 the landing page, L3s
linked pages and PDFs on the event's own host, L3x those on other hosts.

Both master layouts are read: the older one (events joined with " | ") and the
long one (one row per event, with event_key, layer and country columns).

Curated inputs (read, never written): --org-codes org_codes.csv (columns
org_key, organisation, include, merge_into_key, canonical_name, actor_class,
country, is_organisation, is_partial_name, notes) and --event-codes
event_codes.csv (event_id, series_id, duplicate_of, country_override, notes).
Fill in org_codes_template.xlsx, save its first sheet as org_codes.csv and run
again with --org-codes.

    python analytics.py results/movementhub -o results/movementhub_analytics
    python analytics.py results/master -o results/analytics --run results/pdf_agendas
    python analytics.py results/master -o results/analytics --theme-filter none
    python analytics.py results/master -o results/analytics --network-scope all --layers L1-L3
    python analytics.py results/master -o results/analytics --source-codes results/analytics/org_codes_template.xlsx
    python analytics.py --selftest
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import hashlib
import json
import logging
import math
import os
import platform
import random
import re
import sys
import time
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple
from urllib.parse import unquote, urlparse

log = logging.getLogger("analytics")
csv.field_size_limit(10_000_000)

VERSION = "1.0"
HERE = Path(__file__).resolve().parent
DEFAULT_THEMES_FILE = HERE / "themes_colonialism_racism_fascism.txt"

THEME_FILTERS = ["direct", "source", "site", "event", "none"]
# Layers that document the event itself: under the 'direct' theme filter they count when
# any source of their event has a theme hit (as in make_master.py).
DIRECT_LAYERS = {"listing_seed", "listing_calendar", "main_event_page", "landing"}
NETWORK_SCOPES = ["participation", "all"]
DOC_KINDS = ["listing", "page", "programme", "report"]
# Bins of words read on the web for an event (exposure), for standardised prevalence.
WORDS_BINS = [(1000, "<1k"), (5000, "1-5k"), (20000, "5-20k"), (100000, "20-100k"), (None, ">=100k")]
CONF_RANK = {"low": 0, "medium": 1, "high": 2}
FURNITURE_REGIONS = {"nav", "header", "footer", "aside"}
NO_EVENT = "(not tied to an event)"
OWN_LAYERS = {"listing_seed", "listing_calendar", "main_event_page"}
TIE_LEVEL = {"listing_seed": "L1", "listing_calendar": "L1", "main_event_page": "L1", "main_page": "L1",
             "landing": "L2", "linked_pdf": "L3", "linked_page": "L3"}
LAYER_GROUP = {"L1": "own", "L2": "landing", "L3": "linked"}
PERIODS = [(2015, 2017, "2015-17"), (2018, 2020, "2018-20"), (2021, 2022, "2021-22"),
           (2023, 2024, "2023-24"), (2025, 2026, "2025-26")]
DROP_CLASSES = {"place", "site_ui", "noise", "event_name"}
ACTOR_CLASSES = ["movement_org", "ngo", "party", "union", "media", "company", "state_body", "police_court",
                 "place", "person", "project_infrastructure", "event_name", "reference", "site_ui", "noise"]
RELEVANCE = ["substantive", "name_or_title_only", "list_or_tag", "off_topic_sense", "false_match", "unclear"]
SCENARIOS = {"A": "all organisations", "B": "without organisations named in the event title",
             "C": "confidence >= --min-confidence (or include=1)", "D": "C without organisations named in the title",
             "C_curated": "C without curated places, site UI, noise and event names"}
# L3 is split by host: L3s = linked pages and PDFs on the event's own host (or a
# subdomain of it, or it of them), L3x = other hosts. L1-L3s is the default.
LAYER_SETS = {"L1-L3s": {"L1", "L2", "L3s"}, "L1-L3": {"L1", "L2", "L3"}, "L1+L2": {"L1", "L2"}}
DEFAULT_LAYERS = "L1-L3s"
SUBLEVELS = ["L1", "L2", "L3s", "L3x"]

# Colour-blind-safe categorical slots (validated reference palette): themes take
# slots 1-3, the three agenda streams slots 4, 7 and 6, always with a legend.
THEME_COLOUR = {"colonialism": "#2a78d6", "racism": "#eb6834", "fascism": "#1baf7a"}
STREAM_COLOUR = {"cja": "#eda100", "echos": "#6250d6", "hub": "#008300"}
STREAM_MARKER = {"cja": "o", "echos": "s", "hub": "^"}
EXTRA_COLOURS = ["#e87ba4", "#e34948", "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#6250d6", "#008300"]
INK, INK2, GRID, SURFACE = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_WS = re.compile(r"\s+")
_WORD = re.compile(r"[^\W\d_]+")


# ==========================================================================
# Small helpers
# ==========================================================================
def s(v) -> str:
    """A cell value as text ("" for None and NaN; 3.0 -> "3"; booleans -> yes/no)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def truthy(v) -> bool:
    return s(v).casefold() in ("yes", "true", "1", "y")


def yn(b) -> str:
    return "yes" if b else "no"


def sha12(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:12]


def collapse(text: str) -> str:
    return _WS.sub(" ", text or "").strip()


def count_of(k, word: str, plural: str = "") -> str:
    """'1 organisation', '2 organisations' (for user-facing sentences)."""
    return f"{k} {word if k == 1 else (plural or word + 's')}"


def wilson(k: int, n: int, z: float = 1.96) -> Tuple[Optional[float], Optional[float]]:
    if not n:
        return None, None
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return round(max(0.0, (c - h) / d), 4), round(min(1.0, (c + h) / d), 4)


def byar(k: int, z: float = 1.96) -> Tuple[float, float]:
    """Byar's approximation of the 95 % Poisson interval for a count k."""
    lo = 0.0 if k == 0 else k * (1 - 1 / (9 * k) - z / (3 * math.sqrt(k))) ** 3
    k1 = k + 1
    hi = k1 * (1 - 1 / (9 * k1) + z / (3 * math.sqrt(k1))) ** 3
    return lo, hi


def share(k, n) -> Optional[float]:
    return round(k / n, 4) if n else None


def norm_link(u: str) -> str:
    """Link without scheme, www., fragment and trailing slash (for comparisons)."""
    u = s(u)
    u = re.sub(r"^https?://(?:www\.)?", "", u.split("#")[0], flags=re.I)
    return u.rstrip("/").casefold()


def norm_doc(u: str) -> str:
    """Document identity of a link: host lower-cased, no scheme, no www., no fragment,
    no trailing slash (the path keeps its case). Other values (agenda entries
    'pdf:...#entry-3') are returned unchanged."""
    u = s(u)
    m = re.match(r"(?i)^https?://([^/?#]*)(.*)$", u)
    if not m:
        return u
    host = m.group(1).lower()
    host = host[4:] if host.startswith("www.") else host
    rest = m.group(2).split("#")[0].rstrip("/")
    return host + rest


def same_host(a: str, b: str) -> bool:
    """Same host, or one a subdomain of the other (www. removed)."""
    return bool(a and b) and (a == b or a.endswith("." + b) or b.endswith("." + a))


def host_of(u: str) -> str:
    try:
        h = (urlparse(s(u)).hostname or "").lower()
    except ValueError:
        return ""
    return h[4:] if h.startswith("www.") else h


def is_web(u: str) -> bool:
    return bool(re.match(r"https?://", s(u), re.I))


def _join_entry_lines(lines: List[str]) -> str:
    """Join the wrapped lines of an agenda entry as site_actors.seed_blocks does: a
    line ending in a hyphen right after a non-space continues on the next line
    without a space (so a URL broken by a line wrap stays one URL and its tail is
    not searched for theme terms)."""
    txt = ""
    for ln in lines:
        ln = ln.strip()
        if txt.endswith("-") and len(txt) > 1 and not txt[-2].isspace():
            txt += ln
        else:
            txt = f"{txt} {ln}" if txt else ln
    return txt


def year_of(date: str) -> Optional[int]:
    m = re.search(r"(?<!\d)((?:19|20)\d\d)(?!\d)", s(date))
    return int(m.group(1)) if m else None


def period_of(year: Optional[int]) -> str:
    if year is None:
        return "undated"
    for a, b, lab in PERIODS:
        if a <= year <= b:
            return lab
    return "before 2015" if year < PERIODS[0][0] else "after 2026"


def doc_year_of(url: str) -> Optional[int]:
    p = unquote(s(url))
    m = re.search(r"/(?:uploads|files)/((?:19|20)\d\d)/\d\d/", p) or re.search(r"/((?:19|20)\d\d)/\d\d/", p)
    return int(m.group(1)) if m else None


def pdf_display(value: str) -> str:
    """Seed-PDF paths as the pipeline shows them ("pdf:<file name>")."""
    try:
        from seed_pdfs import pdf_display as pd_
        return pd_(value)
    except Exception:  # noqa: BLE001 - optional project module
        if "file:" not in (value or ""):
            return value

        def one(m):
            p = urlparse(m.group(0))
            name = unicodedata.normalize("NFC", unquote(Path(p.path).name))
            return f"pdf:{name}" + (f"#{p.fragment}" if p.fragment else "")
        return re.sub(r"file://[^\s|;#]*(?:#[^\s|;]*)?", one, value)


def stream_of(source: str) -> str:
    """Which listing an event comes from: cja (Riseup pad), echos (Les échos du
    vent agenda), hub (a web calendar), or agenda:<file> for other seed PDFs."""
    v = unicodedata.normalize("NFC", unquote(s(source))).casefold()
    if not v:
        return "other"
    if "chos" in v:
        return "echos"
    if "riseup" in v or "cja" in v:
        return "cja"
    if v.startswith(("http://", "https://")):
        return "hub"
    name = Path(urlparse(v).path).name or v
    return "agenda:" + name.replace("pdf:", "")


def short(text: str, n: int = 60) -> str:
    text = collapse(text)
    return text if len(text) <= n else text[: n - 3] + "..."


# ==========================================================================
# Project modules (read-only, imported lazily so a module being edited does
# not stop the analysis)
# ==========================================================================
class Project:
    def __init__(self):
        import extract
        global _NER_LIKE
        self.extract = extract
        _NER_LIKE = frozenset(getattr(extract, "NER_LIKE_METHODS", None) or _NER_LIKE)
        self.canonical_key = extract.canonical_key
        self.URL_LIKE = getattr(extract, "URL_LIKE", re.compile(r"(?:https?://|www\.)\S+"))
        self.detect_language = getattr(extract, "detect_language", None)
        try:
            import linked_sites
            self.platform_hosts = set(linked_sites.PLATFORM_HOSTS)
            self.host_matches = linked_sites.host_matches
            self.filter_targets = linked_sites.filter_targets
        except Exception as exc:  # noqa: BLE001
            log.warning("linked_sites not importable (%s); platform hosts unknown", exc)
            self.platform_hosts, self.filter_targets = set(), None
            self.host_matches = lambda h, names: any(h == n or h.endswith("." + n) for n in names)
        try:
            import places
            self.places = set(places.PLACES)
        except Exception as exc:  # noqa: BLE001
            log.warning("places not importable (%s)", exc)
            self.places = set()
        self.gazetteer = None
        self.gazetteer_error = ""
        try:
            import gazetteer
            self.gazetteer = gazetteer
        except Exception as exc:  # noqa: BLE001
            self.gazetteer_error = f"{type(exc).__name__}: {exc}"
        self._ck: Dict[str, str] = {}

    def key(self, name: str) -> str:
        k = self._ck.get(name)
        if k is None:
            try:
                k = self.canonical_key(name or "")
            except Exception:  # noqa: BLE001
                k = collapse(name).casefold()
            self._ck[name] = k
        return k

    def country(self, name: str, entry_text: str = "", link: str = "", structured: str = "") -> dict:
        """gazetteer.event_country (both the current dict form and the older tuple form)."""
        if self.gazetteer is None:
            return {}
        fn = self.gazetteer.event_country
        try:
            res = fn(name or "", entry_text or "", link=link or "", structured=structured or "")
        except TypeError:
            res = fn(name or "", entry_text or "")
        except Exception:  # noqa: BLE001 - one odd name must not stop the analysis
            return {}
        if isinstance(res, tuple):
            res = {"countries": list(res[0] or []), "method": res[1] if len(res) > 1 else "", "matched": ""}
        return dict(res or {})

    def country_name(self, iso: str) -> str:
        if self.gazetteer is not None and hasattr(self.gazetteer, "country_name"):
            try:
                return self.gazetteer.country_name(iso)
            except Exception:  # noqa: BLE001
                pass
        return {"INT": "International", "ONLINE": "Online"}.get(iso, iso)

    def is_platform(self, link: str) -> bool:
        """Platform, social-media, pad and form hosts: their links do not make a series."""
        h = host_of(link)
        if not h:
            return False
        if self.host_matches(h, self.platform_hosts):
            return True
        skip = getattr(self.extract, "is_skip_host", None)
        try:
            return bool(skip(h)) if skip else False
        except Exception:  # noqa: BLE001
            return False


# ==========================================================================
# Reading
# ==========================================================================
def _sheet_rows(ws) -> Tuple[List[str], List[dict]]:
    it = ws.iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        return [], []
    fields = [s(h) for h in head]
    rows = []
    for vals in it:
        if vals is None or all(v is None for v in vals):
            continue
        rows.append({f: s(v) for f, v in zip(fields, vals) if f})
    return fields, rows


def read_master(folder: Path, stem: str, sheets: Sequence[str]) -> Tuple[Optional[Path], Dict[str, Tuple[List[str], List[dict]]]]:
    """The named sheets of <stem>.xlsx (missing sheets are empty), else <stem>.csv as
    the first sheet."""
    x = folder / f"{stem}.xlsx"
    out: Dict[str, Tuple[List[str], List[dict]]] = {sh: ([], []) for sh in sheets}
    if x.exists() and x.stat().st_size:
        import openpyxl
        wb = openpyxl.load_workbook(x, read_only=True, data_only=True)
        try:
            for i, sh in enumerate(sheets):
                ws = wb[sh] if sh in wb.sheetnames else (wb.worksheets[0] if i == 0 else None)
                if ws is not None:
                    out[sh] = _sheet_rows(ws)
        finally:
            wb.close()
        return x, out
    c = folder / f"{stem}.csv"
    if c.exists():
        out[sheets[0]] = read_csv(c)
        return c, out
    return None, out


def read_csv(path: Path) -> Tuple[List[str], List[dict]]:
    if not path.exists() or not path.stat().st_size:
        return [], []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rd = csv.DictReader(fh)
        rows = [{k: s(v) for k, v in r.items() if k} for r in rd]
        return list(rd.fieldnames or []), rows


def sha256_of(path: Optional[Path]) -> str:
    if not path or not Path(path).exists():
        return ""
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ==========================================================================
# Writing (every file goes to *.tmp first and is then renamed, so a crash
# never leaves a half-written or 0-byte file)
# ==========================================================================
def _atomic(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    return path.with_name(path.name + ".tmp")


def write_text(path: Path, text: str) -> None:
    tmp = _atomic(path)
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def _csv_value(v):
    if isinstance(v, bool):
        return yn(v)
    if v is None:
        return ""
    if isinstance(v, float):
        return "" if math.isnan(v) else (str(int(v)) if v.is_integer() and abs(v) < 1e15 else f"{v:.6g}")
    return v


def write_csv(path: Path, rows: List[dict], fields: List[str]) -> None:
    tmp = _atomic(path)
    with tmp.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(fields)
        for r in rows:
            w.writerow([_csv_value(r.get(f, "")) for f in fields])
    os.replace(tmp, path)


def _xl(v):
    if isinstance(v, bool):
        return yn(v)
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    if isinstance(v, (int, float)):
        return v
    return _CTRL.sub("", str(v))[:32000]   # Excel's cell limit (as make_master._cell)


WIDE = {"event": 40, "organisation": 30, "sentence": 60, "context_before": 50, "context_after": 50,
        "left_ctx": 45, "right_ctx": 45, "source_link": 40, "example_context_1": 50, "example_context_2": 50,
        "variants": 30, "item": 28, "value": 90, "text": 100, "memo": 40, "notes": 30, "top1_doc": 50,
        "sites": 30, "description": 70, "name_a": 30, "name_b": 30, "top_event": 40, "document": 40}


def write_xlsx(path: Path, sheets: List[Tuple[str, List[dict], List[str]]],
               validations: Optional[Dict[str, Dict[str, List[str]]]] = None) -> None:
    """One sheet per (title, rows, fields), streamed (write-only): bold frozen header,
    filters, widths. `validations` = {sheet: {column: [allowed values]}} adds
    dropdown lists (from a hidden 'lists' range when given)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook(write_only=True)
    for title, rows, fields in sheets:
        ws = wb.create_sheet(title[:31])
        for i, f in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDE.get(f, 14)
        ws.freeze_panes = "A2"
        last = get_column_letter(max(1, len(fields)))
        if rows:
            ws.auto_filter.ref = f"A1:{last}{len(rows) + 1}"
        for col, allowed in ((validations or {}).get(title) or {}).items():
            if col not in fields:
                continue
            letter = get_column_letter(fields.index(col) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{max(len(rows), 1) + 500}")
            ws.data_validations.append(dv)   # write-only sheets keep a plain list
        hdr = []
        for f in fields:
            c = WriteOnlyCell(ws, value=f)
            c.font = bold
            hdr.append(c)
        ws.append(hdr)
        # Wrapped text only on smaller sheets: styled cells are slow to write, and a
        # 20,000-row sheet is read with filters rather than scrolled through.
        wrap_cols = {f for f in fields if WIDE.get(f, 0) >= 40} if len(rows) <= 5000 else set()
        for r in rows:
            cells = []
            for f in fields:
                v = _xl(r.get(f, ""))
                if f in wrap_cols or (isinstance(v, str) and v.startswith("=")):
                    c = WriteOnlyCell(ws, value=v)
                    if isinstance(v, str) and v.startswith("="):
                        c.data_type = "s"          # text, never a formula
                    if f in wrap_cols:
                        c.alignment = wrap
                    v = c
                cells.append(v)
            ws.append(cells)
    tmp = _atomic(path)
    wb.save(tmp)
    os.replace(tmp, path)


def clean_graph(G):
    """A copy whose attributes the GEXF and GraphML writers accept: lists and sets
    become "; "-joined text, control characters are removed, and every attribute
    gets one type (GEXF refuses mixed types). A numeric attribute with gaps
    becomes a float with NaN for the gaps; any other attribute with None/NaN or
    mixed values becomes text with "" for the gaps."""
    import networkx as nx

    def missing(v):
        return v is None or (isinstance(v, float) and math.isnan(v)) or v == ""

    def plain(v):
        if isinstance(v, bool):
            return yn(v)
        if isinstance(v, (list, tuple, set, frozenset)):
            v = "; ".join(str(x) for x in (sorted(v) if isinstance(v, (set, frozenset)) else v))
        if isinstance(v, (int, float)):
            return v
        return _CTRL.sub("", str(v))

    def fixers(items):
        kinds: Dict[str, Set[str]] = defaultdict(set)
        for d in items:
            for k, v in d.items():
                v = plain(v) if not missing(v) else None
                kinds[k].add("none" if v is None else "float" if isinstance(v, float) else
                             "int" if isinstance(v, int) else "str")
        out = {}
        for k, ks in kinds.items():
            if ks <= {"int"}:
                out[k] = lambda v: v
            elif ks <= {"int", "float", "none"}:
                out[k] = lambda v: float("nan") if missing(v) else float(v)
            else:
                out[k] = lambda v: "" if missing(v) else str(plain(v))
        return out
    nf = fixers([d for _, d in G.nodes(data=True)])
    ef = fixers([d for _, _, d in G.edges(data=True)])
    H = nx.Graph()
    for n, d in G.nodes(data=True):
        H.add_node(_CTRL.sub("", str(n)), **{k: nf[k](v) for k, v in d.items()})
    for a, b, d in G.edges(data=True):
        H.add_edge(_CTRL.sub("", str(a)), _CTRL.sub("", str(b)), **{k: ef[k](v) for k, v in d.items()})
    return H


def write_graph(G, base: Path) -> List[Path]:
    import networkx as nx
    H = clean_graph(G)
    out = []
    for ext, fn in ((".gexf", nx.write_gexf), (".graphml", nx.write_graphml)):
        p = base.with_suffix(ext)
        tmp = _atomic(p)
        fn(H, str(tmp))
        os.replace(tmp, p)
        out.append(p)
    return out


# ==========================================================================
# The analysis: inputs and core tables
# ==========================================================================
class InputError(Exception):
    pass


# Words in a source's address or title that make a long list of organisations a
# programme, partner list or call (the F2 rule for masters without doc_kind).
_PROGRAMME_HINT = re.compile(r"program|programm|agenda|line-?up|speaker|referent|intervenant|partners?\b|partenaires|"
                             r"unterst[uü]tzer|signatair|signator|ondertekenaar|aufruf|appel|call-to|b[uü]ndnis|"
                             r"who-we-are|wer-wir-sind|qui-sommes-nous", re.I)
CONTEXT_KINDS = OrderedDict([("linked report", "report"), ("about another year", "other_year"),
                             ("other pages", "site_level"), ("not tied", "not_tied"), ("copy of", "copy"),
                             ("left out by --source-codes", "source_codes")])


def _reason_kind(reason: str) -> str:
    for prefix, kind in CONTEXT_KINDS.items():
        if reason.startswith(prefix):
            return kind
    return "other"


class Analysis:
    """Reads one results folder and builds the analysis-ready tables (events,
    ties, theme units) that every analysis below works on."""

    def __init__(self, args, project: Project):
        self.a = args
        self.p = project
        self.results = Path(args.results)
        self.runs = [Path(r) for r in (args.run or [])]
        self.inputs: Dict[str, Path] = {}
        self.all_inputs: Dict[str, List[Path]] = {}
        self.skipped: Dict[str, str] = {}
        self.warnings: List[str] = []
        self.diag: List[dict] = []
        self.tables: "OrderedDict[str, Tuple[List[dict], List[str]]]" = OrderedDict()
        self.facts: Dict[str, object] = {}
        self.common_words: Set[str] = set()
        self.source_override: Dict[str, bool] = {}
        self.source_codes: List[dict] = []

    # ------------------------------------------------------------ helpers
    def warn(self, msg: str, *args) -> None:
        text = msg % args if args else msg
        log.warning(text)
        self.warnings.append(text)

    def note(self, item: str, value, note: str = "") -> None:
        self.diag.append({"item": item, "value": value, "note": note})

    def table(self, name: str, rows: List[dict], fields: List[str]) -> None:
        self.tables[name] = (rows, fields)

    def optional(self, name: str) -> Optional[Path]:
        """First copy of an optional file in RESULTS_DIR or the --run folders. Every
        copy found is kept in self.all_inputs[name], so that several --run folders
        (masters merged by make_master.py) are all read; later copies are also listed
        in summary.json as "<name> [2]", "<name> [3]" and so on."""
        found: List[Path] = []
        for d in [self.results] + self.runs:
            p = d / name
            if p.exists() and p.stat().st_size and p.resolve() not in {q.resolve() for q in found}:
                found.append(p)
        if not found:
            return None
        self.all_inputs[name] = found
        self.inputs[name] = found[0]
        for i, p in enumerate(found[1:], 2):
            self.inputs[f"{name} [{i}]"] = p
        return found[0]

    def optional_rows(self, name: str) -> List[dict]:
        """Rows of every copy of an optional CSV file (see optional())."""
        return [r for p in self.all_inputs.get(name, []) for r in read_csv(p)[1]]

    # ------------------------------------------------------------ loading
    def load(self) -> None:
        a = self.a
        path, net = read_master(self.results, "master_network", ["event_organisations", "menus_and_footers",
                                                                 "all_events", "sources", "excluded_sources", "about",
                                                                 "context_organisations"])
        if path is None:
            raise InputError(f"master_network.xlsx or master_network.csv not found in {self.results}")
        self.inputs["master_network"] = path
        path, th = read_master(self.results, "master_themes", ["theme_instances", "menus_and_sidebars"])
        if path is None:
            raise InputError(f"master_themes.xlsx or master_themes.csv not found in {self.results}")
        self.inputs["master_themes"] = path
        ev_path = self.results / "events.csv"
        if ev_path.exists():
            self.inputs["events.csv"] = ev_path
            self.ev_fields, self.ev_rows = read_csv(ev_path)
        elif net["all_events"][1]:
            # A make_master.py folder without events.csv: its all_events sheet lists every event.
            self.ev_fields, self.ev_rows = net["all_events"]
            self.note("events table", "master_network all_events sheet", "events.csv not found")
        else:
            raise InputError(f"events.csv not found in {self.results} (and master_network has no all_events sheet)")
        self.net_fields, self.net_rows = net["event_organisations"]
        self.menu_rows = net["menus_and_footers"][1]
        # Rows that are not evidence of taking part in the event (newer masters): linked
        # reports, other pages of a site, rows about another year, copies. A CSV master
        # keeps them in master_network_context.csv.
        self.ctx_fields, self.ctx_rows = net["context_organisations"]
        if not self.ctx_rows and path.suffix == ".csv":
            self.ctx_fields, self.ctx_rows = read_csv(self.results / "master_network_context.csv")
        if self.ctx_rows:
            self.inputs["context_organisations"] = path if path.suffix == ".xlsx" else (
                self.results / "master_network_context.csv")
        self.th_fields, self.th_rows = th["theme_instances"]
        self.th_menu_rows = th["menus_and_sidebars"][1]
        self.sources_sheet = net["sources"][1]
        self.excluded_sources = net["excluded_sources"][1]
        # The theme filter make_master.py already applied (newer masters say so in their about sheet).
        about = {r.get("item", ""): r.get("value", "") for r in net["about"][1]}
        self.master_filter = about.get("theme_filter", "") or ("none" if not self.excluded_sources else "")
        self.master_about = about
        if a.theme_filter is None:
            # Default: the filter make_master.py recorded, so that masters built with
            # --theme-filter none are analysed unfiltered; older masters: source.
            rec = s(about.get("theme_filter"))
            a.theme_filter = rec if rec in THEME_FILTERS else "source"
            how = ("as recorded by make_master.py in the masters' about sheet" if rec in THEME_FILTERS
                   else "default; the masters do not record a filter")
            log.info("Theme filter '%s' (%s; set --theme-filter to change it)", a.theme_filter, how)
            self.note("theme filter used here", a.theme_filter, how)
        if not self.runs and about.get("result_folders"):
            # A make_master.py output folder holds only the masters: look for the
            # per-run files (seed_links.json, pages.csv ...) in the folders it was built from.
            for d in (x.strip() for x in s(about.get("result_folders")).split(";")):
                if d and Path(d).is_dir() and Path(d).resolve() != self.results.resolve():
                    self.runs.append(Path(d))
            if self.runs:
                log.info("Run folders taken from the masters' about sheet: %s", "; ".join(map(str, self.runs)))
                self.note("run folders (from master about sheet)", "; ".join(map(str, self.runs)),
                          "give --run to use other folders")
        need_net = {"event", "event_date", "event_link", "organisation", "confidence", "source_link", "site",
                    "site_role"}
        need_th = {"theme", "term", "event", "event_date", "event_link", "source_link", "sentence"}
        for label, fields, need, rows in (("master_network", self.net_fields, need_net, self.net_rows),
                                          ("master_themes", self.th_fields, need_th, self.th_rows)):
            miss = need - set(fields)
            if fields and miss:
                raise InputError(f"{label}: required columns missing: {', '.join(sorted(miss))}")
        if not {"event", "date", "event_link"} <= set(self.ev_fields):
            raise InputError("events.csv: needs the columns event, date, event_link")
        self.long_net = "event_key" in self.net_fields
        self.long_th = "event_key" in self.th_fields
        self.note("layout master_network", "long (event_key)" if self.long_net else "older (' | '-joined)")
        self.note("layout master_themes", "long (event_key)" if self.long_th else "older (' | '-joined)")
        self.ts = None
        tf = Path(a.themes_file) if a.themes_file else DEFAULT_THEMES_FILE
        if tf.exists():
            self.inputs["themes_file"] = tf
            self.ts = self.p.extract.ThemeSet.load(tf)
            self.sections = self._theme_sections(tf)
        else:
            self.warn("theme file %s not found: term families, title flags and entry rates are skipped", tf)
            self.sections = {}
        names = list(self.ts.names) if self.ts else []
        for r in self.th_rows:
            if r.get("theme") and r["theme"] not in names:
                names.append(r["theme"])
        # Focus themes drive the theme filter and every prevalence and any-theme measure;
        # descriptive themes (global_south, authoritarianism) are reported on their own.
        # The rows say which is which (focus_theme); older masters: the '#: descriptive'
        # marks of the theme file.
        if "focus_theme" in self.th_fields:
            desc = {r.get("theme", "") for r in self.th_rows + self.th_menu_rows
                    if s(r.get("focus_theme")).casefold() == "no"}
            focus = {r.get("theme", "") for r in self.th_rows + self.th_menu_rows
                     if s(r.get("focus_theme")).casefold() != "no"}
            # themes without rows follow the file's marks
            desc |= set(getattr(self.ts, "descriptive", ()) or ()) - focus if self.ts else set()
            how = "focus_theme column of master_themes"
        else:
            desc = set(getattr(self.ts, "descriptive", ()) or ()) if self.ts else set()
            how = "'#: descriptive' marks of the theme file" if desc else "no descriptive themes"
        desc.discard("")
        self.descriptive = [x for x in names if x in desc]
        self.themes = [x for x in names if x not in desc] or ["colonialism", "racism", "fascism"]
        self.note("focus themes", "; ".join(self.themes))
        self.note("descriptive themes (reported apart)", "; ".join(self.descriptive) or "none", how)
        for name in ("seed_links.json", "pages.csv", "sites.csv", "sources.csv", "summary.json"):
            self.optional(name)
        self.sources_rows = (self.optional_rows("sources.csv") if "sources.csv" in self.inputs
                             else self.sources_sheet)
        self._check_master_filter()
        self.org_codes = self._read_codes(a.org_codes, "org_codes") if a.org_codes else []
        self.event_codes = self._read_codes(a.event_codes, "event_codes") if a.event_codes else []
        self.source_codes = self._read_codes(a.source_codes, "source_codes", "sources_to_code") \
            if getattr(a, "source_codes", None) else []
        self.curated = bool(self.org_codes)

    def _check_master_filter(self) -> None:
        """make_master.py (newer versions) applies the inclusion criterion itself, so
        sources left out there cannot be brought back here. Say so when a weaker
        filter is asked for."""
        mf, tf = self.master_filter, self.a.theme_filter
        if mf:
            self.note("theme filter applied by make_master.py", mf,
                      f"{len(self.excluded_sources)} sources listed in excluded_sources")
        # make_master.py's own counts (about sheet): what its filter kept of what was found
        self.master_counts = {}
        for what in ("network_rows", "organisations", "events"):
            kept, before = s(self.master_about.get(f"{what}_kept")), s(self.master_about.get(f"{what}_before_filter"))
            in_file = s(self.master_about.get(f"{what}_in_file"))
            if kept.isdigit() and before.isdigit():
                self.master_counts[what] = (int(kept), int(before), int(in_file) if in_file.isdigit() else None)
                self.note(f"make_master.py theme filter: {what} kept", f"{kept}/{before}",
                          f"{int(before) - int(kept)} left out before analytics")
        if mf and mf != "none" and tf != mf:
            strength = {"none": 0, "site": 1, "event": 1, "direct": 1.5, "source": 2}
            if strength.get(tf, 0) <= strength.get(mf, 0):
                self.warn("master_network was built with --theme-filter %s, so organisations from sources "
                          "without a theme hit are not in it. --theme-filter %s here cannot bring them back; "
                          "rebuild the masters with make_master.py --theme-filter none first", mf, tf)

    def _read_codes(self, path: str, label: str, sheet: str = "") -> List[dict]:
        p = Path(path)
        if not p.exists():
            raise InputError(f"--{label.replace('_', '-')}: {p} not found")
        self.inputs[label] = p
        if p.suffix.lower() == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            try:
                ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
                rows = _sheet_rows(ws)[1]
                if label == "org_codes" and "no_series" in wb.sheetnames and ws.title != "no_series":
                    rows += _sheet_rows(wb["no_series"])[1]    # the template's second organisation sheet
                return rows
            finally:
                wb.close()
        return read_csv(p)[1]

    @staticmethod
    def _theme_sections(path: Path) -> Dict[Tuple[str, str], str]:
        sec, cur, out = "", None, {}
        for line in path.read_text(encoding="utf-8").splitlines():
            t = line.strip()
            if t.startswith("# ---"):
                sec = t[5:].split(":")[0].strip()
            elif re.fullmatch(r"\[(.+)\]", t):
                cur = t[1:-1].strip()
            elif t and not t.startswith(("#", "!")) and cur:
                out[(cur, t)] = sec
        return out

    # ------------------------------------------------------------ events
    def build_events(self) -> None:
        p = self.p
        self.events: List[dict] = []
        self.ev_by_id: Dict[str, dict] = {}
        self.by_key: Dict[str, str] = {}
        self.by_nld: Dict[tuple, List[str]] = defaultdict(list)
        self.by_nl: Dict[tuple, List[str]] = defaultdict(list)
        self.by_ld: Dict[tuple, List[str]] = defaultdict(list)
        self.by_link: Dict[str, List[str]] = defaultdict(list)
        for r in self.ev_rows:
            name, date, link, key = r.get("event", ""), r.get("date", ""), r.get("event_link", ""), r.get("event_key", "")
            base = f"k\x1f{key}" if key else f"{link}\x1f{date}\x1f{name}"
            eid, n = sha12(base), 1
            while eid in self.ev_by_id:
                n += 1
                eid = sha12(f"{base}\x1f{n}")
            parts = [x.strip() for x in date.split(" to ")]
            start = parts[0] if parts else ""
            listed = [x for x in re.split(r";\s+", r.get("listed_on", "")) if x]
            streams = list(dict.fromkeys(stream_of(x) for x in listed)) or ["other"]
            ev = {"event_id": eid, "event": name, "date": date, "start": start,
                  "end": parts[1] if len(parts) > 1 else "", "year": year_of(start or date),
                  "event_link": link, "event_key": key, "link_host": host_of(link),
                  "listed_on": r.get("listed_on", ""), "stream_primary": streams[0], "_streams": streams,
                  "entry_text": "", "location_structured": r.get("location_structured", ""), "_row": r}
            ev["period"] = period_of(ev["year"])
            self.events.append(ev)
            self.ev_by_id[eid] = ev
            if key:
                self.by_key.setdefault(key, eid)
            self.by_nld[(name, link, date)].append(eid)
            self.by_nl[(name, link)].append(eid)
            self.by_ld[(link, date)].append(eid)
            self.by_link[link].append(eid)
        if len(self.events) != len({(e['event_link'], e['date'], e['event']) for e in self.events}):
            self.note("events with the same name, date and link", len(self.events) - len(
                {(e['event_link'], e['date'], e['event']) for e in self.events}), "kept apart by event_key")
        self.note("events (events.csv rows)", len(self.events))
        # event_codes: series, duplicates, country overrides
        self.alias: Dict[str, str] = {}
        self.ev_codes = {}
        # Events merged by the pipeline keep their old keys in merged_from: codes written
        # for an old key (its event_id) apply to the merged event.
        old_ids: Dict[str, str] = {}
        for ev in self.events:
            for old in (x.strip() for x in s(ev["_row"].get("merged_from")).split(";")):
                if old:
                    old_ids.setdefault(sha12(f"k\x1f{old}"), ev["event_id"])
                    self.by_key.setdefault(old, ev["event_id"])
        remapped = 0
        for c in self.event_codes:
            eid = c.get("event_id", "")
            if eid not in self.ev_by_id and eid in old_ids:
                c = dict(c, event_id=old_ids[eid])
                eid = c["event_id"]
                remapped += 1
            if eid not in self.ev_by_id:
                if eid:
                    self.warn("event_codes: event_id %s is not in events.csv (ignored)", eid)
                continue
            self.ev_codes[eid] = c
            dup = old_ids.get(c.get("duplicate_of", ""), c.get("duplicate_of", ""))
            if dup and dup in self.ev_by_id and dup != eid:
                self.alias[eid] = dup
        if remapped:
            self.note("event_codes rows mapped through merged_from", remapped,
                      "codes written for an event key that the pipeline has since merged")
        self._seed_entries()
        self._countries()
        for ev in self.events:
            c = self.ev_codes.get(ev["event_id"], {})
            link = ev["event_link"]
            if c.get("series_id"):
                ev["series_id"] = c["series_id"]
            elif is_web(link) and not p.is_platform(link):
                ev["series_id"] = norm_link(link)
            else:
                ev["series_id"] = ev["event_id"]
            text = ev["entry_text"]
            ev["entry_words"] = (int(ev["_row"]["entry_words"]) if s(ev["_row"].get("entry_words")).isdigit()
                                 else len(_WORD.findall(p.URL_LIKE.sub(" ", text))) if text else 0)
            ev["entry_lang"] = ""
            if text and p.detect_language:
                try:
                    ev["entry_lang"] = p.detect_language(p.URL_LIKE.sub(" ", text))
                except Exception:  # noqa: BLE001
                    ev["entry_lang"] = ""
        self.series_of = {e["event_id"]: e["series_id"] for e in self.events}
        # Series whose editions carry different names (a shared link used by several events)
        by_series = defaultdict(list)
        for ev in self.events:
            by_series[ev["series_id"]].append(ev)
        try:
            from events import series_tokens
        except Exception:  # noqa: BLE001 - older or edited module
            def series_tokens(name: str) -> Set[str]:
                name = re.sub(r"[(\[][^)\]]*[)\]]", " ", name or "")
                return set(re.findall(r"[^\W\d_]{3,}", name.casefold()))
        multi = []
        for sid, evs in sorted(by_series.items()):
            names = list(dict.fromkeys(collapse(e["event"]) for e in evs))
            if len(evs) < 2 or len(names) < 2:
                continue
            # names that share too few words (Jaccard below 0.4, as events.same_series) are other events
            toks = [series_tokens(n) for n in names]
            group = list(range(len(names)))
            for i in range(len(names)):
                for j in range(i + 1, len(names)):
                    a_, b_ = toks[i], toks[j]
                    if a_ and b_ and len(a_ & b_) / len(a_ | b_) >= 0.4:
                        gi, gj = group[i], group[j]
                        group = [gi if g == gj else g for g in group]
            # Two name groups are mostly one event listed in two languages; three or more
            # usually mean that different events share one link.
            n_groups = len(set(group))
            if n_groups >= 3:
                multi.append({"series_id": sid, "n_events": len(evs), "n_names": n_groups,
                              "names": " | ".join(short(n, 80) for n in names[:6])})
        self.multi_name_series = multi
        self.note("series whose editions carry three or more unrelated names", len(multi),
                  "listed in the multi_name_series table; split them with event_codes.csv series_id")

    def _seed_entries(self) -> None:
        """Agenda entries of the seed PDFs (seed_links.json): entry text and the
        full set of agendas an event is listed in (events.csv keeps one source in
        older runs)."""
        self.entries: List[dict] = []
        path = self.inputs.get("seed_links.json")
        if not path:
            self.skipped["entry_rates"] = "seed_links.json not found (give the run folder with --run)"
            self.skipped["agenda_overlap"] = "seed_links.json not found"
            for ev in self.events:
                ev["streams"] = ev["_streams"]
            return
        try:
            pdfs = {}
            for p in self.all_inputs.get("seed_links.json", [path]):   # several --run folders
                data = json.loads(p.read_text(encoding="utf-8"))
                for k, v in (data.get("pdfs") or {}).items():
                    pdfs.setdefault(k, v)
        except (OSError, ValueError) as exc:
            self.skipped["entry_rates"] = f"seed_links.json unreadable ({exc})"
            for ev in self.events:
                ev["streams"] = ev["_streams"]
            return
        keys = links = None
        try:
            from site_actors import assign_entry_keys
            keys, links = assign_entry_keys(pdfs, set())
            how = "site_actors.assign_entry_keys"
        except Exception as exc:  # noqa: BLE001 - module being edited or older version
            log.info("site_actors.assign_entry_keys not usable (%s); simple entry keys", exc)
            how = "simple link#date keys"
        for pdf, d in pdfs.items():
            st = stream_of(pdf_display(pdf))
            for e in d.get("entries") or []:
                idx = (pdf, e.get("index"))
                if keys is not None:
                    key, link = keys.get(idx, ""), links.get(idx, "")
                else:
                    good = [u for u in e.get("links") or []
                            if self.p.filter_targets is None or self.p.filter_targets({u: []}, set())]
                    link = good[0] if good else ((e.get("links") or [""])[0])
                    when = e.get("start") or e.get("month") or f"entry-{(e.get('index') or 0) + 1}"
                    key = f"{good[0]}#{when}" if good else f"{pdf}#entry-{(e.get('index') or 0) + 1}"
                self.entries.append({"pdf": pdf, "stream": st, "index": e.get("index"), "name": e.get("name", ""),
                                     "start": e.get("start") or "", "month": e.get("month") or "",
                                     "lines": e.get("lines") or [], "key": pdf_display(key), "raw_key": key,
                                     "link": pdf_display(link) if link else pdf_display(pdf)})
        by_key = defaultdict(set)
        for en in self.entries:
            by_key[en["key"]].add(en["pdf"])
        both = sum(1 for v in by_key.values() if len(v) > 1)
        self.note("seed entries", len(self.entries), "; ".join(f"{k}: {v}" for k, v in sorted(
            Counter(en["stream"] for en in self.entries).items())))
        self.note("seed entry keys (rebuilt)", len(by_key), how)
        self.note("seed entry keys listed in two or more agenda PDFs", both)
        # entries -> events
        ev_key = {e["event_key"]: e["event_id"] for e in self.events if e["event_key"]}
        by_ls = defaultdict(list)
        by_ns = defaultdict(list)
        for e in self.events:
            by_ls[(norm_link(pdf_display(e["event_link"])), e["start"])].append(e["event_id"])
            by_ns[(collapse(e["event"]).casefold(), e["start"])].append(e["event_id"])
        matched = 0
        self.entries_of: Dict[str, List[dict]] = defaultdict(list)
        for en in self.entries:
            eid = ev_key.get(en["key"]) or ev_key.get(en["raw_key"])
            if not eid:
                when = en["start"] or en["month"]
                c = by_ls.get((norm_link(en["link"]), when)) or []
                if len(c) != 1:
                    c = by_ns.get((collapse(en["name"]).casefold(), when)) or c
                eid = c[0] if len(c) == 1 else None
            en["event_id"] = eid
            if eid:
                matched += 1
                self.entries_of[eid].append(en)
        self.note("seed entries matched to events.csv", f"{matched}/{len(self.entries)}")
        for ev in self.events:
            ens = self.entries_of.get(ev["event_id"], [])
            if ens:
                ev["entry_text"] = "\n".join(ens[0]["lines"])
            ev["streams"] = list(dict.fromkeys(ev["_streams"] + [en["stream"] for en in ens]))
            if ev["stream_primary"] == "other" and ens:
                ev["stream_primary"] = ens[0]["stream"]
        multi = sum(1 for ev in self.events if len([x for x in ev["streams"] if x != "hub"]) > 1)
        self.facts["events_in_two_agendas"] = multi

    def _countries(self) -> None:
        p = self.p
        rows_have = "country_iso" in self.ev_fields and any(e["_row"].get("country_iso") for e in self.events)
        self.country_from = "events.csv" if rows_have else ("gazetteer" if p.gazetteer else "none")
        if not rows_have and p.gazetteer is None:
            self.skipped["country"] = f"no country columns in events.csv and gazetteer.py not importable ({p.gazetteer_error})"
        for ev in self.events:
            r = ev["_row"]
            if rows_have:
                isos = [x for x in re.split(r";\s*", r.get("country_iso", "")) if x]
                method, text = r.get("country_source", ""), r.get("location_text", "")
            else:
                res = p.country(ev["event"], ev["entry_text"], ev["event_link"] if is_web(ev["event_link"]) else "",
                                ev["location_structured"])
                isos = [str(x) for x in res.get("countries") or []]
                method = s(res.get("method"))
                text = res.get("matched", "")
                text = "; ".join(text) if isinstance(text, (list, tuple)) else s(text)
            c = self.ev_codes.get(ev["event_id"], {})
            if c.get("country_override"):
                isos = [x.strip().upper() for x in re.split(r"[;,]", c["country_override"]) if x.strip()]
                method, text = "override", "event_codes.csv"
            ev["countries"] = isos
            ev["country_method"] = method if isos else "none"
            ev["location_text"] = text if isos else ""
            ev["country_names"] = [p.country_name(x) for x in isos]
        n = sum(1 for e in self.events if e["countries"])
        self.note("events with a country", f"{n}/{len(self.events)}", f"from {self.country_from}")

    # ------------------------------------------------------------ refs
    def _explode(self, r: dict, long_format: bool) -> List[Tuple[str, str, str, str]]:
        if long_format:
            return [(r.get("event", ""), r.get("event_date", ""), r.get("event_link", ""), r.get("event_key", ""))]
        names = r.get("event", "").split(" | ")
        dates = r.get("event_date", "").split(" | ")
        links = r.get("event_link", "").split(" | ")
        if not any(names) and not any(links):
            return [(NO_EVENT, "", "", "")]
        if len(names) > len(links) == 1 and len(dates) <= 1:
            names = [" | ".join(names)]      # one event whose name contains " | "
        n = max(len(names), len(dates), len(links))
        if len(names) != n and len(links) == n:  # re-align names on the event table
            out = []
            for i, ln in enumerate(links):
                d = dates[i] if i < len(dates) else ""
                c = self.by_ld.get((ln, d)) or self.by_link.get(ln) or []
                nm = self.ev_by_id[c[0]]["event"] if len(set(c)) == 1 else (names[i] if i < len(names) else "")
                out.append((nm, d, ln, ""))
            return out
        pad = lambda xs: xs + [""] * (n - len(xs))  # noqa: E731
        return [(a, b, c, "") for a, b, c in zip(pad(names), pad(dates), pad(links))]

    def _match(self, name: str, date: str, link: str, key: str) -> Tuple[Optional[str], str]:
        if (not name and not link) or name.startswith("("):
            return None, "pseudo"
        if key:
            eid = self.by_key.get(key) or self.by_key.get(pdf_display(key))
            if eid:
                return self.alias.get(eid, eid), "key"
        c = self.by_nld.get((name, link, date), [])
        if len(c) == 1:
            return self.alias.get(c[0], c[0]), "exact"
        c = self.by_nl.get((name, link), [])
        if len(c) == 1:
            return self.alias.get(c[0], c[0]), "name+link"
        c = self.by_ld.get((link, date), [])
        if len(c) == 1:
            return self.alias.get(c[0], c[0]), "link+date"
        return None, "unmatched"

    @staticmethod
    def _layer(row: dict, ev_link: str, pseudo: bool) -> str:
        if row.get("layer"):
            return row["layer"]
        url = row.get("source_link", "")
        if row.get("site_role") == "seed" or url.startswith(("file:", "pdf:")):
            return "listing_seed"
        if "calendar_month=" in url:
            return "listing_calendar"
        if row.get("site_role") == "main":
            return "main_page" if pseudo else "main_event_page"
        if ev_link and norm_link(ev_link) == norm_link(url):
            return "landing"
        pdf = truthy(row.get("is_pdf")) or bool(re.search(r"\.pdf(?:\?|$)", url, re.I))
        return "linked_pdf" if pdf else "linked_page"

    @staticmethod
    def _unit(row: dict, layer: str, eid: Optional[str], name: str) -> str:
        if row.get("source_unit"):
            return row["source_unit"]
        if layer == "listing_seed":
            return f"{row.get('source_link', '')}#{eid or name}"
        return row.get("source_link", "")

    def _doc_year(self, url: str, page_date: str) -> Optional[int]:
        y = doc_year_of(url)
        if y is None and page_date:
            y = year_of(page_date)
        if y is None:
            y = self.page_year.get(url)
        return y

    @staticmethod
    def _doc_ids(row: dict, unit: str, layer: str) -> Tuple[str, str]:
        """(source unit, document) of a row. Web links are compared without scheme, www.,
        fragment and trailing slash; a copy (duplicate_document_of) is the same source and
        document as its original, a translation (translation_of) the same document only."""
        url = row.get("source_link", "")
        if layer == "listing_seed":
            return unit, unit
        dup = s(row.get("duplicate_document_of"))
        tr = s(row.get("translation_of"))
        if layer == "listing_calendar":
            return unit, norm_doc(url)
        unit = norm_doc(dup or unit)
        return unit, norm_doc(tr or dup or url)

    def _sublevel(self, layer: str, url: str, eid: Optional[str]) -> str:
        """L1, L2, L3s (a linked page or PDF on the event's own host) or L3x."""
        lvl = TIE_LEVEL.get(layer, "L3")
        if lvl != "L3":
            return lvl
        ev = self.ev_by_id.get(eid) if eid else None
        return "L3s" if ev and same_host(host_of(url), ev["link_host"]) else "L3x"

    def _doc_kinds(self) -> None:
        """doc_kind of every tie: from the masters, else make_master.classify_document,
        else the F2 rule (an L3 source naming more than 100 organisations whose address
        and title have no programme or partner word is a report). --source-codes include
        (1/0) overrides it. Then participation and context_reason for older masters."""
        words: Dict[str, int] = {}
        for r in self.sources_sheet or self.sources_rows:
            w = s(r.get("words"))
            if w.isdigit():
                for u in (r.get("source_unit", ""), r.get("source_link", "")):
                    if u:
                        words.setdefault(u, int(w))
                        words.setdefault(norm_doc(u), int(w))
        self._words = words
        min_rank = CONF_RANK[self.a.min_confidence]
        orgs = defaultdict(set)
        for t in self.ties:
            if not t["in_menus"] and CONF_RANK.get(t["confidence"], 0) >= min_rank:
                orgs[t["unit"]].add(t["raw_key"])
        self._unit_orgs = {u: len(v) for u, v in orgs.items()}
        missing = [t for t in self.ties if not t["doc_kind"]]
        self.doc_kind_from = "masters (doc_kind column)"
        self._classify = None
        if missing:
            try:
                from make_master import classify_document
                self._classify = classify_document
                self.doc_kind_from = "make_master.classify_document (masters without doc_kind)"
            except Exception as exc:  # noqa: BLE001 - older or edited module: the F2 rule
                self.doc_kind_from = f"F2 rule (make_master.classify_document not usable: {type(exc).__name__})"
        self._unit_kind: Dict[str, Tuple[str, str]] = {}
        for t in missing:
            t["doc_kind"], t["doc_kind_reason"] = self.doc_kind_of(t["unit"], t["layer"], t["source_link"],
                                                                   t["source_title"], t["is_pdf"])
        # --source-codes: include 1 keeps a source in the network whatever its doc_kind, 0 leaves it out
        self.source_override: Dict[str, bool] = {}
        for c in self.source_codes:
            u = s(c.get("source_unit")) or s(c.get("source_link"))
            inc = s(c.get("include"))
            if u and inc in ("0", "1"):
                self.source_override[u] = inc == "1"
                self.source_override[norm_doc(u)] = inc == "1"
        n_over = 0
        for t in self.ties:
            over = self.source_override.get(t["unit"])
            if self.has_participation:
                part = t["participation_master"]
                reasons = [x for x in t["context_reason"].split("; ") if x]
            else:
                reasons = []
                if t["doc_kind"] == "report":
                    reasons.append("linked report or study")
                if t["attribution"] == "other_year":
                    reasons.append(f"about another year ({t['source_year'] or '?'})")
                if t["attribution"] == "site_level":
                    reasons.append("other pages of a site several events link to")
                if t["attribution"] == "not_tied":
                    reasons.append("not tied to an event")
                part = not reasons
            if over is not None:
                n_over += 1
                reasons = [x for x in reasons if not x.startswith("linked report")]
                if over:
                    part = part or not reasons
                else:
                    part = False
                    reasons.append("left out by --source-codes")
            t["participation"] = part
            t["context_reason"] = "; ".join(reasons)
            t["ctx_kinds"] = {_reason_kind(x) for x in reasons}
            t["reference"] = (t["doc_kind"] == "report" and over is not True) or over is False
        refs = {x["unit"]: x for x in self.ties if x["reference"]}
        n_ev = len({x["unit"] for x in self.ties if x["reference"] and x["event_id"]})
        self.note("doc_kind from", self.doc_kind_from, f"{len(missing)} rows classified here" if missing else "")
        self.note("reference documents (doc_kind report) tied to an event", n_ev,
                  f"{len(refs)} sources in all; " + "; ".join(f"{k} {v}" for k, v in sorted(Counter(
                      ("pdf" if x["is_pdf"] else "page") for x in refs.values()).items())))
        if self.source_codes:
            self.note("--source-codes rows applied", n_over, f"{len(self.source_override) // 2 or 0} sources coded")
        self.note("participation from", "masters (participation column)" if self.has_participation else
                  "derived here: doc_kind report and placeholder events are context")

    def doc_kind_of(self, unit: str, layer: str, url: str, title: str, is_pdf: bool) -> Tuple[str, str]:
        if unit in self._unit_kind:
            return self._unit_kind[unit]
        n_orgs = self._unit_orgs.get(unit, 0)
        w = self._words.get(unit) or self._words.get(url) or self._words.get(norm_doc(url)) or 0
        if self._classify is not None:
            try:
                res = tuple(self._classify(url, title, w, n_orgs, is_pdf, layer))
            except Exception as exc:  # noqa: BLE001
                res = ("page", f"classify_document failed ({type(exc).__name__})")
        elif layer in ("listing_seed", "listing_calendar"):
            res = ("listing", "agenda entry or calendar cell")
        elif TIE_LEVEL.get(layer, "L3") == "L3" and n_orgs > 100 and not _PROGRAMME_HINT.search(
                f"{unquote(url)} {title}"):
            res = ("report", f"F2 rule: L3 source naming {n_orgs} organisations, no programme or partner word")
        else:
            res = ("programme" if is_pdf and _PROGRAMME_HINT.search(f"{unquote(url)} {title}") else "page", "F2 rule")
        self._unit_kind[unit] = res
        return res

    def build_refs(self) -> None:
        self.page_year: Dict[str, int] = {}
        self.pages = self.optional_rows("pages.csv")
        for pg in self.pages:
            y = year_of(pg.get("date") or pg.get("page_date") or "")
            if y:
                self.page_year[pg.get("url", "")] = y
        # ---- network ties
        self.ties: List[dict] = []
        counts = {"event_organisations": Counter(), "context_organisations": Counter(), "menus_and_footers": Counter()}
        self.unmatched_names: Counter = Counter()
        has_part = "participation" in self.net_fields or bool(self.ctx_rows)
        for sheet, rows in (("event_organisations", self.net_rows), ("context_organisations", self.ctx_rows),
                            ("menus_and_footers", self.menu_rows)):
            for r in rows:
                refs = self._explode(r, self.long_net)
                for name, date, link, key in refs:
                    eid, how = self._match(name, date, link, key)
                    counts[sheet][how] += 1
                    if how == "unmatched":
                        self.unmatched_names[name] += 1
                        continue
                    pseudo = eid is None
                    layer = self._layer(r, link, pseudo)
                    url = r.get("source_link", "")
                    unit, document = self._doc_ids(r, self._unit(r, layer, eid, name), layer)
                    sht = r.get("source_has_theme", "")
                    att = s(r.get("attribution"))
                    if not att:      # older masters: derived from the placeholder names
                        att = ("not_tied" if name == NO_EVENT or not name else "site_level") if pseudo else "event"
                    part = s(r.get("participation")).casefold()
                    if not part:
                        part = "no" if sheet == "context_organisations" else "yes"
                    self.ties.append({
                        "event_id": eid or "", "match": how, "pseudo_label": name if pseudo else "",
                        "organisation": collapse(r.get("organisation", "")),
                        "raw_key": r.get("org_key") or self.p.key(r.get("organisation", "")),
                        "confidence": r.get("confidence", "") or "low", "source_link": url, "unit": unit,
                        "document": document, "source_title": r.get("source_title") or r.get("page_title", ""),
                        "is_pdf": truthy(r.get("is_pdf")) or layer == "linked_pdf",
                        "site": r.get("site", ""), "site_role": r.get("site_role", ""), "layer": layer,
                        "tie_level": TIE_LEVEL.get(layer, "L3"),
                        "tie_sublevel": self._sublevel(layer, url, eid),
                        "doc_kind": s(r.get("doc_kind")), "doc_kind_reason": s(r.get("doc_kind_reason")),
                        "attribution": att, "source_year": s(r.get("source_year")),
                        "duplicate_document_of": s(r.get("duplicate_document_of")),
                        "translation_of": s(r.get("translation_of")),
                        "participation_master": part == "yes", "context_reason": s(r.get("context_reason")),
                        "named_in_event_title": truthy(r.get("named_in_event_title")),
                        "in_menus": sheet == "menus_and_footers" or truthy(r.get("only_in_menus_or_footers")),
                        "found_by": r.get("found_by", ""), "mentions": int(r.get("mentions_on_source") or 1)
                        if s(r.get("mentions_on_source")).isdigit() else 1,
                        "source_has_theme": truthy(sht) if sht else None,
                        "example": r.get("example_context", ""),
                        "doc_year": self._doc_year(url, r.get("page_date", "")), "sheet": sheet})
        for sheet, c in counts.items():
            if sheet != "context_organisations" or self.ctx_rows:
                self.note(f"event refs, {sheet}", "; ".join(f"{k} {v}" for k, v in sorted(c.items())) or "0")
        self.has_participation = has_part
        self._doc_kinds()
        self.ref_counts = counts
        if self.unmatched_names:
            self.note("unmatched event refs (dropped)", sum(self.unmatched_names.values()),
                      "; ".join(f"{k!r} x{v}" for k, v in self.unmatched_names.most_common(10)))
        # ---- theme instances
        self.instances: List[dict] = []
        for sheet, rows in (("theme_instances", self.th_rows), ("menus_and_sidebars", self.th_menu_rows)):
            # Long layout: one row per (hit, event), so rows of one hit are joined again.
            # A term found twice in the same sentence gives two hits with identical
            # columns; a row whose event is already in the group starts a new instance.
            groups: "OrderedDict[tuple, List[dict]]" = OrderedDict()
            for i, r in enumerate(rows):
                gk = ((sheet, r.get("source_unit") or r.get("source_link", ""), r.get("theme", ""),
                       r.get("term", ""), r.get("sentence", ""), r.get("context_before", ""),
                       r.get("context_after", ""), r.get("region", ""), r.get("duplicate_of", ""),
                       r.get("boilerplate", "")) if self.long_th else (sheet, i))
                refs = [(name, date, link, key, r.get("layer", ""))
                        for name, date, link, key in self._explode(r, self.long_th)]
                ev_ids = {rf[:4] for rf in refs}
                cands = groups.setdefault(gk, [])
                g = next((c for c in cands if not ev_ids & c["seen"]), None)
                if g is None:
                    g = {"row": r, "sheet": sheet, "refs": [], "seen": set()}
                    cands.append(g)
                g["refs"].extend(refs)
                g["seen"] |= ev_ids
            for cands in groups.values():
                for g in cands:
                    self.instances.append(self._instance(g))
        # Descriptive themes (global_south, authoritarianism) are kept apart: they do not
        # drive the theme filter or any prevalence measure (table descriptive_themes).
        desc = set(self.descriptive)
        self.desc_instances = [i for i in self.instances if i["theme"] in desc]
        self.instances = [i for i in self.instances if i["theme"] not in desc]
        if self.desc_instances:
            self.note("descriptive theme rows (kept apart)", len(self.desc_instances),
                      "; ".join(f"{k} {v}" for k, v in sorted(Counter(i["theme"] for i in self.desc_instances).items())))
        self.theme_refs = [rf for inst in self.instances for rf in inst["refs"]]
        tc = Counter(rf["match"] for rf in self.theme_refs)
        self.note("event refs, theme rows", "; ".join(f"{k} {v}" for k, v in sorted(tc.items())))
        self.theme_ref_counts = tc

    def _instance(self, g: dict) -> dict:
        r = g["row"]
        theme, term, sent = r.get("theme", ""), r.get("term", ""), r.get("sentence", "")
        # Same rule as make_master.py: menus, headers, footers and sidebars are furniture.
        # A passage repeated in the main content of many pages (boilerplate=True with
        # menu_or_sidebar=False) counts as a theme hit of every page carrying it. Older theme
        # rows without the menu_or_sidebar column fall back to the boilerplate flag.
        repeated = truthy(r.get("boilerplate"))
        furniture = (r.get("region", "") in FURNITURE_REGIONS or truthy(r.get("menu_or_sidebar"))
                     or g["sheet"] == "menus_and_sidebars"
                     or (repeated and "menu_or_sidebar" not in r))
        dup = bool(r.get("duplicate_of"))
        family, section = self._family(theme, term)
        low = sent.casefold()
        tl = term.casefold()
        inst = {"theme": theme, "term": term, "sentence": sent, "context_before": r.get("context_before", ""),
                "context_after": r.get("context_after", ""), "source_link": r.get("source_link", ""),
                "site": r.get("site", ""), "site_role": r.get("site_role", ""), "is_pdf": truthy(r.get("is_pdf")),
                "region": r.get("region", ""), "furniture": furniture, "repeated": repeated, "dup": dup, "clean": not furniture and not dup,
                # the same sentence is the same passage (make_master.passage_id), wherever it is
                "passage_id": s(r.get("passage_id")) or sha12(" ".join(sent.casefold().split())),
                "page_title": r.get("page_title", ""), "doc_kind": s(r.get("doc_kind")),
                "term_family": family, "language_section": section,
                "subtheme": subtheme(theme, term), "stance": stance(term, sent[:max(0, low.find(tl))] if tl in low else ""),
                "in_handle": bool(tl and re.search(r"[@#][\w.-]*" + re.escape(tl), low)),
                "digit_glued": bool(tl and re.search(r"\d" + re.escape(tl), low)),
                "orgs_text": r.get("organisations_in_paragraph", ""),
                "org_keys_text": r.get("organisations_in_paragraph_keys", ""),
                "page_date": r.get("page_date", ""), "refs": []}
        refs = g["refs"]
        inst["doc_year"] = self._doc_year(inst["source_link"], inst["page_date"])
        w = 1 / max(1, len(refs))
        names = []
        for name, date, link, key, layer in refs:
            eid, how = self._match(name, date, link, key)
            pseudo = eid is None
            lay = self._layer({**r, "layer": layer}, link, pseudo)
            unit, document = self._doc_ids(r, self._unit(r, lay, eid, name), lay)
            ev = self.ev_by_id.get(eid) if eid else None
            names.append(ev["event"] if ev else name)
            kind = inst["doc_kind"] or self.doc_kind_of(unit, lay, inst["source_link"], inst["page_title"],
                                                        inst["is_pdf"] or lay == "linked_pdf")[0]
            over = self.source_override.get(unit)
            inst["refs"].append({"inst": inst, "event_id": eid or "", "match": how,
                                 "pseudo_label": name if pseudo else "", "layer": lay,
                                 "tie_level": TIE_LEVEL.get(lay, "L3"),
                                 "tie_sublevel": self._sublevel(lay, inst["source_link"], eid), "unit": unit,
                                 "document": document, "w": w, "doc_kind": kind,
                                 "reference": (kind == "report" and over is not True) or over is False})
        inst["in_event_title"] = bool(tl) and any(tl in n.casefold() for n in names)
        inst["event_names"] = names
        return inst

    def _family(self, theme: str, term: str) -> Tuple[str, str]:
        if not self.ts:
            return "", ""
        t = term.strip()
        for raw, pat in zip(self.ts.raw.get(theme, []), self.ts.themes.get(theme, [])):
            try:
                if pat.fullmatch(t) or pat.fullmatch(t.translate(getattr(self.p.extract, "_GREEK_FOLD", {}))):
                    return raw, self.sections.get((theme, raw), "")
            except Exception:  # noqa: BLE001
                continue
        return "(unmapped)", ""


# ==========================================================================
# Term model: sub-themes and stance (rule-based, first match wins)
# ==========================================================================
SUBTHEME_RULES = {
    "colonialism": [("decolonial", r"decolon|dekolon|décolon|descolon"), ("anticolonial", r"anti.?colon|antikolon"),
                    ("neocolonial", r"neo.?colon|néo.?colon|neokolon"), ("extractivism", r"extracti|extrativ|estratti|extraktiv"),
                    ("global_south", r"global south|süden|sud global|sur global|sud del mondo"),
                    ("reparations", r"reparation|réparation|climate debt|dette|klimaschuld"),
                    ("imperialism", r"imperial|impérial|empire")],
    "racism": [("antiracism", r"anti.?ras|anti.?rac|antirazz"), ("whiteness", r"white (?:supremacy|privilege)"),
               ("xenophobia", r"xeno"), ("antisemitism", r"antisemit|anti-semit"), ("islamophobia", r"islamoph"),
               ("environmental", r"environmental rac"), ("structural", r"racial|racis(?:ed|ation)|racisé|structural")],
    "fascism": [("antifascism", r"antifa|anti.?fa[sš]c|anti.?fasch"),
                ("compound", r"(?:techno|eco|öko|oko|fossil|klima).?fa[sš]"),
                ("nazism", r"nazi"), ("far_right", r"far.?right|extrême.?droite|extreme.?droite|rechts"),
                ("authoritarianism", r"autorit|authorit")],
}
_AGAINST = re.compile(r"(?:^|\W)(?:against|contre|gegen|contro|tegen|stop)\W", re.I)


def subtheme(theme: str, term: str) -> str:
    t = term.casefold()
    for name, pat in SUBTHEME_RULES.get(theme, []):
        if re.search(pat, t):
            return name
    return "core"


def stance(term: str, left: str) -> str:
    t = term.casefold()
    if "anti" in t:
        return "anti"
    if re.match(r"(?:de|dé)-?", t) and re.match(r"(?:de|dé)-?(?:colon|kolon)", t):
        return "de"
    if _AGAINST.search(" " + left[-40:] + " "):
        return "against"
    return "unmarked"


# ==========================================================================
# Theme filter (R1), organisation identity, event coverage
# ==========================================================================
def _derive(an: Analysis) -> None:
    """Theme-filter sets, organisation keys, per-event coverage and theme flags."""
    p = an.p
    # ---- R1: sources (pages, PDFs, agenda entries) with a theme hit in their content
    units, sites, evs = set(), set(), set()
    for inst in an.instances:
        if inst["furniture"]:
            continue
        sites.add(inst["site"])
        for rf in inst["refs"]:
            units.add(rf["unit"])
            if rf["event_id"]:
                evs.add(rf["event_id"])
    an.theme_units_set, an.theme_sites, an.theme_events = units, sites, evs
    tf = an.a.theme_filter
    trusted = 0
    for t in an.ties:
        src = t["source_has_theme"] if t["source_has_theme"] is not None else t["unit"] in units
        # Agenda entries of a seed PDF are judged one by one under 'site' too (as in make_master.py).
        seed = t["site_role"] == "seed" or "#entry-" in t["unit"]
        ev_hit = bool(t["event_id"]) and t["event_id"] in evs
        t["pass"] = {"none": True, "source": src, "site": src if seed else t["site"] in sites,
                     "event": ev_hit if t["event_id"] else src,
                     "direct": src or (ev_hit and t["layer"] in DIRECT_LAYERS)}
        if (t["source_has_theme"] is None and tf != "none" and an.master_filter == tf
                and t["sheet"] != "menus_and_footers" and not t["pass"][tf]):
            # The masters were already filtered with this filter but do not say per row how
            # each source was judged (e.g. rebuilt from an old folder): trust that selection.
            t["pass"][tf] = True
            trusted += 1
    if trusted:
        log.info("%d network rows have no source_has_theme value; they were kept because make_master.py "
                 "already applied the theme filter '%s' to these masters", trusted, tf)
        an.note("rows kept on make_master.py's theme filter", trusted,
                "source_has_theme blank; masters already filtered with the same filter")
    real = [t for t in an.ties if t["sheet"] != "menus_and_footers"]
    kept = [t for t in real if t["pass"][tf]]
    an.filter_counts = {
        "rows": (len(kept), len(real)),
        "organisations": (len({t["raw_key"] for t in kept}), len({t["raw_key"] for t in real})),
        "events": (len({t["event_id"] for t in kept if t["event_id"]}), len({t["event_id"] for t in real if t["event_id"]})),
        "sources": (len({t["unit"] for t in kept}), len({t["unit"] for t in real}))}
    if tf != "none":
        fc = an.filter_counts
        log.info("Theme filter '%s': kept %d of %d network rows, %d of %d organisations, %d of %d events, "
                 "%d of %d sources (use --theme-filter none to keep everything)", tf,
                 *fc["rows"], *fc["organisations"], *fc["events"], *fc["sources"])
    for k, (a_, b_) in an.filter_counts.items():
        an.note(f"theme filter '{tf}': {k} kept", f"{a_}/{b_}", f"{b_ - a_} left out")
    # ---- organisations
    _org_identity(an)
    # ---- event coverage
    pages_by_event = defaultdict(lambda: [0, 0])
    for pg in an.pages:
        for nm in [x for x in re.split(r";\s+", pg.get("via_event", "")) if x]:
            pe = pages_by_event[nm]
            pe[0] += 1
            pe[1] += int(pg.get("words") or 0) if s(pg.get("words")).isdigit() else 0
    read_hosts = set()
    sites_rows = an.optional_rows("sites.csv")
    for sr in sites_rows:
        n = int(sr.get("pages_analysed") or 0) if s(sr.get("pages_analysed")).isdigit() else 0
        if sr.get("role") == "linked" and n >= 1:
            read_hosts.add(sr.get("site", "").lower())
    row_web = {t["event_id"] for t in an.ties if t["event_id"] and t["layer"] != "listing_seed"}
    row_web |= {rf["event_id"] for rf in an.theme_refs if rf["event_id"] and rf["layer"] != "listing_seed"}
    csv_followed = "link_followed" in an.ev_fields
    if csv_followed:
        an.web_method = "link_followed column of events.csv, plus events with web-layer rows"
    elif an.pages or sites_rows:
        an.web_method = "sites.csv and pages.csv, plus events with web-layer rows"
    else:
        an.web_method = "proxy: events with web-layer rows in the masters"
    if sites_rows:
        an.note("sites.csv read from", "; ".join(str(x) for x in an.all_inputs.get("sites.csv", [])),
                "used for link_followed" if not csv_followed else "")
    for ev in an.events:
        r = ev["_row"]
        host = ev["link_host"]
        via = pages_by_event.get(ev["event"], [0, 0])
        on_site = bool(host) and p.host_matches(host, read_hosts)
        if csv_followed:
            followed = truthy(r.get("link_followed"))
        elif sites_rows:
            followed = on_site
        else:
            followed = False
        # An event with web-layer rows had its link read, whatever sites.csv says (sites.csv
        # from another run, or a link host with site role 'main'): link_followed >= has_web_text.
        followed = followed or (ev["event_id"] in row_web and is_web(ev["event_link"]))
        ev["has_link"] = is_web(ev["event_link"]) and not ev["event_link"].startswith(("file:", "pdf:"))
        ev["link_followed"] = followed
        ev["has_web_text"] = followed or on_site or via[0] > 0 or ev["event_id"] in row_web
        ev["n_pages"] = int(r["pages_analysed"]) if s(r.get("pages_analysed")).isdigit() else via[0]
        ev["n_words_web"] = int(r["words_analysed"]) if s(r.get("words_analysed")).isdigit() else via[1]
    an.note("events with web text (has_web_text)", f"{sum(e['has_web_text'] for e in an.events)}/{len(an.events)}",
            an.web_method)
    # ---- event theme flags: every row outside menus and footers. Rows with duplicate_of
    # count here (the pipeline marks the second sighting of a passage, which depends on
    # crawl order and may belong to another event); the dedup applies only to the
    # passage, document and concentration counts.
    flags = defaultdict(lambda: defaultdict(set))   # eid -> theme -> {L1, L2, L3}
    for inst in an.instances:
        if inst["furniture"]:
            continue
        for rf in inst["refs"]:
            if rf["event_id"]:
                lvl = "own" if rf["layer"] in OWN_LAYERS else LAYER_GROUP.get(rf["tie_level"], "linked")
                if rf["layer"] == "main_page":
                    lvl = "linked"
                flags[rf["event_id"]][inst["theme"]].add(lvl)
                # core: the listing, the landing page and the event's own host, without reports
                if rf["tie_sublevel"] in ("L1", "L2", "L3s") and not rf["reference"]:
                    flags[rf["event_id"]][inst["theme"]].add("core")
    for ev in an.events:
        f = flags.get(ev["event_id"], {})
        title_hits = an.ts.find(ev["event"]) if an.ts else {}
        for th in an.themes:
            lv = f.get(th, set())
            ev[f"{th}_own"] = "own" in lv
            ev[f"{th}_landing"] = "landing" in lv
            ev[f"{th}_linked"] = "linked" in lv
            ev[f"{th}_any"] = bool(lv)
            ev[f"{th}_core"] = "core" in lv
            ev[f"{th}_in_title"] = th in title_hits
        for m in ("own", "landing", "linked", "any", "core", "in_title"):
            ev[f"any_{m}"] = any(ev[f"{th}_{m}"] for th in an.themes)
        ev["words_read_bin"] = words_bin(ev["n_words_web"])


def words_bin(words: int) -> str:
    for hi, lab in WORDS_BINS:
        if hi is None or (words or 0) < hi:
            return lab
    return WORDS_BINS[-1][1]


def _fold_umlaut(k: str) -> str:
    return k.replace("ae", "a").replace("oe", "o").replace("ue", "u")


def _org_identity(an: Analysis) -> None:
    p = an.p
    surfaces: Dict[str, Counter] = defaultdict(Counter)
    for t in an.ties:
        surfaces[t["raw_key"]][t["organisation"]] += 1
    an.para_orgs: Dict[int, List[Tuple[str, str]]] = {}
    for i, inst in enumerate(an.instances):
        names = [x.strip() for x in inst["orgs_text"].split("; ") if x.strip()]
        keys = [x.strip() for x in inst["org_keys_text"].split("; ")] if inst["org_keys_text"] else []
        pairs = [(nm, keys[j] if j < len(keys) and keys[j] else p.key(nm)) for j, nm in enumerate(names)]
        an.para_orgs[i] = pairs
        for nm, k in pairs:
            surfaces[k][nm] += 1
    # Umlaut transliteration variants ("Ende Gelaende" / "Ende Gelände") merge automatically,
    # but only when the folded key exists and is written with an umlaut somewhere.
    umlaut = {}
    for k in surfaces:
        if re.search(r"ae|oe|ue", k):
            f = _fold_umlaut(k)
            if f != k and f in surfaces and any(re.search(r"[äöüÄÖÜ]", sf) for sf in surfaces[f]):
                umlaut[k] = f
    an.umlaut_merges = umlaut
    codes: Dict[str, dict] = {}
    stale = same = 0
    for c in an.org_codes:
        name = c.get("organisation", "")
        derived = p.key(name) if name else p.key(c.get("org_key", ""))
        if not derived:
            continue
        if c.get("org_key") and c["org_key"] != derived and p.key(c["org_key"]) != derived:
            stale += 1
        if derived in codes:
            # Two rows now give the same key (e.g. a spelling variant the current
            # canonical_key already folds): fill the gaps of the first row, never overwrite it.
            same += 1
            first = codes[derived]
            for f, v in c.items():
                if s(v) and not s(first.get(f)):
                    first[f] = v
        else:
            codes[derived] = dict(c)
    if stale:
        an.warn("org_codes: %d rows whose org_key no longer matches the organisation name (keys re-derived "
                "from the name; check them)", stale)
    if same:
        an.note("org_codes rows folded into an earlier row with the same key", same,
                "the first row's values win; empty cells are filled from the later row")
    merge = {}
    for k, c in codes.items():
        tgt = c.get("merge_into_key", "")
        if tgt:
            tk = p.key(tgt) or tgt
            if tk != k:
                merge[k] = tk

    def resolve(k: str) -> str:
        k = umlaut.get(k, k)
        seen = set()
        while k in merge and k not in seen:
            seen.add(k)
            k = umlaut.get(merge[k], merge[k])
        return k
    an.resolve = resolve
    an.codes = codes
    final_surfaces: Dict[str, Counter] = defaultdict(Counter)
    variants: Dict[str, Set[str]] = defaultdict(set)
    for k, cnt in surfaces.items():
        fk = resolve(k)
        final_surfaces[fk].update(cnt)
        variants[fk].add(k)
    an.org_name = {}
    for fk, cnt in final_surfaces.items():
        c = codes.get(fk, {})
        an.org_name[fk] = c.get("canonical_name") or sorted(cnt.items(), key=lambda x: (-x[1], x[0]))[0][0]
    an.org_variants = variants
    an.org_surfaces = final_surfaces
    for t in an.ties:
        t["org_key"] = resolve(t["raw_key"])
        c = codes.get(t["org_key"]) or codes.get(t["raw_key"]) or {}
        t["include"] = s(c.get("include"))
        t["actor_class"] = c.get("actor_class", "")
        ev = an.ev_by_id.get(t["event_id"])
        t["series_id"] = an.series_of.get(t["event_id"], "") if ev else ""
        t["year"] = ev["year"] if ev else None
        t["countries"] = ev["countries"] if ev else []
    for i, inst in enumerate(an.instances):
        inst["org_keys"] = list(dict.fromkeys(resolve(k) for _, k in an.para_orgs[i] if k))
    if umlaut:
        an.note("umlaut variants merged automatically", len(umlaut),
                "; ".join(f"{a} -> {b}" for a, b in sorted(umlaut.items())))
    if codes:
        an.note("org_codes rows read", len(codes), f"{len(merge)} merges")


def org_dropped(an: Analysis, k: str) -> bool:
    c = an.codes.get(k, {})
    return s(c.get("include")) in ("0", "no")


def org_class(an: Analysis, k: str) -> str:
    return an.codes.get(k, {}).get("actor_class", "")


# ==========================================================================
# Core analysis-ready tables
# ==========================================================================
def core_tables(an: Analysis) -> None:
    th = an.themes
    n_by_level = defaultdict(lambda: defaultdict(set))
    for t in an.ties:
        if t["event_id"] and not t["in_menus"] and t["pass"][an.a.theme_filter] and in_scope(t, an.a.network_scope):
            n_by_level[t["event_id"]][t["tie_level"]].add(t["org_key"])
            if t["tie_level"] == "L3":
                n_by_level[t["event_id"]][t["tie_sublevel"]].add(t["org_key"])
    rows = []
    for ev in an.events:
        r = {k: ev[k] for k in ("event_id", "event", "date", "start", "end", "year", "period", "event_link",
                                "event_key", "link_host", "series_id", "stream_primary")}
        r.update({"streams": "; ".join(ev["streams"]), "n_streams": len(ev["streams"]),
                  "countries": "; ".join(ev["countries"]), "country_names": "; ".join(ev["country_names"]),
                  "country_method": ev["country_method"], "location_text": ev["location_text"],
                  "entry_words": ev["entry_words"], "entry_lang": ev["entry_lang"],
                  "has_link": ev["has_link"], "link_followed": ev["link_followed"],
                  "has_web_text": ev["has_web_text"], "n_pages": ev["n_pages"], "n_words_web": ev["n_words_web"],
                  "words_read_bin": ev["words_read_bin"]})
        for lv in ("L1", "L2", "L3", "L3s", "L3x"):
            r[f"n_orgs_{lv}"] = len(n_by_level[ev["event_id"]][lv])
        for t in th:
            for m in ("own", "landing", "linked", "any", "in_title", "core"):
                r[f"{t}_{m}"] = ev[f"{t}_{m}"]
        rows.append(r)
    fields = ["event_id", "event", "date", "start", "end", "year", "period", "event_link", "event_key", "link_host",
              "series_id", "stream_primary", "streams", "n_streams", "countries", "country_names", "country_method",
              "location_text", "entry_words", "entry_lang", "has_link", "link_followed", "has_web_text", "n_pages",
              "n_words_web", "n_orgs_L1", "n_orgs_L2", "n_orgs_L3"] + [
        f"{t}_{m}" for t in th for m in ("own", "landing", "linked", "any", "in_title")] + [
        "words_read_bin", "n_orgs_L3s", "n_orgs_L3x"] + [f"{t}_core" for t in th]
    an.table("events_enriched", rows, fields)
    # ---- ties
    rows = []
    for t in an.ties:
        rows.append({"event_id": t["event_id"], "pseudo_label": t["pseudo_label"], "series_id": t["series_id"],
                     "org_key": t["org_key"], "organisation": t["organisation"],
                     "canonical_org": an.org_name.get(t["org_key"], t["organisation"]),
                     "actor_class": t["actor_class"], "confidence": t["confidence"], "tie_level": t["tie_level"],
                     "layer": t["layer"], "source_link": t["source_link"], "document": t["document"],
                     "site": t["site"], "named_in_event_title": t["named_in_event_title"], "in_menus": t["in_menus"],
                     "year": t["year"], "countries": "; ".join(t["countries"]),
                     "doc_year": t["doc_year"] if t["doc_year"] else "undated",
                     "source_has_theme": t["pass"]["source"], "passes_theme_filter": t["pass"][an.a.theme_filter],
                     "found_by": t["found_by"], "match": t["match"], "tie_sublevel": t["tie_sublevel"],
                     "doc_kind": t["doc_kind"], "attribution": t["attribution"],
                     "participation": t["participation"], "context_reason": t["context_reason"],
                     "duplicate_document_of": t["duplicate_document_of"], "translation_of": t["translation_of"],
                     "sheet": t["sheet"]})
    an.table("ties", rows, ["event_id", "pseudo_label", "series_id", "org_key", "organisation", "canonical_org",
                            "actor_class", "confidence", "tie_level", "layer", "source_link", "document", "site",
                            "named_in_event_title", "in_menus", "year", "countries", "doc_year", "source_has_theme",
                            "passes_theme_filter", "found_by", "match", "tie_sublevel", "doc_kind", "attribution",
                            "participation", "context_reason", "duplicate_document_of", "translation_of", "sheet"])
    # ---- theme units: (passage_id, theme, event)
    units: "OrderedDict[tuple, dict]" = OrderedDict()
    for inst in an.instances:
        for rf in inst["refs"]:
            key = (inst["passage_id"], inst["theme"], rf["event_id"] or rf["pseudo_label"])
            u = units.get(key)
            ev = an.ev_by_id.get(rf["event_id"])
            if u is None:
                u = units[key] = {
                    "passage_id": inst["passage_id"], "theme": inst["theme"], "_terms": [], "_fam": [],
                    "language_section": inst["language_section"], "subtheme": inst["subtheme"],
                    "stance": inst["stance"], "event_id": rf["event_id"], "pseudo_label": rf["pseudo_label"],
                    "w": 0.0, "layer": rf["layer"], "tie_level": rf["tie_level"], "document": rf["document"],
                    "site": inst["site"], "is_pdf": inst["is_pdf"],
                    "doc_year": inst["doc_year"] or "undated", "event_year": ev["year"] if ev else "",
                    "furniture": inst["furniture"], "repeated": inst["repeated"], "duplicate": inst["dup"],
                    "clean": inst["clean"], "n_rows": 0,
                    "in_event_title": inst["in_event_title"], "in_handle": inst["in_handle"],
                    "orgs_in_paragraph_keys": "; ".join(inst["org_keys"]),
                    "context_before": inst["context_before"], "sentence": inst["sentence"],
                    "context_after": inst["context_after"], "source_link": inst["source_link"],
                    "tie_sublevel": rf["tie_sublevel"], "doc_kind": rf["doc_kind"]}
            u["_terms"].append(inst["term"])
            u["_fam"].append(inst["term_family"])
            u["w"] += rf["w"]
            u["n_rows"] += 1
            u["clean"] = u["clean"] or inst["clean"]
            u["furniture"] = u["furniture"] and inst["furniture"]
            u["repeated"] = u["repeated"] and inst["repeated"]
    rows = []
    for u in units.values():
        u["terms"] = "; ".join(dict.fromkeys(u.pop("_terms")))
        u["term_families"] = "; ".join(dict.fromkeys(u.pop("_fam")))
        u["w"] = round(u["w"], 4)
        rows.append(u)
    an.units = rows
    an.table("theme_units", rows, ["passage_id", "theme", "terms", "term_families", "language_section", "subtheme",
                                   "stance", "event_id", "pseudo_label", "w", "layer", "tie_level", "document", "site",
                                   "is_pdf", "doc_year", "event_year", "furniture", "repeated", "duplicate", "clean", "n_rows",
                                   "in_event_title", "in_handle", "orgs_in_paragraph_keys", "context_before",
                                   "sentence", "context_after", "source_link", "tie_sublevel", "doc_kind"])


# ==========================================================================
# A2 coverage frame and missingness
# ==========================================================================
def a2_coverage(an: Analysis) -> None:
    th = an.themes
    tf = an.a.theme_filter
    with_ties = {t["event_id"] for t in an.ties if t["event_id"] and not t["in_menus"] and t["pass"][tf]}
    groups = defaultdict(list)
    for ev in an.events:
        groups[ev["stream_primary"]].append(ev)
    groups["ALL"] = list(an.events)
    measures = [("with_link", lambda e: e["has_link"]), ("link_followed", lambda e: e["link_followed"]),
                ("has_web_text", lambda e: e["has_web_text"]), ("with_org_ties", lambda e: e["event_id"] in with_ties),
                ("with_theme_any", lambda e: e["any_any"]), ("country_resolved", lambda e: bool(e["countries"])),
                ("date_day_precision", lambda e: bool(re.fullmatch(r"\d{4}-\d\d-\d\d", e["start"])))]
    rows = []
    for st in sorted(groups, key=lambda x: (x == "ALL", x)):
        evs = groups[st]
        ys = [e["year"] for e in evs if e["year"]]
        r = {"stream": st, "years": f"{min(ys)}-{max(ys)}" if ys else "", "n_events": len(evs)}
        for name, fn in measures:
            k = sum(1 for e in evs if fn(e))
            lo, hi = wilson(k, len(evs))
            r.update({name: k, f"{name}_share": share(k, len(evs)), f"{name}_lo": lo, f"{name}_hi": hi})
        rows.append(r)
    fields = ["stream", "years", "n_events"] + [f"{m}{x}" for m, _ in measures for x in ("", "_share", "_lo", "_hi")]
    an.table("coverage_by_stream", rows, fields)
    # ---- missingness
    n = len(an.events)
    clean = [i for i in an.instances if i["clean"]]
    web_insts = [i for i in an.instances if i["site_role"] != "seed"]
    seed_clean = [i for i in clean if any(rf["layer"] == "listing_seed" for rf in i["refs"])]
    names = [e["event"] for e in an.events if len(e["event"]) >= 25]
    leak = 0
    for inst in seed_clean:
        own = {nm.casefold() for nm in inst["event_names"]}
        ctx = (inst["context_before"] + " " + inst["context_after"]).casefold()
        if any(nm.casefold() in ctx for nm in names if nm.casefold() not in own):
            leak += 1
    an.facts["seed_context_leak"] = (leak, len(seed_clean))
    net = [t for t in an.ties if t["sheet"] == "event_organisations"]
    n_refs = sum(sum(c.values()) for c in an.ref_counts.values())
    rows = [
        ("events without a country", sum(1 for e in an.events if not e["countries"]), n),
        ("events without a date", sum(1 for e in an.events if e["year"] is None), n),
        ("network ties to pseudo events ('(not tied to an event)', '(other pages of ...)')",
         sum(1 for t in net if not t["event_id"]), len(net)),
        ("network event refs that match no event (phantom splits, dropped)",
         sum(c["unmatched"] for c in an.ref_counts.values()), n_refs),
        ("clean theme rows with no organisation in the paragraph", sum(1 for i in clean if not i["orgs_text"]),
         len(clean)),
        ("web theme rows without a document year", sum(1 for i in web_insts if not i["doc_year"]), len(web_insts)),
        ("low-confidence network rows", sum(1 for t in net if t["confidence"] == "low"), len(net)),
        ("events listed in both agenda PDFs", an.facts.get("events_in_two_agendas", 0), n),
        ("seed theme rows whose context names another event (caveat C6)", leak, len(seed_clean)),
        (f"network rows left out by the theme filter '{tf}'", an.filter_counts["rows"][1] - an.filter_counts["rows"][0],
         an.filter_counts["rows"][1]),
    ]
    an.table("missingness", [{"field": f, "k": k, "n": m, "pct": round(100 * k / m, 1) if m else None}
                             for f, k, m in rows], ["field", "k", "n", "pct"])
    _ = th


# ==========================================================================
# A3 measurement ladder and concentration
# ==========================================================================
def _clean_inst_docs(an: Analysis):
    """Clean instances with their document (rule 7: one agenda entry = one document)."""
    out = []
    for inst in an.instances:
        if inst["clean"]:
            out.append((inst, inst["refs"][0]["document"] if inst["refs"] else inst["source_link"]))
    return out


def a3_ladder(an: Analysis) -> None:
    th = an.themes
    cd = _clean_inst_docs(an)
    counts = {}
    for t in th:
        insts = [i for i in an.instances if i["theme"] == t]
        cl = [(i, d) for i, d in cd if i["theme"] == t]
        evs_any = {rf["event_id"] for i, _ in cl for rf in i["refs"] if rf["event_id"]}
        evs_own = {rf["event_id"] for i, _ in cl for rf in i["refs"] if rf["event_id"] and rf["layer"] in OWN_LAYERS}
        counts[t] = OrderedDict([("raw_rows", len(insts)), ("clean_rows", len(cl)),
                                 ("unique_passages", len({i["passage_id"] for i, _ in cl})),
                                 ("documents", len({d for _, d in cl})), ("sites", len({i["site"] for i, _ in cl})),
                                 ("events_any", len(evs_any)), ("events_own", len(evs_own))])
    rows = []
    for lev in ["raw_rows", "clean_rows", "unique_passages", "documents", "sites", "events_any", "events_own"]:
        tot = sum(counts[t][lev] for t in th)
        for t in th:
            rows.append({"theme": t, "level": lev, "count": counts[t][lev], "share_of_theme_total": share(counts[t][lev], tot)})
    an.table("measurement_ladder", rows, ["theme", "level", "count", "share_of_theme_total"])
    an.ladder = counts
    # ---- concentration
    rows = []
    for t in ["ALL"] + th:
        cl = [(i, d) for i, d in cd if t == "ALL" or i["theme"] == t]
        c = Counter(d for _, d in cl)
        tot = sum(c.values())
        ev_w = Counter()
        for i, _ in cl:
            for rf in i["refs"]:
                if rf["event_id"]:
                    ev_w[rf["event_id"]] += rf["w"]
        if not tot:
            rows.append({"theme": t, "documents": 0})
            continue
        top = sorted(c.items(), key=lambda x: (-x[1], x[0]))
        te = sorted(ev_w.items(), key=lambda x: (-x[1], x[0]))
        rows.append({"theme": t, "documents": len(c), "clean_rows": tot, "top1_doc": top[0][0],
                     "top1_share": share(top[0][1], tot), "top5_share": share(sum(v for _, v in top[:5]), tot),
                     "eff_docs": round(1 / sum((v / tot) ** 2 for v in c.values()), 2),
                     "top_event": an.ev_by_id[te[0][0]]["event"] if te else "",
                     "top_event_share": round(te[0][1] / sum(ev_w.values()), 4) if te else None})
    an.table("concentration", rows, ["theme", "documents", "clean_rows", "top1_doc", "top1_share", "top5_share",
                                     "eff_docs", "top_event", "top_event_share"])
    an.concentration = {r["theme"]: r for r in rows}
    # ---- attribution: layer x theme
    rows = []
    agg = defaultdict(lambda: [0, set()])
    for inst in an.instances:
        if not inst["clean"]:
            continue
        for rf in inst["refs"]:
            a = agg[(rf["layer"], inst["theme"])]
            a[0] += 1
            a[1].add(rf["document"])
    for (lay, t), (k, docs) in sorted(agg.items()):
        rows.append({"layer": lay, "tie_level": TIE_LEVEL.get(lay, "L3"), "theme": t, "clean_refs": k,
                     "documents": len(docs)})
    an.table("attribution", rows, ["layer", "tie_level", "theme", "clean_refs", "documents"])
    # ---- term families
    fam = defaultdict(lambda: {"rows": 0, "clean": 0, "passages": set(), "docs": set(), "events": set(),
                               "forms": Counter(), "title": 0, "handle": 0, "years": []})
    for inst in an.instances:
        f = fam[(inst["theme"], inst["term_family"], inst["language_section"])]
        f["rows"] += 1
        f["forms"][inst["term"].casefold()] += 1
        f["title"] += inst["in_event_title"]
        f["handle"] += inst["in_handle"]
        if inst["clean"]:
            f["clean"] += 1
            f["passages"].add(inst["passage_id"])
            f["docs"].add(inst["refs"][0]["document"] if inst["refs"] else inst["source_link"])
            for rf in inst["refs"]:
                if rf["event_id"]:
                    f["events"].add(rf["event_id"])
                    y = an.ev_by_id[rf["event_id"]]["year"]
                    if y:
                        f["years"].append(y)
    rows = []
    for (t, tfam, sec), f in sorted(fam.items(), key=lambda x: (x[0][0], -x[1]["rows"], x[0][1])):
        rows.append({"theme": t, "term_family": tfam, "language_section": sec, "rows": f["rows"],
                     "clean_rows": f["clean"], "passages": len(f["passages"]), "documents": len(f["docs"]),
                     "events": len(f["events"]),
                     "surface_forms": "; ".join(f"{k} ({v})" for k, v in f["forms"].most_common(8)),
                     "in_title_rows": f["title"], "handle_rows": f["handle"],
                     "first_year": min(f["years"]) if f["years"] else "", "last_year": max(f["years"]) if f["years"] else ""})
    an.table("term_families", rows, ["theme", "term_family", "language_section", "rows", "clean_rows", "passages",
                                     "documents", "events", "surface_forms", "in_title_rows", "handle_rows",
                                     "first_year", "last_year"])
    an.facts["unmapped_terms"] = sum(1 for i in an.instances if i["term_family"] == "(unmapped)")


def doc_hygiene(an: Analysis) -> None:
    """Sources that name nearly the same organisations (Jaccard >= 0.9, at least 15 each)
    but are tied to different series: copies or near copies the pipeline did not mark
    (duplicate_document_of), which would tie the same names to two series. Plus the
    series whose editions carry different names (multi_name_series)."""
    org_sets: Dict[str, Set[str]] = defaultdict(set)
    ser: Dict[str, Set[str]] = defaultdict(set)
    for t in an.ties:
        if t["event_id"] and not t["in_menus"] and t["layer"] not in ("listing_seed", "listing_calendar"):
            org_sets[t["unit"]].add(t["org_key"])
            ser[t["unit"]].add(t["series_id"])
    big = sorted(u for u, v in org_sets.items() if len(v) >= 15)
    inv: Dict[str, List[str]] = defaultdict(list)
    for u in big:
        for k in org_sets[u]:
            inv[k].append(u)
    rows = []
    for i, u in enumerate(big):
        cand = Counter(v for k in org_sets[u] for v in inv[k] if v > u)
        for v, inter in cand.items():
            a, b = org_sets[u], org_sets[v]
            if inter / (len(a) + len(b) - inter) >= 0.9 and ser[u] != ser[v]:
                rows.append({"source_a": u, "source_b": v, "n_a": len(a), "n_b": len(b),
                             "jaccard": round(inter / (len(a) + len(b) - inter), 3),
                             "series_a": "; ".join(sorted(ser[u]))[:300], "series_b": "; ".join(sorted(ser[v]))[:300]})
    rows.sort(key=lambda r: (-r["jaccard"], -r["n_a"], r["source_a"], r["source_b"]))
    an.table("similar_sources", rows, ["source_a", "source_b", "n_a", "n_b", "jaccard", "series_a", "series_b"])
    an.note("source pairs with nearly the same organisations, tied to different series", len(rows),
            "Jaccard >= 0.9, at least 15 organisations each; see similar_sources")
    an.table("multi_name_series", an.multi_name_series, ["series_id", "n_events", "n_names", "names"])


def a3b_documents(an: Analysis) -> None:
    """Per event: distinct documents and passages with a (focus) theme hit, by tie level,
    and whether one document gives more than half of the event's theme rows (a long
    report can carry hundreds of rows). Plus the descriptive themes, reported apart."""
    per: Dict[str, dict] = {}
    for inst in an.instances:
        if inst["furniture"]:      # as the event theme flags: repeated passages count, furniture does not
            continue
        for rf in inst["refs"]:
            if not rf["event_id"]:
                continue
            p = per.setdefault(rf["event_id"], {"rows": 0, "docs": Counter(), "pass": set(),
                                                "sd": defaultdict(set), "sp": defaultdict(set)})
            p["rows"] += 1
            p["docs"][rf["document"]] += 1
            p["pass"].add(inst["passage_id"])
            p["sd"][rf["tie_sublevel"]].add(rf["document"])
            p["sp"][rf["tie_sublevel"]].add(inst["passage_id"])
    rows = []
    for eid, p in per.items():
        ev = an.ev_by_id[eid]
        top, k = sorted(p["docs"].items(), key=lambda x: (-x[1], x[0]))[0]
        r = {"event_id": eid, "event": ev["event"], "year": ev["year"], "theme_rows": p["rows"],
             "documents": len(p["docs"]), "passages": len(p["pass"])}
        for lv in SUBLEVELS:
            r[f"documents_{lv}"] = len(p["sd"].get(lv, ()))
            r[f"passages_{lv}"] = len(p["sp"].get(lv, ()))
        r.update({"top_document": pdf_display(top), "top_document_rows": k, "top_document_share": share(k, p["rows"]),
                  "one_document_dominates": k / p["rows"] > 0.5 and len(p["docs"]) >= 2})
        rows.append(r)
    rows.sort(key=lambda r: (not r["one_document_dominates"], -r["theme_rows"], r["event_id"]))
    an.table("event_theme_documents", rows, ["event_id", "event", "year", "theme_rows", "documents", "passages"]
             + [f"{x}_{lv}" for lv in SUBLEVELS for x in ("documents", "passages")]
             + ["top_document", "top_document_rows", "top_document_share", "one_document_dominates"])
    flagged = [r for r in rows if r["one_document_dominates"]]
    an.facts["events_one_document_dominates"] = len(flagged)
    an.note("events where one document gives more than half of the theme rows", len(flagged),
            "; ".join(f"{short(r['event'], 40)} ({r['top_document_rows']}/{r['theme_rows']})" for r in flagged[:5]))
    # ---- descriptive themes (global_south, authoritarianism): counted, never filtering
    focus_ev = {rf["event_id"] for i in an.instances if not i["furniture"] for rf in i["refs"] if rf["event_id"]}
    drows = []
    for th in an.descriptive:
        insts = [i for i in an.desc_instances if i["theme"] == th]
        cl = [i for i in insts if i["clean"]]
        evs = {rf["event_id"] for i in insts if not i["furniture"] for rf in i["refs"] if rf["event_id"]}
        forms = Counter(i["term"].casefold() for i in cl)
        drows.append({"theme": th, "raw_rows": len(insts), "clean_rows": len(cl),
                      "passages": len({i["passage_id"] for i in cl}),
                      "documents": len({i["refs"][0]["document"] if i["refs"] else i["source_link"] for i in cl}),
                      "sites": len({i["site"] for i in cl}), "events_any": len(evs),
                      "events_with_a_focus_theme": len(evs & focus_ev), "events_without_a_focus_theme": len(evs - focus_ev),
                      "surface_forms": "; ".join(f"{k} ({v})" for k, v in forms.most_common(8))})
    if not an.descriptive:
        an.skipped["descriptive_themes"] = "no descriptive theme in the masters or the theme file"
    an.table("descriptive_themes", drows, ["theme", "raw_rows", "clean_rows", "passages", "documents", "sites",
                                           "events_any", "events_with_a_focus_theme", "events_without_a_focus_theme",
                                           "surface_forms"])


# ==========================================================================
# A4 theme prevalence (event as unit)
# ==========================================================================
# Event-level measures. own, any_all and in_title are counted over all events of the group;
# landing, linked and any_web only over the events whose link was read (has_web_text).
# core: the listing (L1), the landing page (L2) and pages and PDFs on the event's own host
# (L3s), without linked reports; less dependent on how much linked text was read.
MEASURES = ["own", "landing", "linked", "any_web", "any_all", "in_title", "core"]
WEB_MEASURES = {"landing", "linked", "any_web", "core"}
MEASURE_FLAG = {"any_web": "any", "any_all": "any"}


def _std_share(an: Analysis, base: List[dict], t: str, flag: str) -> Optional[float]:
    """Direct standardisation over the words-read bins: the group's share in each bin,
    weighted by the bin distribution of all events with web text (a bin the group has
    no event in takes the pooled share of that bin)."""
    std, pooled = an.std_bins
    if not base or not std:
        return None
    by = defaultdict(lambda: [0, 0])
    for e in base:
        b = by[e["words_read_bin"]]
        b[0] += 1
        b[1] += bool(e[f"{t}_{flag}"])
    tot = sum(std.values())
    val = 0.0
    for b, nb in std.items():
        n, k = by[b] if b in by else (0, 0)
        rate = k / n if n else pooled.get((t, flag, b), 0.0)
        val += nb / tot * rate
    return round(val, 4)


def _prev_rows(an: Analysis, dimension: str, groups: Dict[tuple, List[dict]], measures=MEASURES) -> List[dict]:
    rows = []
    for (grp, st) in sorted(groups, key=lambda x: (str(x[0]), str(x[1]))):
        evs = groups[(grp, st)]
        web = [e for e in evs if e["has_web_text"]]
        for t in an.themes + ["any"]:
            for m in measures:
                base = web if m in WEB_MEASURES else evs
                flag = MEASURE_FLAG.get(m, m)
                k = sum(1 for e in base if e[f"{t}_{flag}"])
                lo, hi = wilson(k, len(base))
                rows.append({"dimension": dimension, "group": grp, "stream": st, "theme": t, "measure": m,
                             "denominator": "events with web text" if m in WEB_MEASURES else "all events",
                             "n_events": len(base), "k_events": k, "share": share(k, len(base)),
                             "wilson_lo": lo, "wilson_hi": hi,
                             "share_std": _std_share(an, base, t, flag) if m in WEB_MEASURES and
                             dimension != "words_read_bin" else None})
    return rows


PREV_FIELDS = ["dimension", "group", "stream", "theme", "measure", "denominator", "n_events", "k_events", "share", "wilson_lo",
               "wilson_hi", "share_std"]


def a4_prevalence(an: Analysis) -> None:
    # Exposure: events with web text by words read (the standard for share_std)
    web_all = [e for e in an.events if e["has_web_text"]]
    std = Counter(e["words_read_bin"] for e in web_all)
    pooled = {}
    for t in an.themes + ["any"]:
        for flag in {MEASURE_FLAG.get(m, m) for m in WEB_MEASURES}:
            for b, nb in std.items():
                k = sum(1 for e in web_all if e["words_read_bin"] == b and e[f"{t}_{flag}"])
                pooled[(t, flag, b)] = k / nb
    an.std_bins = (dict(std), pooled)
    order = [lab for _, lab in WORDS_BINS]
    gb = defaultdict(list)
    for e in an.events:
        gb[(e["words_read_bin"], "all")].append(e)
    rows = _prev_rows(an, "words_read_bin", gb, sorted(WEB_MEASURES))
    rows.sort(key=lambda r: (order.index(r["group"]), r["theme"], r["measure"]))
    an.table("prev_words_read", rows, PREV_FIELDS)
    grad = {r["group"]: r for r in rows if r["theme"] == "any" and r["measure"] == "any_web" and r["n_events"]}
    an.facts["words_read_gradient"] = {b: (grad[b]["k_events"], grad[b]["n_events"], grad[b]["share"])
                                       for b in order if b in grad}
    g = defaultdict(list)
    gy = defaultdict(list)
    for ev in an.events:
        for st in (ev["stream_primary"], "all"):
            g[(ev["period"], st)].append(ev)
            gy[(ev["year"] or "undated", st)].append(ev)
    # Sensitivity: an event listed in several agendas counts in each of them.
    if any(len(ev["streams"]) > 1 for ev in an.events):
        for ev in an.events:
            for st in ev["streams"]:
                g[(ev["period"], f"{st} (all listings)")].append(ev)
    an.table("prev_stream_period", _prev_rows(an, "period", g), PREV_FIELDS)
    an.prev_sp = an.tables["prev_stream_period"][0]
    an.table("prev_stream_year", _prev_rows(an, "year", gy), PREV_FIELDS)
    # ---- country: every country of an event; groups below 15 events pooled as "other"
    cnt = Counter(c for ev in an.events for c in (ev["countries"] or ["(unresolved)"]))
    gc = defaultdict(list)
    for ev in an.events:
        for c in ev["countries"] or ["(unresolved)"]:
            lab = c if (cnt[c] >= 15 or c == "(unresolved)") else "other (n<15 each)"
            gc[(lab, "all")].append(ev)
    rows = _prev_rows(an, "country", gc, ["own", "any_web", "any_all"])
    for r in rows:
        r["group_name"] = an.p.country_name(r["group"]) if re.fullmatch(r"[A-Z]{2,6}", r["group"]) else r["group"]
        r["headline_ok"] = r["n_events"] >= 30
    an.table("prev_country", rows, PREV_FIELDS[:2] + ["group_name"] + PREV_FIELDS[2:] + ["headline_ok"])
    gl = defaultdict(list)
    for ev in an.events:
        gl[(ev["entry_lang"] or "(no entry text)", ev["stream_primary"])].append(ev)
    an.table("prev_lang", _prev_rows(an, "entry_lang", gl, ["own", "in_title"]), PREV_FIELDS)
    # ---- entry-corpus rates (agenda entries of the seed PDFs)
    if an.entries and an.ts:
        p = an.p
        agg = defaultdict(lambda: {"entries": 0, "words": 0, "hits": Counter(), "with": Counter()})
        total_hits = 0
        for en in an.entries:
            text = _join_entry_lines(en["lines"])
            words = len(_WORD.findall(p.URL_LIKE.sub(" ", text)))
            hits = an.ts.find(text)
            y = year_of(en["start"] or en["month"])
            for key in ((en["stream"], period_of(y)), (en["stream"], "all"), ("all", "all")):
                a = agg[key]
                a["entries"] += 1
                a["words"] += words
                for t in an.themes:
                    k = len(hits.get(t, []))
                    a["hits"][t] += k
                    a["with"][t] += k > 0
                a["hits"]["any"] += sum(len(v) for v in hits.values())
                a["with"]["any"] += bool(hits)
            total_hits += sum(len(v) for v in hits.values())
        rows = []
        for (st, per), a in sorted(agg.items()):
            for t in an.themes + ["any"]:
                k = a["hits"][t]
                lo, hi = byar(k)
                w = a["words"]
                ke = a["with"][t]
                wl, wh = wilson(ke, a["entries"])
                rows.append({"stream": st, "period": per, "theme": t, "entries": a["entries"], "words": w, "hits": k,
                             "per10k": round(1e4 * k / w, 3) if w else None,
                             "byar_lo": round(1e4 * lo / w, 3) if w else None,
                             "byar_hi": round(1e4 * hi / w, 3) if w else None,
                             "entries_with_theme": ke, "share": share(ke, a["entries"]), "wilson_lo": wl,
                             "wilson_hi": wh})
        an.table("entry_rates", rows, ["stream", "period", "theme", "entries", "words", "hits", "per10k", "byar_lo",
                                       "byar_hi", "entries_with_theme", "share", "wilson_lo", "wilson_hi"])
        seed_rows = sum(1 for i in an.instances if i["site_role"] == "seed")
        an.facts["entry_hits"] = (total_hits, seed_rows)
        an.note("entry re-run theme hits vs seed theme rows in master_themes", f"{total_hits} vs {seed_rows}",
                "equal" if total_hits == seed_rows else "differ (line joins, menu filter or a newer theme file)")
        if total_hits != seed_rows:
            an.warn("Agenda entries re-read give %d theme hits, master_themes has %d seed rows "
                    "(differences come from line joins or a changed theme file)", total_hits, seed_rows)
    elif "entry_rates" not in an.skipped:
        an.skipped["entry_rates"] = "no agenda entries or no theme file"
    # ---- per-page rates (pages.csv / sources.csv)
    src = an.sources_rows or an.pages
    # A seed PDF (or calendar page) is listed once as a whole and once per agenda entry
    # (or cell); count only the parts, so words are not counted twice.
    split_links = {r.get("source_link", "") for r in src if "#" in s(r.get("source_unit"))}
    src = [r for r in src if not (r.get("source_unit") and "#" not in r["source_unit"]
                                  and r.get("source_link", "") in split_links)]
    if src:
        agg = defaultdict(lambda: {"n": 0, "words": 0, "hits": Counter(), "with": Counter()})
        for pg in src:
            role = pg.get("site_role", "") or "?"
            kind = pg.get("kind", "") or ("pdf" if truthy(pg.get("is_pdf")) else "page")
            w = int(pg["words"]) if s(pg.get("words")).isdigit() else 0
            for key in ((role, kind), ("all", "all")):
                a = agg[key]
                a["n"] += 1
                a["words"] += w
                anyh = 0
                for t in an.themes:
                    v = s(pg.get(f"theme:{t}"))
                    k = int(v) if v.isdigit() else 0
                    a["hits"][t] += k
                    a["with"][t] += k > 0
                    anyh += k
                a["hits"]["any"] += anyh
                a["with"]["any"] += anyh > 0
        rows = []
        for (role, kind), a in sorted(agg.items()):
            for t in an.themes + ["any"]:
                k = a["hits"][t]
                lo, hi = byar(k)
                wl, wh = wilson(a["with"][t], a["n"])
                rows.append({"site_role": role, "kind": kind, "theme": t, "sources": a["n"], "words": a["words"],
                             "hits": k, "per10k": round(1e4 * k / a["words"], 3) if a["words"] else None,
                             "byar_lo": round(1e4 * lo / a["words"], 3) if a["words"] else None,
                             "byar_hi": round(1e4 * hi / a["words"], 3) if a["words"] else None,
                             "sources_with_theme": a["with"][t], "share": share(a["with"][t], a["n"]),
                             "wilson_lo": wl, "wilson_hi": wh})
        an.table("page_rates", rows, ["site_role", "kind", "theme", "sources", "words", "hits", "per10k", "byar_lo",
                                      "byar_hi", "sources_with_theme", "share", "wilson_lo", "wilson_hi"])
    else:
        an.skipped["page_rates"] = "pages.csv / sources.csv not found"


# ==========================================================================
# A5 / A6 bipartite network, projection, backbone, centralities
# ==========================================================================
def _org_conf(an: Analysis) -> Dict[str, int]:
    if not hasattr(an, "_conf"):
        c: Dict[str, int] = {}
        for t in an.ties:
            c[t["org_key"]] = max(c.get(t["org_key"], 0), CONF_RANK.get(t["confidence"], 0))
        an._conf = c
    return an._conf


def in_scope(t: dict, scope: str) -> bool:
    """participation: evidence of taking part (default); reports: also rows whose only
    context reason is a linked report; all: every row (context rows too)."""
    if scope == "all" or t["participation"]:
        return True
    return scope == "reports" and bool(t["ctx_kinds"]) and t["ctx_kinds"] <= {"report"}


def scenario_ties(an: Analysis, sc: str, layers: Optional[str] = None, tf: Optional[str] = None,
                  scope: Optional[str] = None) -> List[dict]:
    tf = tf or an.a.theme_filter
    scope = scope or an.a.network_scope
    min_rank = CONF_RANK[an.a.min_confidence]
    conf = _org_conf(an)
    keep_levels = LAYER_SETS[layers or an.a.layers]
    out = []
    for t in an.ties:
        if not t["event_id"] or t["in_menus"] or not t["pass"][tf]:
            continue
        if t["tie_level"] not in keep_levels and t["tie_sublevel"] not in keep_levels:
            continue
        if not in_scope(t, scope):
            continue
        k = t["org_key"]
        if org_dropped(an, k):
            continue
        title = t["named_in_event_title"]
        conf_ok = conf.get(k, 0) >= min_rank or s(an.codes.get(k, {}).get("include")) in ("1", "yes")
        if sc in ("B", "D") and title:
            continue
        if sc in ("C", "D", "C_curated") and not conf_ok:
            continue
        if sc == "C_curated" and org_class(an, k) in DROP_CLASSES:
            continue
        out.append(t)
    return out


def bipartite(an: Analysis, ties: List[dict], level: str = "series"):
    import networkx as nx
    G = nx.Graph()
    agg: "OrderedDict[tuple, dict]" = OrderedDict()
    ev_info: Dict[str, dict] = {}
    for t in ties:
        node = t["series_id"] if level == "series" else t["event_id"]
        e, o = "E:" + node, "O:" + t["org_key"]
        ev = an.ev_by_id[t["event_id"]]
        ei = ev_info.setdefault(e, {"label": ev["event"], "editions": set(), "years": set(), "countries": set(),
                                    "streams": set()})
        ei["editions"].add(t["event_id"])
        if ev["year"]:
            ei["years"].add(ev["year"])
        ei["countries"].update(ev["countries"])
        ei["streams"].update(ev["streams"])
        a = agg.setdefault((e, o), {"editions": set(), "sources": set(), "docs": set(), "levels": set(),
                                    "sub": set(), "title": False, "conf": 0, "years": set()})
        a["editions"].add(t["event_id"])
        a["sources"].add(t["unit"])
        a["docs"].add(t["document"])
        a["levels"].add(t["tie_level"])
        a["sub"].add(t["tie_sublevel"])
        a["title"] = a["title"] or t["named_in_event_title"]
        a["conf"] = max(a["conf"], CONF_RANK.get(t["confidence"], 0))
        if ev["year"]:
            a["years"].add(ev["year"])
    inv = {v: k for k, v in CONF_RANK.items()}
    for e, ei in ev_info.items():
        G.add_node(e, kind="event", bipartite=0, label=short(ei["label"], 80), n_editions=len(ei["editions"]),
                   first_year=min(ei["years"]) if ei["years"] else "", last_year=max(ei["years"]) if ei["years"] else "",
                   countries="; ".join(sorted(ei["countries"])), streams="; ".join(sorted(ei["streams"])))
    for (e, o), a in agg.items():
        if o not in G:
            k = o[2:]
            G.add_node(o, kind="organisation", bipartite=1, label=an.org_name.get(k, k),
                       actor_class=org_class(an, k))
        G.add_edge(e, o, n_editions=len(a["editions"]), n_sources=len(a["sources"]), n_docs=len(a["docs"]),
                   tie_levels="; ".join(sorted(a["levels"])), min_tie_level=min(a["levels"]),
                   tie_sublevel="; ".join(x for x in SUBLEVELS if x in a["sub"]),
                   title_named=a["title"], confidence=inv[a["conf"]],
                   first_year=min(a["years"]) if a["years"] else "", last_year=max(a["years"]) if a["years"] else "")
    return G


class Projection:
    """Newman-weighted organisation projection held as arrays (numpy): an event with
    k >= 2 organisations adds 1/(k-1) to every pair of them; `shared` counts the
    events. Pair i < j of the sorted organisation list is edge (a[x], b[x]) with
    weight W[x] (rounded to 6 places) and C[x] shared events. Only what the analysis
    needs is offered (degree, membership, edge count, one edge's weights); networkx
    graphs are built for the backbone only, a full projection graph of millions of
    pairs would not fit in memory."""

    def __init__(self, G):
        import numpy as np
        self.G = G
        self.orgs = sorted(n for n in G if n.startswith("O:"))
        idx = {o: i for i, o in enumerate(self.orgs)}
        n = len(self.orgs)
        codes, ws = [], []
        for e in sorted(x for x in G if x.startswith("E:")):
            nb = np.array(sorted(idx[m] for m in G[e]), dtype=np.int64)
            k = len(nb)
            if k < 2:
                continue
            i, j = np.triu_indices(k, 1)
            codes.append(nb[i] * n + nb[j])
            ws.append(np.full(len(i), 1.0 / (k - 1)))
        if codes:
            # bincount adds the weights in event order, as the pair-by-pair loop did
            u, inv = np.unique(np.concatenate(codes), return_inverse=True)
            self.W = np.round(np.bincount(inv, weights=np.concatenate(ws)), 6)
            self.C = np.bincount(inv)
            self.a, self.b = u // n, u % n
        else:
            self.a = self.b = self.C = np.zeros(0, dtype=np.int64)
            self.W = np.zeros(0)
        deg = np.bincount(self.a, minlength=n) + np.bincount(self.b, minlength=n)
        self._deg = {o: int(deg[i]) for i, o in enumerate(self.orgs)}
        self._idx = idx

    def __contains__(self, node) -> bool:
        return node in self._deg

    def __len__(self) -> int:
        return len(self.orgs)

    def degree(self, node) -> int:
        return self._deg[node]

    def number_of_edges(self) -> int:
        return int(len(self.a))

    def edge(self, u: str, v: str) -> Optional[dict]:
        """{'weight', 'shared'} of the pair u-v, or None."""
        import numpy as np
        i, j = sorted((self._idx[u], self._idx[v]))
        x = np.flatnonzero((self.a == i) & (self.b == j))
        return {"weight": float(self.W[x[0]]), "shared": int(self.C[x[0]])} if len(x) else None


def project(G) -> Projection:
    """Newman-weighted organisation projection of a bipartite event x organisation graph."""
    return Projection(G)


def _disparity_mask(n: int, a, b, W, alpha: float):
    """Disparity filter (Serrano et al. 2009) on edge arrays: keep edge ij if
    (1 - w_ij/s_i)^(k_i - 1) < alpha from either endpoint with degree >= 2."""
    import numpy as np
    deg = np.bincount(a, minlength=n) + np.bincount(b, minlength=n)
    st = np.bincount(a, weights=W, minlength=n) + np.bincount(b, weights=W, minlength=n)
    keep = np.zeros(len(a), bool)
    for x in (a, b):
        k = deg[x]
        ok = k >= 2
        p = np.ones(len(a))
        p[ok] = (1 - W[ok] / st[x[ok]]) ** (k[ok] - 1)
        keep |= ok & (p < alpha)
    return keep


def disparity(P, alpha: float):
    """Disparity backbone as a networkx graph (nodes keep the attributes of the
    bipartite graph; edges keep weight and shared). P is a Projection, or a weighted
    networkx graph."""
    import networkx as nx
    import numpy as np
    if not isinstance(P, Projection):
        nodes = sorted(P)
        idx = {o: i for i, o in enumerate(nodes)}
        rows = sorted((min(idx[u], idx[v]), max(idx[u], idx[v]), float(d.get("weight", 1.0)), int(d.get("shared", 1)))
                      for u, v, d in P.edges(data=True))
        a = np.array([r[0] for r in rows], dtype=np.int64)
        b = np.array([r[1] for r in rows], dtype=np.int64)
        W = np.array([r[2] for r in rows], dtype=float)
        C = np.array([r[3] for r in rows], dtype=np.int64)
        attrs = P.nodes
    else:
        nodes, a, b, W, C, attrs = P.orgs, P.a, P.b, P.W, P.C, P.G.nodes
    k = _disparity_mask(len(nodes), a, b, W, alpha)
    B = nx.Graph()
    for i, j, w, c in sorted((nodes[x], nodes[y], float(w), int(c)) for x, y, w, c in zip(a[k], b[k], W[k], C[k])):
        B.add_node(i, **attrs[i])
        B.add_node(j, **attrs[j])
        B.add_edge(i, j, weight=round(w, 6), shared=c)
    return B


_STATS_CACHE: Dict[tuple, dict] = {}


def graph_stats(G, alpha: float) -> dict:
    """Size, components, projection and backbone figures of a bipartite graph. Cached
    by its tie set: identical selections (e.g. the 'none' rows of filtered masters)
    are computed once."""
    import networkx as nx
    import numpy as np
    key = (alpha, frozenset((u, v) if u < v else (v, u) for u, v in G.edges()), frozenset(G.nodes()))
    hit = _STATS_CACHE.get(key)
    if hit is not None:
        return dict(hit)
    orgs = [n for n in G if n.startswith("O:")]
    evs = [n for n in G if n.startswith("E:")]
    comps = sorted((len(c) for c in nx.connected_components(G)), reverse=True) if len(G) else []
    P = project(G)
    n = len(P.orgs)
    top_share = None
    if evs and P.number_of_edges():
        big = max(sorted(evs), key=G.degree)
        S = np.zeros(n, bool)
        for m in G[big]:
            S[P._idx[m]] = True
        top_share = round(float((S[P.a] & S[P.b]).sum()) / P.number_of_edges(), 4)
    out = {"orgs": len(orgs), "nodes_events": len(evs), "orgs_in_2plus": sum(1 for o in orgs if G.degree(o) >= 2),
           "edges": G.number_of_edges(), "components": len(comps), "giant_nodes": comps[0] if comps else 0,
           "proj_edges": P.number_of_edges(), "top_event_share_of_proj_edges": top_share}
    for al, tag in ((alpha, "backbone"), (0.2, "backbone02")):
        k = _disparity_mask(n, P.a, P.b, P.W, al)
        out[tag + "_nodes"] = int(len(np.union1d(P.a[k], P.b[k])))
        out[tag + "_edges"] = int(k.sum())
    _STATS_CACHE[key] = out
    return dict(out)


SENS_FIELDS = ["theme_filter", "master_theme_filter", "scenario", "description", "variant", "layers", "node_level",
               "rows", "events", "series", "orgs", "orgs_in_2plus", "edges", "components", "giant_nodes",
               "proj_edges", "top_event_share_of_proj_edges", "backbone_nodes", "backbone_edges",
               "backbone_a02_nodes", "backbone_a02_edges"]
SCOPE_LABEL = {"participation": "participation (default)", "all": "all rows (--network-scope all)"}


def cap_ties(ties: List[dict], n: int) -> List[dict]:
    """At most n organisations per source: most mentions first, then confidence, then key."""
    best: Dict[str, Dict[str, tuple]] = defaultdict(dict)
    for t in ties:
        sc = (t["mentions"], CONF_RANK.get(t["confidence"], 0))
        if sc > best[t["unit"]].get(t["org_key"], (-1, -1)):
            best[t["unit"]][t["org_key"]] = sc
    keep = {u: {k for k, _ in sorted(d.items(), key=lambda x: (-x[1][0], -x[1][1], x[0]))[:n]}
            for u, d in best.items()}
    return [t for t in ties if t["org_key"] in keep[t["unit"]]]


def l4_ties(an: Analysis, sc: str) -> List[dict]:
    """Context rows with attribution site_level ('(other pages of SITE)') attributed to the
    site's series when the site's event-tied rows all belong to one series (tie level L4)."""
    site_series: Dict[str, Set[str]] = defaultdict(set)
    site_event: Dict[str, str] = {}
    for t in an.ties:
        if t["event_id"] and not t["in_menus"]:
            site_series[t["site"]].add(t["series_id"])
            site_event[t["site"]] = min(site_event.get(t["site"], t["event_id"]), t["event_id"])
    min_rank = CONF_RANK[an.a.min_confidence]
    conf = _org_conf(an)
    out = []
    for t in an.ties:
        if (t["event_id"] or t["in_menus"] or t["attribution"] != "site_level" or not t["pass"][an.a.theme_filter]
                or len(site_series.get(t["site"], ())) != 1 or org_dropped(an, t["org_key"])):
            continue
        k = t["org_key"]
        if sc in ("C", "D", "C_curated") and not (conf.get(k, 0) >= min_rank
                                                  or s(an.codes.get(k, {}).get("include")) in ("1", "yes")):
            continue
        if sc == "C_curated" and org_class(an, k) in DROP_CLASSES:
            continue
        eid = site_event[t["site"]]
        out.append(dict(t, event_id=eid, series_id=an.series_of[eid], tie_level="L4", tie_sublevel="L4"))
    return out


def _sens_row(an: Analysis, ties: List[dict], f: str, sc: str, variant: str, layers: str, level: str) -> dict:
    G = bipartite(an, ties, level)
    st = graph_stats(G, an.a.backbone_alpha)
    return {"theme_filter": f, "master_theme_filter": an.master_filter or "unknown", "scenario": sc,
            "description": SCENARIOS[sc], "variant": variant, "layers": layers, "node_level": level,
            "rows": len(ties), "events": len({t["event_id"] for t in ties}),
            "series": len({t["series_id"] for t in ties}), "orgs": st["orgs"], "orgs_in_2plus": st["orgs_in_2plus"],
            "edges": st["edges"], "components": st["components"], "giant_nodes": st["giant_nodes"],
            "proj_edges": st["proj_edges"], "top_event_share_of_proj_edges": st["top_event_share_of_proj_edges"],
            "backbone_nodes": st["backbone_nodes"], "backbone_edges": st["backbone_edges"],
            "backbone_a02_nodes": st["backbone02_nodes"], "backbone_a02_edges": st["backbone02_edges"]}


def a5_network(an: Analysis) -> None:
    tf = an.a.theme_filter
    scs = ["A", "B", "C", "D"] + (["C_curated"] if an.curated else [])
    an.default_sc = "C_curated" if an.curated else "C"
    scope = an.a.network_scope
    rows = []
    for f in dict.fromkeys([tf, "none"]):
        for sc in scs:
            for layers in LAYER_SETS:
                ties = scenario_ties(an, sc, layers, f)
                for level in ("edition", "series"):
                    rows.append(_sens_row(an, ties, f, sc, SCOPE_LABEL[scope], layers, level))
    # Variants of the default network (series level, default layers): what the defaults leave out
    dsc, lay = an.default_sc, an.a.layers
    # Rows per tie level and kind of source in the default scenario (all layers), every row and
    # the participation rows: L1 is structurally small (agenda entries name few organisations).
    tl = []
    for sc_scope in ("all", "participation"):
        agg = defaultdict(lambda: [0, set(), set()])
        for x in scenario_ties(an, dsc, "L1-L3", scope=sc_scope):
            a_ = agg[(x["tie_sublevel"], "pdf" if x["layer"] == "linked_pdf" else
                      "listing" if x["tie_level"] == "L1" else "page")]
            a_[0] += 1
            a_[1].add(x["org_key"])
            a_[2].add(x["unit"])
        for (lv, kind), (n, o, u) in sorted(agg.items()):
            tl.append({"scenario": dsc, "rows_in": "every row" if sc_scope == "all" else "participation rows",
                       "tie_sublevel": lv, "kind": kind, "rows": n, "orgs": len(o), "sources": len(u)})
    an.table("tie_levels", tl, ["scenario", "rows_in", "tie_sublevel", "kind", "rows", "orgs", "sources"])
    base = scenario_ties(an, dsc)
    variants = [("+context rows (every row, --network-scope all)", scenario_ties(an, dsc, scope="all")),
                ("reports included", scenario_ties(an, dsc, scope="reports" if scope != "all" else "all")),
                ("cap 50 organisations per source", cap_ties(base, 50)),
                (f"{dsc} without common_word? suggestions", [t for t in base if t["org_key"] not in an.common_words]),
                ("+L4: site-level rows to the site's only series", base + l4_ties(an, dsc))]
    for label, ties in variants:
        rows.append(_sens_row(an, ties, tf, dsc, label, lay, "series"))
    an.l4_added = len(variants[-1][1]) - len(base)
    an.note("L4 sensitivity: site-level rows attributed to the site's only series", an.l4_added,
            "rows added to the default network in the '+L4' sensitivity row (never in the default)")
    an.table("sensitivity", rows, SENS_FIELDS)
    an.sensitivity = rows
    # ---- default scenario: series level, default layers and scope
    ties = base
    an.default_ties = ties
    G = bipartite(an, ties, "series")
    an.G = G
    P = project(G)
    B = disparity(P, an.a.backbone_alpha)
    an.P, an.B = P, B
    top = next((r["top_event_share_of_proj_edges"] for r in rows if r["scenario"] == dsc and r["layers"] == lay
                and r["node_level"] == "series" and r["theme_filter"] == tf
                and r["variant"] == SCOPE_LABEL[scope]), None)
    an.facts["default_graph"] = {"scenario": an.default_sc, "theme_filter": tf, "network_scope": scope, "layers": lay,
                                 "series": sum(1 for n in G if n.startswith("E:")),
                                 "orgs": sum(1 for n in G if n.startswith("O:")),
                                 "orgs_in_2plus": sum(1 for n in G if n.startswith("O:") and G.degree(n) >= 2),
                                 "edges": G.number_of_edges(), "proj_edges": P.number_of_edges(),
                                 "top_event_share_of_proj_edges": top, "backbone_nodes": B.number_of_nodes(),
                                 "backbone_edges": B.number_of_edges()}
    n_orgs = an.facts["default_graph"]["orgs"]
    if n_orgs < 5:
        # e.g. a run of the seed agendas only: most entry names are low-confidence NER finds
        n_a = next((r["orgs"] for r in rows if r["theme_filter"] == tf and r["scenario"] == "A"
                    and r["layers"] == lay and r["node_level"] == "series"), 0)
        an.warn("The default network (scenario %s) has only %s (scenario A: %s). Centrality and "
                "backbone figures are not meaningful; see the sensitivity table or try --min-confidence low",
                an.default_sc, count_of(n_orgs, "organisation"), n_a)
    a6_centrality(an, ties, G, P, B)
    a6b_communities(an, ties, G, B)
    an.exports = {an.default_sc: (G, B)}
    if an.default_sc != "A":
        GA = bipartite(an, scenario_ties(an, "A"), "series")
        an.exports["A"] = (GA, disparity(project(GA), an.a.backbone_alpha))


def a6b_communities(an: Analysis, ties: List[dict], G, B) -> None:
    """Louvain communities of the default backbone (weight, --seed), with their largest
    members, countries and how many members rest on one source only."""
    import networkx as nx
    if not B.number_of_edges():
        an.skipped["communities"] = "the default backbone has no ties"
        an.table("communities", [], ["community", "size"])
        return
    comms = nx.community.louvain_communities(B, weight="weight", seed=an.a.seed)
    comms = sorted((sorted(c) for c in comms), key=lambda c: (-len(c), c[0]))
    mod = round(nx.community.modularity(B, comms, weight="weight"), 4)
    an.facts["communities"] = {"n": len(comms), "n_10plus": sum(1 for c in comms if len(c) >= 10),
                               "modularity": mod, "largest": len(comms[0])}
    an.note("backbone communities (Louvain)", len(comms), f"modularity {mod}, seed {an.a.seed}")
    srcs: Dict[str, Set[str]] = defaultdict(set)
    ev_of: Dict[str, Set[str]] = defaultdict(set)
    for t in ties:
        srcs["O:" + t["org_key"]].add(t["unit"])
        ev_of["O:" + t["org_key"]].add(t["event_id"])
    rows = []
    for i, c in enumerate(comms, 1):
        for n in c:
            B.nodes[n]["community"] = i
            if n in G:
                G.nodes[n]["community"] = i
        top = sorted(c, key=lambda n: (-B.degree(n, weight="weight"), n))[:8]
        cc = Counter(x for n in c for e in sorted(ev_of.get(n, ())) for x in an.ev_by_id[e]["countries"])
        tot = sum(cc.values())
        series = {m for n in c if n in G for m in G[n]}
        names = Counter(G.nodes[m].get("label", m[2:]) for n in c if n in G for m in G[n])
        rows.append({"community": i, "size": len(c),
                     "top_members": "; ".join(B.nodes[n].get("label", n[2:]) for n in top),
                     "dominant_countries": "; ".join(f"{k} {100 * v / tot:.0f} %" for k, v in cc.most_common(3)) if tot
                     else "", "share_one_source": share(sum(1 for n in c if len(srcs.get(n, ())) <= 1), len(c)),
                     "series_covered": len(series),
                     "top_series": "; ".join(short(k, 50) for k, _ in names.most_common(3))})
    an.table("communities", rows, ["community", "size", "top_members", "dominant_countries", "share_one_source",
                                   "series_covered", "top_series"])


def _betweenness(an: Analysis, G) -> Tuple[Dict[str, float], str]:
    import networkx as nx
    kopt = s(an.a.betweenness_k)
    n = G.number_of_nodes()
    if not n:
        return {}, "empty"
    if kopt.isdigit() and int(kopt) < n:
        k = int(kopt)
    elif kopt in ("", "auto") and n > 5000:
        k = 500
    else:
        k = None
    t0 = time.time()
    if k is None:
        bc = nx.betweenness_centrality(G, normalized=True)
        how = "exact"
    else:
        bc = nx.betweenness_centrality(G, k=k, normalized=True, seed=an.a.seed)
        how = f"sampled k={k}, seed={an.a.seed}"
    # the time goes to the log only, so the diagnostics table stays identical between runs
    log.info("Bipartite betweenness (%s) on %d nodes in %.1f s", how, G.number_of_nodes(), time.time() - t0)
    an.note("bipartite betweenness", how, f"{G.number_of_nodes()} nodes")
    return bc, how


def a6_centrality(an: Analysis, ties: List[dict], G, P, B) -> None:
    bc, how = _betweenness(an, G)
    an.betweenness_method = how
    orgs = sorted(n for n in G if n.startswith("O:"))
    # (2,2)-core: organisations in >= 2 series and series with >= 2 core organisations
    core = set(G.nodes)
    changed = True
    while changed:
        changed = False
        for n in list(core):
            if sum(1 for m in G[n] if m in core) < 2:
                core.discard(n)
                changed = True
    by_org = defaultdict(list)
    for t in ties:
        by_org[t["org_key"]].append(t)
    rows = []
    for o in orgs:
        k = o[2:]
        ts = by_org[k]
        series = sorted(G[o])
        members = {e: set(G[e]) - {o} for e in series}
        neigh = set().union(*members.values()) if members else set()
        hd = None
        if neigh:
            big = max(series, key=lambda e: (len(members[e]), e))
            others = set().union(*(members[e] for e in series if e != big)) if len(series) > 1 else set()
            hd = round(len(members[big] - others) / len(neigh), 4)
        evs = {t["event_id"] for t in ts}
        cc = Counter(c for e in sorted(evs) for c in an.ev_by_id[e]["countries"])
        tot = sum(cc.values())
        years = {an.ev_by_id[e]["year"] for e in evs if an.ev_by_id[e]["year"]}
        lv = {lev: len({t["series_id"] for t in ts if t["tie_level"] == lev}) for lev in ("L1", "L2", "L3")}
        lv.update({lev: len({t["series_id"] for t in ts if t["tie_sublevel"] == lev}) for lev in ("L3s", "L3x")})
        rows.append({"org_key": k, "organisation": an.org_name.get(k, k), "actor_class": org_class(an, k),
                     "curated": yn(an.curated), "n_events": len(evs), "n_series": G.degree(o), "n_years": len(years),
                     "n_countries": len(cc), "countries": "; ".join(c for c, _ in cc.most_common()),
                     "n_sites": len({t["site"] for t in ts}), "n_sources": len({t["unit"] for t in ts}),
                     "n_documents": len({t["document"] for t in ts}),
                     "n_L1": lv["L1"], "n_L2": lv["L2"], "n_L3": lv["L3"], "n_L3s": lv["L3s"], "n_L3x": lv["L3x"],
                     "bip_degree": G.degree(o), "bip_betweenness": round(bc.get(o, 0.0), 6),
                     "backbone_degree": B.degree(o) if o in B else 0,
                     "backbone_strength": round(B.degree(o, weight="weight"), 4) if o in B else 0.0,
                     "in_backbone": o in B, "projection_degree": P.degree(o) if o in P else 0,
                     "hub_dependence": hd, "hub_dependent": hd is not None and hd >= 0.8,
                     "core_2_2": o in core,
                     "participation_coef_country": round(1 - sum((v / tot) ** 2 for v in cc.values()), 4) if tot else None})
    rows.sort(key=lambda r: (-r["n_series"], -r["bip_betweenness"], r["org_key"]))
    an.org_centrality = rows
    an.table("org_centrality", rows, ["org_key", "organisation", "actor_class", "curated", "n_events", "n_series",
                                      "n_years", "n_countries", "countries", "n_sites", "n_sources", "n_documents",
                                      "n_L1", "n_L2", "n_L3", "n_L3s", "n_L3x",
                                      "bip_degree", "bip_betweenness", "backbone_degree", "backbone_strength",
                                      "in_backbone", "projection_degree", "hub_dependence", "hub_dependent",
                                      "core_2_2", "participation_coef_country"])


# ==========================================================================
# A7 theme x organisation association and bridging
# ==========================================================================
_LINKERS = re.compile(r"(?i)\b(?:gegen|against|contre|contro|contra|tegen|auf|on|sur|for|für|pour|per|para|"
                      r"voor|without|ohne|sans|sin|senza|zonder|of|von|de|des|du|del|della|van|kein|keine|no|"
                      r"stop|stopp|to|up|à|al|zu|zum|zur|y|and|und|et|e|en|og|och)\b")


def theme_named_themes(an: Analysis, name: str) -> Set[str]:
    """The focus themes whose term the name is ("Rassismus", "Nazis", "Le racisme") or is
    headed by in a name of two words ("SOS Racisme", "Coordination Antifasciste"). Such
    a name co-occurs with that theme by construction, so the (organisation, theme) pair
    is left out of org_theme; its other themes are kept. Names of three or more words
    ("Jeune Garde Antifasciste") and names with a linker before the term ("Bündnis gegen
    Rassismus", "Stand Up To Racism") are kept for every theme."""
    if not an.ts or not name:
        return set()
    text = collapse(name)
    found = {th: hs for th, hs in an.ts.find(text).items() if th in an.themes and hs}
    if not found:
        return set()
    words = _ART.sub("", text.casefold()).split()
    if len(words) == 1:
        return set(found)
    if len(words) >= 3:
        return set()
    out = set()
    for th, hs in found.items():
        for a, b, _ in hs:
            if b >= len(text.rstrip(" .")) - 1 and not _LINKERS.search(text[:a]):
                out.add(th)
    return out


def theme_named(an: Analysis, name: str) -> bool:
    """Is the name named after one of the focus themes (see theme_named_themes)?"""
    return bool(theme_named_themes(an, name))


def a7_org_theme(an: Analysis) -> None:
    th = an.themes
    # passage level (primary): organisations named in the paragraph of a clean theme passage
    pas = defaultdict(lambda: {"passages": set(), "docs": set(), "sites": set(), "series": set(), "example": ""})
    named = {}
    for inst in an.instances:
        if not inst["clean"]:
            continue
        doc = inst["refs"][0]["document"] if inst["refs"] else inst["source_link"]
        series = {an.series_of[rf["event_id"]] for rf in inst["refs"] if rf["event_id"]}
        for k in inst["org_keys"]:
            if org_dropped(an, k):
                continue
            if k not in named:
                named[k] = theme_named_themes(an, an.org_name.get(k, k))
            if inst["theme"] in named[k]:
                continue
            a = pas[(k, inst["theme"])]
            a["passages"].add(inst["passage_id"])
            a["sites"].add(inst["site"])
            a["docs"].add(doc)
            a["series"].update(series)
            a["example"] = a["example"] or inst["passage_id"]
    # event level (secondary): series of the default scenario whose L1/L2 text carries the theme.
    # Measured on all sources (theme filter "none") for the organisations the filter keeps,
    # so the inclusion criterion does not inflate the shares it is compared with.
    included = {t["org_key"] for t in an.default_ties}
    base = [t for t in scenario_ties(an, an.default_sc, None, "none") if t["org_key"] in included]
    # When make_master.py already applied a theme filter, the 'none' ties above are the
    # filtered ones: every series then has a theme hit by construction, which inflates
    # share, lift and shrunk. Those three are left blank in that case.
    prefiltered = an.master_filter != "none"
    if prefiltered:
        an.skipped["org_theme event-level share, lift, shrunk"] = (
            f"the masters were built with theme filter '{an.master_filter or 'unknown'}', so the unfiltered base "
            "rate cannot be measured; rebuild them with make_master.py --theme-filter none and run analytics.py "
            "with --theme-filter source")
    series_theme = defaultdict(set)
    for ev in an.events:
        for t in th:
            if ev[f"{t}_own"] or ev[f"{t}_landing"]:
                series_theme[ev["series_id"]].add(t)
    org_series = defaultdict(set)
    org_title_series = defaultdict(set)
    for t in base:
        org_series[t["org_key"]].add(t["series_id"])
        if t["named_in_event_title"]:
            org_title_series[t["org_key"]].add(t["series_id"])
    all_series = set().union(*org_series.values()) if org_series else set()
    p0 = {t: share(sum(1 for x in all_series if t in series_theme.get(x, ())), len(all_series)) or 0.0 for t in th}
    rows = []
    keys = sorted(set(k for k, _ in pas) | set(org_series))
    excluded = sorted(k for k in keys if named.setdefault(k, theme_named_themes(an, an.org_name.get(k, k))))
    if excluded:
        an.note("organisations named after a theme (left out of org_theme for that theme only)", len(excluded),
                "; ".join(f"{an.org_name.get(k, k)} ({', '.join(sorted(named[k]))})" for k in excluded[:12]))
    for k in keys:
        for t in th:
            if t in named[k]:
                continue   # named after this theme: co-occurs with it by construction
            a = pas.get((k, t))
            ser = org_series.get(k, set())
            n = len(ser)
            kk = sum(1 for x in ser if t in series_theme.get(x, ()))
            ev_ok = n >= 3 and kk >= 2
            if not a and not ev_ok:
                continue
            r = {"org_key": k, "organisation": an.org_name.get(k, k), "actor_class": org_class(an, k), "theme": t,
                 "passages": len(a["passages"]) if a else 0, "documents": len(a["docs"]) if a else 0,
                 "series_passage": len(a["series"]) if a else 0, "n_series": n, "k_series": kk,
                 "k_title": sum(1 for x in org_title_series.get(k, ()) if t in series_theme.get(x, ())),
                 "example_passage_id": a["example"] if a else ""}
            if ev_ok and not prefiltered:
                lo, _ = wilson(kk, n)
                r.update({"share": share(kk, n), "wilson_lo": lo,
                          "lift": round((kk / n) / p0[t], 3) if p0[t] else None,
                          "shrunk": round((kk + 5 * p0[t]) / (n + 5), 4)})
            # robust: two documents from two websites, or the event-level rule (only when the
            # masters were not theme-filtered: otherwise every series has a hit by construction)
            r["sites"] = len(a["sites"]) if a else 0
            r["robust"] = (r["documents"] >= 2 and r["sites"] >= 2) or (ev_ok and kk >= 2 and not prefiltered)
            rows.append(r)
    an.table("org_theme", rows, ["org_key", "organisation", "actor_class", "theme", "passages", "documents",
                                 "series_passage", "n_series", "k_series", "k_title", "share", "wilson_lo", "lift",
                                 "shrunk", "robust", "example_passage_id", "sites"])
    an.org_theme = rows
    an.facts["org_theme"] = {"orgs_in_passages": len({k for k, _ in pas}),
                             "robust_pairs_passage": sum(1 for (k, t), a in pas.items()
                                                         if len(a["docs"]) >= 2 and len(a["sites"]) >= 2),
                             "p0": None if prefiltered else p0}
    cen = {r["org_key"]: r for r in getattr(an, "org_centrality", [])}
    robust = defaultdict(list)
    docs_of = defaultdict(set)
    for r in rows:
        if r["robust"]:
            robust[r["org_key"]].append(r["theme"])
    for (k, _), a in pas.items():
        docs_of[k] |= a["docs"]
    nser = {k: len(v) for k, v in org_series.items()}
    br = []
    for k, ts_ in robust.items():
        if len(ts_) < 2:
            continue        # bridging: robustly tied to two or more themes
        c = cen.get(k, {})
        br.append({"org_key": k, "organisation": an.org_name.get(k, k), "actor_class": org_class(an, k),
                   "n_themes_robust": len(ts_), "themes": "; ".join(ts_), "bridging": len(ts_) >= 2,
                   "n_series": nser.get(k, 0), "documents": len(docs_of.get(k, ())),
                   "bip_betweenness": c.get("bip_betweenness"),
                   "participation_coef_country": c.get("participation_coef_country")})
    br.sort(key=lambda r: (-r["n_series"], -r["documents"], -r["n_themes_robust"], r["org_key"]))
    an.table("org_bridging", br, ["org_key", "organisation", "actor_class", "n_themes_robust", "themes", "bridging",
                                  "n_series", "documents", "bip_betweenness", "participation_coef_country"])
    an.facts["bridging_orgs"] = [r["organisation"] for r in br if r["bridging"]]
    # ---- theme co-occurrence at three levels
    sets = {"passage": defaultdict(set), "document": defaultdict(set), "series": defaultdict(set)}
    site_of = {"passage": {}, "document": {}, "series": defaultdict(set)}
    for inst in an.instances:
        if not inst["clean"]:
            continue
        t = inst["theme"]
        doc = inst["refs"][0]["document"] if inst["refs"] else inst["source_link"]
        sets["passage"][t].add(inst["passage_id"])
        site_of["passage"][inst["passage_id"]] = inst["site"]
        sets["document"][t].add(doc)
        site_of["document"][doc] = inst["site"]
        for rf in inst["refs"]:
            if rf["event_id"]:
                x = an.series_of[rf["event_id"]]
                sets["series"][t].add(x)
                site_of["series"][x].add(inst["site"])
    n_series_all = len({e["series_id"] for e in an.events})
    rows = []
    for lev, d in sets.items():
        for i, a in enumerate(th):
            for b in th[i + 1:]:
                A, Bs = d.get(a, set()), d.get(b, set())
                ab = A & Bs
                uni = A | Bs
                sites = sorted({x for y in uni for x in ((site_of[lev][y] if lev == "series" else {site_of[lev][y]}))})
                loo = []
                for st in sites:
                    def ok(y):
                        v = site_of[lev][y]
                        return st not in (v if lev == "series" else {v})
                    loo.append(sum(1 for y in ab if ok(y)))
                r = {"level": lev, "theme_a": a, "theme_b": b, "n_a": len(A), "n_b": len(Bs), "n_ab": len(ab),
                     "jaccard": share(len(ab), len(uni)), "p_b_given_a": share(len(ab), len(A)),
                     "lift": round(len(ab) * n_series_all / (len(A) * len(Bs)), 3)
                     if lev == "series" and A and Bs else None,
                     "loo_site_min_n_ab": min(loo) if loo else None, "loo_site_max_n_ab": max(loo) if loo else None}
                rows.append(r)
    an.table("theme_cooccurrence", rows, ["level", "theme_a", "theme_b", "n_a", "n_b", "n_ab", "jaccard",
                                          "p_b_given_a", "lift", "loo_site_min_n_ab", "loo_site_max_n_ab"])


# ==========================================================================
# A1 organisation authority list, merge candidates, precision sample
# ==========================================================================
_COMPANY = re.compile(r"\b(?:AG|SE|PLC|GmbH|Ltd|B\.?V\.?|S\.?A\.?|N\.?V\.?|Inc|Corp|Bank|Group)\.?$", re.I)
FOSSIL_ARMS = {"shell", "rwe", "uniper", "eph", "totalenergies", "total", "equinor", "vattenfall", "leag", "bp", "eni",
               "rheinmetall", "leonardo", "thales", "exxon", "exxonmobil", "chevron", "engie", "repsol", "omv",
               "wintershall", "bayer", "tesla", "glencore", "heidelberg materials", "lafarge", "holcim"}
STATE = re.compile(r"^(?:nato|eu|european union|european commission|commission|un|united nations|unfccc|cop\d*|"
                   r"bundesregierung|government|gouvernement|regierung)$|minist|bundesamt|police|polizei|"
                   r"gendarmerie|tribunal|court|gericht|parliament|parlement|bundestag", re.I)
SITE_UI = re.compile(r"^(?:faq|newsletter|archive|archives|archiv|calendar|kalender|agenda|outlook|login|log in|"
                     r"impressum|imprint|datenschutz|privacy|cookie|cookies|contact|kontakt|menu|home|search|"
                     r"donate|spenden|shop|blog|news|press|presse|about|about us|team|events?|tech)$", re.I)
# Interface words that mark a name as site UI wherever they occur ('Outlook 365',
# 'Add to Google Calendar'); words that are also common in organisation names
# (archive, news, tech) only count as the whole name (SITE_UI above).
SITE_UI_ANY = re.compile(r"\b(?:calendar|kalender|calendrier|outlook|ical|icalendar|add to|login|log in|sign in|"
                         r"newsletter|faq|cookies?|impressum|datenschutz|privacy policy|google maps)\b", re.I)


# Methods that are as weak as NER on their own (extract.NER_LIKE_METHODS).
# Project() replaces it with extract.NER_LIKE_METHODS when that exists.
_NER_LIKE = frozenset({"ner", "pattern_after", "link_text", "website_elsewhere"})


def _route(found_by: Set[str]) -> str:
    return "ner_only" if found_by and found_by <= _NER_LIKE else ("rule" if found_by else "unknown")


_WEEKDAYS_FALLBACK = frozenset("""monday tuesday wednesday thursday friday saturday sunday montag dienstag mittwoch
donnerstag freitag samstag sonnabend sonntag lundi mardi mercredi jeudi vendredi samedi dimanche lunes martes
miércoles jueves viernes sábado domingo lunedì martedì mercoledì giovedì venerdì sabato domenica maandag dinsdag
woensdag donderdag vrijdag zaterdag zondag segunda terça quarta quinta sexta mandag tirsdag onsdag torsdag fredag
lørdag søndag""".split())


def common_word_suggestions(an: Analysis, max_contexts: int = 6) -> Set[str]:
    """Organisation keys to pre-fill as 'common_word?' in the curation template: weekdays
    (8 languages), and one-word title-case names that are not acronyms when the
    language's part-of-speech tagger reads them as an ordinary word (noun, adjective,
    verb ...) in at least half of up to six of their example contexts
    (extract.common_word_votes). Institutional compounds ("Bundesregierung"),
    hyphenated organisation compounds ("Rosa-Luxemburg-Stiftung") and organisational
    head nouns are exempt. Only suggestions: the default network keeps these names."""
    t0 = time.time()
    ex = an.p.extract
    weekdays = frozenset(getattr(ex, "WEEKDAYS", None) or _WEEKDAYS_FALLBACK)
    single = getattr(ex, "single_word_name", None)
    exempt = getattr(ex, "_pos_exempt", None)
    heads = set(getattr(ex, "HEAD_TO_TYPE", {}) or {})
    votes_fn = getattr(ex, "common_word_votes", None)
    tf = an.a.theme_filter
    found: Set[str] = set()
    ctx: Dict[str, List[Tuple[str, str]]] = defaultdict(list)
    for t in an.ties:
        if not t["pass"][tf]:
            continue
        k, name = t["org_key"], an.org_name.get(t["org_key"], t["organisation"])
        if k in weekdays or name.casefold() in weekdays:
            found.add(k)
            continue
        if " " in k.strip() or k in heads:
            continue
        ok = single(name) if single else bool(re.fullmatch(r"[A-ZÄÖÜÉ][a-zäöüßéèàç\-]+", name))
        if not ok or (exempt and (exempt(name) or exempt(k))):
            continue
        ex_ = collapse(t["example"])
        lst = ctx[k]
        if ex_ and len(lst) < max_contexts and t["organisation"] in ex_ and all(c != ex_ for _, c in lst):
            lst.append((t["organisation"], ex_))
    an.facts["common_word_candidates"] = len(ctx)
    if votes_fn is None:
        an.skipped["common_word? POS votes"] = "extract.common_word_votes not available (weekdays only)"
        return found
    detect = an.p.detect_language or (lambda x: "")
    items = []
    for k, lst in sorted(ctx.items()):
        for surface, c in lst:
            try:
                lang = detect(c)
            except Exception:  # noqa: BLE001
                lang = ""
            if lang:
                items.append((k, surface, c, lang))
    try:
        votes = votes_fn(items, max_per_key=max_contexts)
    except Exception as exc:  # noqa: BLE001 - no tagger: weekdays only
        an.warn("common_word? suggestions: POS votes failed (%s: %s); weekdays only", type(exc).__name__, exc)
        return found
    found |= {k for k, (n, c) in votes.items() if n >= 1 and c >= 0.5 * n}
    log.info("common_word? suggestions: %d names (%d candidates tagged) in %.1f s", len(found), len(ctx),
             time.time() - t0)
    an.note("common_word? suggestions", len(found), f"{len(ctx)} one-word names tagged; weekdays always")
    return found


def a1_authority(an: Analysis, cap_rows: Optional[int] = None) -> None:
    p = an.p
    tf = an.a.theme_filter
    ties = [t for t in an.ties if t["pass"][tf]]
    by = defaultdict(list)
    for t in ties:
        by[t["org_key"]].append(t)
    para = defaultdict(lambda: [set(), set()])
    for inst in an.instances:
        if inst["clean"]:
            for k in inst["org_keys"]:
                para[k][0].add(inst["passage_id"])
                para[k][1].add(inst["theme"])
    inv = {v: k for k, v in CONF_RANK.items()}
    rows = []
    for k in sorted(by):
        ts = by[k]
        real = [t for t in ts if t["event_id"]]
        evs = {t["event_id"] for t in real}
        name = an.org_name.get(k, k)
        fb = set()
        for t in ts:
            fb.update(x for x in re.split(r"[;+,\s]+", t["found_by"].lower()) if x)
        years = sorted({an.ev_by_id[e]["year"] for e in evs if an.ev_by_id[e]["year"]})
        cc = Counter(c for e in sorted(evs) for c in an.ev_by_id[e]["countries"])
        lv = Counter(t["tie_level"] for t in real)
        ex = sorted(ts, key=lambda t: (t["in_menus"], -len(t["example"] or "")))
        exs = list(dict.fromkeys(collapse(t["example"]) for t in ex if t["example"]))[:2]
        surf = an.org_surfaces.get(k, Counter())
        flags = []
        if " " not in k and fb <= _NER_LIKE:
            flags.append("single_word_ner_only")
        if re.fullmatch(r"[A-Z0-9&.-]{2,6}", name):
            flags.append("unexpanded_acronym")
        if real and all(t["reference"] for t in real):
            flags.append("only_in_reference_docs")
        sug = ""
        if SITE_UI.match(k) or SITE_UI_ANY.search(k) or SITE_UI_ANY.search(name):
            sug = "site_ui?"
        elif k in an.common_words:
            sug = "common_word?"
        elif k in p.places or _is_place_name(p, name):
            sug = "place?"
        elif _COMPANY.search(name) or k in FOSSIL_ARMS:
            sug = "company?"
        elif STATE.search(k):
            sug = "state?"
        elif evs and all(k in p.key(an.ev_by_id[e]["event"]) for e in evs):
            sug = "event_name?"
        c = an.codes.get(k, {})
        rows.append({"org_key": k, "organisation": name,
                     "variants": "; ".join(v for v, _ in surf.most_common() if v != name)[:500],
                     "max_confidence": inv[max(CONF_RANK.get(t["confidence"], 0) for t in ts)],
                     "found_by": "; ".join(sorted(fb)), "route": _route(fb),
                     "n_events": len(evs), "n_series": len({an.series_of[e] for e in evs}),
                     "n_years": len(years), "years": "; ".join(map(str, years)),
                     "countries": "; ".join(x for x, _ in cc.most_common()),
                     "n_sources": len({t["unit"] for t in ts}), "n_sites": len({t["site"] for t in ts}),
                     "sites": "; ".join(x for x, _ in Counter(t["site"] for t in ts).most_common(5)),
                     "tie_levels": "; ".join(f"{x}:{lv[x]}" for x in ("L1", "L2", "L3") if lv[x]),
                     "only_in_event_title": bool(real) and all(t["named_in_event_title"] for t in real),
                     "only_menus_footers": all(t["in_menus"] for t in ts),
                     "theme_passages": len(para[k][0]), "themes_in_paragraph": "; ".join(sorted(para[k][1])),
                     "example_context_1": exs[0] if exs else "", "example_context_2": exs[1] if len(exs) > 1 else "",
                     "suggested_class": sug, "flags": "; ".join(flags), "variant_group": "",
                     "merge_into_suggested": "",
                     **{f: c.get(f, "") for f in ("include", "merge_into_key", "canonical_name", "actor_class",
                                                  "country", "is_organisation", "is_partial_name", "coder", "notes")}})
    cands = merge_candidates(an, rows)
    best: Dict[str, Tuple[float, str]] = {}
    for c in cands:
        for a_, b_ in ((c["key_a"], c["key_b"]), (c["key_b"], c["key_a"])):
            if c["score"] > best.get(a_, (0, ""))[0]:
                best[a_] = (c["score"], b_)
    n_of = {r["org_key"]: r["n_events"] for r in rows}
    for r in rows:
        if r["org_key"] in best:
            other = best[r["org_key"]][1]
            r["variant_group"] = min(r["org_key"], other)
            # suggest merging the less frequent key into the more frequent one
            if (n_of.get(other, 0), other) > (r["n_events"], r["org_key"]):
                r["merge_into_suggested"] = other
    # Curation priority: 1 in two or more series of the default network, 2 in one series
    # there through the listing or the landing page, 3 the rest, 4 no series at all.
    G = getattr(an, "G", None)
    l12 = {t["org_key"] for t in getattr(an, "default_ties", []) if t["tie_level"] in ("L1", "L2")}
    for r in rows:
        node = "O:" + r["org_key"]
        deg = G.degree(node) if G is not None and node in G else 0
        r["priority"] = 1 if deg >= 2 else 2 if deg == 1 and r["org_key"] in l12 else 3 if r["n_series"] else 4
        r["in_default_network"] = deg > 0
    rows.sort(key=lambda r: (r["priority"], -r["n_series"], -r["n_events"], r["org_key"]))
    an.authority = rows
    an.facts["curation_priority"] = dict(sorted(Counter(r["priority"] for r in rows).items()))
    an.sources_to_code = sources_to_code(an)
    an.merge_cands = cands
    an.precision_sample = precision_sample(an, rows)
    an.precision_rows = precision_readback(an, rows)


SOURCE_FIELDS = ["source_unit", "title", "layer", "host_relation", "n_events", "n_orgs", "doc_kind",
                 "doc_kind_reason", "suggested_class", "include", "notes"]
HOST_RELATION = {"L1": "listing", "L2": "landing page", "L3s": "event's own host", "L3x": "other host"}


def sources_to_code(an: Analysis) -> List[dict]:
    """One row per source (page, PDF; not agenda entries or calendar cells) with its
    organisations: what the analysis took it for (doc_kind) and an include column to
    override that (1 = evidence of participation, 0 = leave out; --source-codes)."""
    tf = an.a.theme_filter
    agg: "OrderedDict[str, dict]" = OrderedDict()
    for t in an.ties:
        if not t["pass"][tf] or t["in_menus"] or t["layer"] in ("listing_seed", "listing_calendar"):
            continue
        a = agg.setdefault(t["unit"], {"t": t, "orgs": set(), "events": set(), "sub": Counter()})
        a["orgs"].add(t["org_key"])
        if t["event_id"]:
            a["events"].add(t["event_id"])
        a["sub"][t["tie_sublevel"]] += 1
    rows = []
    for u, a in agg.items():
        t = a["t"]
        sub = a["sub"].most_common(1)[0][0]
        over = an.source_override.get(u)
        kind = t["doc_kind"]
        sug = ("reference?" if kind == "report" else "programme?" if kind == "programme" else
               "list?" if len(a["orgs"]) > 100 else "")
        rows.append({"source_unit": u, "title": short(t["source_title"], 120), "layer": t["layer"],
                     "host_relation": HOST_RELATION.get(sub, sub) if a["events"] else "no event",
                     "n_events": len(a["events"]), "n_orgs": len(a["orgs"]), "doc_kind": kind,
                     "doc_kind_reason": t["doc_kind_reason"], "suggested_class": sug,
                     "include": "" if over is None else ("1" if over else "0"), "notes": ""})
    rows.sort(key=lambda r: (-r["n_orgs"], r["source_unit"]))
    return rows


def _is_place_name(p: Project, name: str) -> bool:
    # Place names are short; longer names are not looked up (the gazetteer is the slow part).
    if p.gazetteer is None or len(name) < 3 or len(name.split()) > 3:
        return False
    try:
        res = p.country(name)
    except Exception:  # noqa: BLE001
        return False
    m = res.get("matched", "")
    m = m[0] if isinstance(m, (list, tuple)) and m else m
    return res.get("method") in ("name", "city") and collapse(s(m)).casefold() == collapse(name).casefold()


_ART = re.compile(r"^(?:the|la|le|les|l'|d'|de|het|die|der|das|el|il|lo)\s+")


def merge_candidates(an: Analysis, rows: List[dict]) -> List[dict]:
    """Pairs of organisation keys that may be the same organisation. Compared only
    inside blocks (same first word, or a shared 3-letter piece of it), so the
    search stays fast for thousands of names."""
    keys = [r["org_key"] for r in rows]
    name = {r["org_key"]: r["organisation"] for r in rows}
    blocks = defaultdict(set)
    for k in keys:
        first = (_ART.sub("", k).split() or [""])[0]
        blocks["w:" + first].add(k)
        for i in range(max(1, len(first) - 2)):
            blocks["g:" + first[i:i + 3]].add(k)
    pairs = set()
    for b, ks in blocks.items():
        if len(ks) < 2 or len(ks) > 300:
            continue
        ks = sorted(ks)
        for i, a in enumerate(ks):
            for c in ks[i + 1:]:
                pairs.add((a, c))
    # acronyms: initials of a multi-word name equal to a one-word key
    single = {k for k in keys if " " not in k and 2 <= len(k) <= 6}
    for k in keys:
        toks = [t for t in k.split() if t not in ("of", "the", "de", "la", "le", "and", "et", "und", "for", "du", "des")]
        if len(toks) >= 2:
            ini = "".join(t[0] for t in toks)
            if ini in single and ini != k:
                pairs.add(tuple(sorted((k, ini))))

    def norm(k: str, how: str) -> str:
        if how == "umlaut":
            return _fold_umlaut(k)
        if how == "elision":
            return re.sub(r"\b[dl]'\s*", "", k)
        if how == "article":
            return _ART.sub("", k)
        if how == "year":
            return re.sub(r"\s*(?:19|20)\d\d$", "", k)
        if how == "inclusive":
            return re.sub(r"[.·•-]?(?:e[.·•-]?s|ne[.·•-]?s|e)\b", "", k)
        return k
    hows = ("umlaut", "elision", "article", "year", "inclusive")
    normed = {k: tuple(norm(k, h) for h in hows) for k in keys}   # once per key, not per pair
    out = []
    for a, b in sorted(pairs):
        reason = ""
        for i, how in enumerate(hows):
            if normed[a][i] and normed[a][i] == normed[b][i]:
                reason = how
                break
        # the cheap upper bounds first (length, then letters): the full ratio is computed
        # only when it can reach 0.88
        sm = None
        if reason or 2 * min(len(a), len(b)) / (len(a) + len(b)) >= 0.88:
            sm = difflib.SequenceMatcher(None, a, b)
        ratio = sm.ratio() if sm and (reason or sm.quick_ratio() >= 0.88) else 0.0
        ta, tb = set(a.split()), set(b.split())
        if not reason:
            if ratio >= 0.88:
                reason = "similar spelling"
            elif (" " in a) != (" " in b) and "".join(t[0] for t in (a if " " in a else b).split()) == (
                    b if " " in a else a):
                reason = "acronym"
            elif ta and tb and (ta < tb or tb < ta) and max(len(x) for x in ta & tb or {""}) >= 4:
                reason = "contained"
        if reason:
            ratio = ratio or difflib.SequenceMatcher(None, a, b).ratio()
            out.append({"key_a": a, "name_a": name[a], "key_b": b, "name_b": name[b], "reason": reason,
                        "score": round(ratio, 3), "accept": "yes" if reason == "umlaut" else ""})
    out.sort(key=lambda r: (r["reason"] != "umlaut", -r["score"], r["key_a"], r["key_b"]))
    return out


def precision_sample(an: Analysis, rows: List[dict]) -> List[dict]:
    rng = random.Random(an.a.seed)
    strata = defaultdict(list)
    for r in rows:
        strata[f"{r['max_confidence']} x {r['route']}"].append(r)
    out = []
    for st in sorted(strata):
        pool = sorted(strata[st], key=lambda r: r["org_key"])
        n = min(len(pool), an.a.precision_per_stratum)
        pick = rng.sample(pool, n) if n < len(pool) else pool
        for r in sorted(pick, key=lambda r: r["org_key"]):
            out.append({"stratum": st, "stratum_N": len(pool), "n": n, "weight": round(len(pool) / n, 4) if n else None,
                        "org_key": r["org_key"], "organisation": r["organisation"], "found_by": r["found_by"],
                        "n_events": r["n_events"], "example_context_1": r["example_context_1"],
                        "is_organisation": r.get("is_organisation", ""), "is_partial_name": r.get("is_partial_name", ""),
                        "coder": r.get("coder", "")})
    return out


def _stratified_interval(p: float, var: float, n: int) -> Tuple[Optional[float], Optional[float]]:
    """95 % interval for a stratified proportion: Wilson with the effective sample size
    p(1-p)/var (Korn and Graubard), capped at the number of coded names."""
    if not n:
        return None, None
    n_eff = min(n, p * (1 - p) / var) if var > 0 else n
    return wilson(p * n_eff, n_eff)


def precision_readback(an: Analysis, rows: List[dict]) -> List[dict]:
    """Precision per stratum from coded org_codes (is_organisation filled in)."""
    coded = {k: c for k, c in an.codes.items() if s(c.get("is_organisation"))}
    if not coded:
        an.skipped["precision_by_stratum"] = "no is_organisation codes in --org-codes"
        return []
    strata = defaultdict(lambda: {"N": 0, "n": 0, "k": 0, "kl": 0})
    for r in rows:
        st = f"{r['max_confidence']} x {r['route']}"
        a = strata[st]
        a["N"] += 1
        c = coded.get(r["org_key"])
        if c:
            a["n"] += 1
            yes = truthy(c.get("is_organisation"))
            a["k"] += yes
            a["kl"] += yes or truthy(c.get("is_partial_name")) or s(c.get("is_organisation")).lower() == "partial"
    out = []
    totN = sum(a["N"] for a in strata.values() if a["n"])
    est = estl = var = varl = 0.0
    for st in sorted(strata):
        a = strata[st]
        lo, hi = wilson(a["k"], a["n"])
        out.append({"stratum": st, "N": a["N"], "n": a["n"], "k": a["k"], "precision": share(a["k"], a["n"]),
                    "wilson_lo": lo, "wilson_hi": hi, "k_lenient": a["kl"], "precision_lenient": share(a["kl"], a["n"])})
        if a["n"] and totN:
            w, n_h, fpc = a["N"] / totN, a["n"], 1 - a["n"] / a["N"]
            ph, phl = a["k"] / n_h, a["kl"] / n_h
            est += w * ph
            estl += w * phl
            if n_h > 1:     # unbiased within-stratum variance, finite population correction
                var += w * w * fpc * ph * (1 - ph) / (n_h - 1)
                varl += w * w * fpc * phl * (1 - phl) / (n_h - 1)
    n_all = sum(a["n"] for a in strata.values())
    out.append({"stratum": "ALL (stratified estimate over strata with coded names; N = their size; interval: "
                           "Wilson with the effective sample size)", "N": totN, "n": n_all,
                "k": sum(a["k"] for a in strata.values()), "precision": round(est, 4),
                "k_lenient": sum(a["kl"] for a in strata.values()), "precision_lenient": round(estl, 4)})
    lo, hi = _stratified_interval(est, var, n_all)
    out[-1].update({"wilson_lo": lo, "wilson_hi": hi, "se": round(math.sqrt(var), 4)})
    lo, hi = _stratified_interval(estl, varl, n_all)
    out[-1].update({"lenient_lo": lo, "lenient_hi": hi})
    an.facts["precision"] = out[-1]
    return out


# ==========================================================================
# A8 coding workbook (keyword in context)
# ==========================================================================
def _kwic(inst: dict, words: int = 14) -> Tuple[str, str, str]:
    sent, term = inst["sentence"], inst["term"]
    i = sent.casefold().find(term.casefold())
    if i < 0:
        left, mid, right = sent, term, ""
    else:
        left, mid, right = sent[:i], sent[i:i + len(term)], sent[i + len(term):]
    lw = (collapse(inst["context_before"]) + " " + collapse(left)).split()
    rw = (collapse(right) + " " + collapse(inst["context_after"])).split()
    return " ".join(lw[-words:]), mid, " ".join(rw[:words])


def a8_kwic(an: Analysis) -> dict:
    rng = random.Random(an.a.seed)
    units = OrderedDict()
    for inst in an.instances:
        if not inst["clean"]:
            continue
        key = (inst["passage_id"], inst["theme"])
        if key in units:
            continue
        refs = inst["refs"]
        ev = next((an.ev_by_id[rf["event_id"]] for rf in refs if rf["event_id"]), None)
        lev = min((rf["tie_level"] for rf in refs), default="L3")
        lw, mid, rw = _kwic(inst)
        units[key] = {"passage_id": inst["passage_id"], "theme": inst["theme"], "subtheme": inst["subtheme"],
                      "stance": inst["stance"], "term_family": inst["term_family"], "left_ctx": lw, "term": mid,
                      "right_ctx": rw, "context_before": inst["context_before"], "sentence": inst["sentence"],
                      "context_after": inst["context_after"],
                      "event": "; ".join(dict.fromkeys(inst["event_names"]))[:1000],
                      "event_level": any(rf["event_id"] for rf in refs),
                      "event_year": ev["year"] if ev else "", "doc_year": inst["doc_year"] or "undated",
                      "countries": "; ".join(ev["countries"]) if ev else "",
                      "stream": ev["stream_primary"] if ev else "", "layer": refs[0]["layer"] if refs else "",
                      "layer_group": LAYER_GROUP.get(lev, "linked"), "period": ev["period"] if ev else "undated",
                      "site": inst["site"], "source_link": inst["source_link"],
                      "document": refs[0]["document"] if refs else inst["source_link"],
                      "orgs": inst["orgs_text"], "in_event_title": inst["in_event_title"],
                      "in_handle": inst["in_handle"]}
    allu = list(units.values())
    # per-document cap for the per-theme sheets
    capped = []
    for t in an.themes:
        by_doc = defaultdict(list)
        for u in allu:
            if u["theme"] == t:
                by_doc[u["document"]].append(u)
        for d in sorted(by_doc):
            pool = sorted(by_doc[d], key=lambda u: u["passage_id"])
            if len(pool) > an.a.doc_cap:
                pool = sorted(rng.sample(pool, an.a.doc_cap), key=lambda u: u["passage_id"])
            capped.extend(pool)
    # stratified sample: theme x layer group x period, proportional, at least 2 per stratum
    strata = defaultdict(list)
    for u in capped:
        strata[(u["theme"], u["layer_group"], u["period"])].append(u)
    N = len(capped)
    size = min(an.a.sample_size, N)
    alloc = {}
    for st, pool in strata.items():
        alloc[st] = min(len(pool), max(2, round(size * len(pool) / N))) if N else 0
    while sum(alloc.values()) > size and any(v > 2 for v in alloc.values()):
        st = max((x for x in alloc if alloc[x] > 2), key=lambda x: (alloc[x] - size * len(strata[x]) / N, x))
        alloc[st] -= 1
    frame = []
    for st in sorted(strata):
        pool = sorted(strata[st], key=lambda u: u["passage_id"])
        n = alloc[st]
        pick = {u["passage_id"] for u in (rng.sample(pool, n) if n < len(pool) else pool)}
        w = round(len(pool) / n, 4) if n else None
        for u in pool:
            u["in_sample"] = u["passage_id"] in pick
            u["sample_weight"] = w if u["in_sample"] else ""
            u["stratum"] = " x ".join(st)
        frame.append({"stratum": " x ".join(st), "theme": st[0], "layer_group": st[1], "period": st[2],
                      "N_h": len(pool), "n_h": n, "weight": w})
    sampled = sorted([u for u in capped if u.get("in_sample")], key=lambda u: u["passage_id"])
    dbl = set(u["passage_id"] + u["theme"] for u in rng.sample(sampled, round(0.25 * len(sampled)))) if sampled else set()
    for u in capped:
        u["double_coded"] = (u["passage_id"] + u["theme"]) in dbl
    for u in allu:
        u.setdefault("in_sample", False)
        u.setdefault("double_coded", False)
    an.facts["kwic"] = {"units": len(allu), "capped": len(capped), "sampled": len(sampled), "strata": len(frame),
                        "double_coded": len(dbl), "sampled_not_event_level": sum(1 for u in sampled if not u["event_level"])}
    return {"all": allu, "capped": capped, "frame": frame,
            "sample": sorted(sampled, key=lambda u: (u["stratum"], u["passage_id"]))}


KWIC_FIELDS = ["passage_id", "theme", "subtheme", "stance", "term_family", "left_ctx", "term", "right_ctx",
               "sentence", "context_before", "context_after", "event", "event_level", "event_year", "doc_year", "countries",
               "stream", "layer",
               "site", "source_link", "orgs", "in_event_title", "in_handle", "in_sample", "sample_weight",
               "double_coded", "relevance", "frame_function", "code_1", "code_2", "code_3", "climate_linkage", "memo",
               "coder"]


# ==========================================================================
# Charts (matplotlib, Agg backend, 150 dpi; colour-blind-safe slots with a
# legend, markers as a second encoding and an n= label on every bar or point)
# ==========================================================================
def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.edgecolor": INK2, "axes.labelcolor": INK, "text.color": INK,
                         "xtick.color": INK2, "ytick.color": INK2, "axes.grid": True, "grid.color": GRID,
                         "grid.linewidth": 0.6, "axes.axisbelow": True, "axes.spines.top": False,
                         "axes.spines.right": False, "figure.facecolor": SURFACE, "axes.facecolor": SURFACE,
                         "legend.frameon": False, "svg.hashsalt": "analytics"})
    return plt


def _save(plt, fig, path: Path) -> None:
    tmp = _atomic(path)
    fig.savefig(tmp, dpi=150, format="png", bbox_inches="tight", metadata={"Software": None})
    plt.close(fig)
    os.replace(tmp, path)


def _stream_colour(st: str, i: int = 0) -> str:
    return STREAM_COLOUR.get(st, EXTRA_COLOURS[i % len(EXTRA_COLOURS)])


def chart_coverage(an: Analysis, plt, path: Path) -> None:
    """01: events per year stacked by stream; below it (same x, own y axis) the share
    of events whose linked site was read, with Wilson intervals."""
    years = sorted({e["year"] for e in an.events if e["year"]})
    if not years:
        return
    streams = sorted({e["stream_primary"] for e in an.events})
    fig, (ax, ax2) = plt.subplots(2, 1, figsize=(9, 6.5), sharex=True, gridspec_kw={"height_ratios": [3, 2]})
    bottom = [0] * len(years)
    for i, st in enumerate(streams):
        vals = [sum(1 for e in an.events if e["year"] == y and e["stream_primary"] == st) for y in years]
        ax.bar(years, vals, bottom=bottom, color=_stream_colour(st, i), label=st, width=0.8,
               edgecolor=SURFACE, linewidth=1)
        bottom = [a + b for a, b in zip(bottom, vals)]
    for x, n in zip(years, bottom):
        ax.annotate(f"n={n}", (x, n), xytext=(0, 2), textcoords="offset points", ha="center", fontsize=7, color=INK2)
    ax.set_ylabel("events")
    ax.set_title("Events per year by listing stream", loc="left")
    ax.legend(title="stream", ncol=len(streams))
    ks = [sum(1 for e in an.events if e["year"] == y and e["has_web_text"]) for y in years]
    ns = bottom
    sh = [k / n if n else 0 for k, n in zip(ks, ns)]
    ci = [wilson(k, n) for k, n in zip(ks, ns)]
    ax2.errorbar(years, sh, yerr=[[max(0, v - (c[0] or 0)) for v, c in zip(sh, ci)],
                                  [max(0, (c[1] or 0) - v) for v, c in zip(sh, ci)]],
                 color=INK2, marker="o", ms=5, lw=2, capsize=2)
    for x, v, k, n in zip(years, sh, ks, ns):
        ax2.annotate(f"{k}/{n}", (x, v), xytext=(0, 6), textcoords="offset points", ha="center", fontsize=7,
                     color=INK2)
    ax2.set_ylim(0, 1.05)
    ax2.set_ylabel("share with web text")
    ax2.set_title("Share of events whose linked site was read"
                  + (" (proxy: events with web-layer rows)" if an.web_method.startswith("proxy") else ""), loc="left")
    ax2.set_xticks(years)
    ax2.tick_params(axis="x", rotation=45)
    _save(plt, fig, path)


def chart_ladder(an: Analysis, plt, path: Path) -> None:
    """02: each theme's share of the total at every level of the measurement ladder."""
    levels = ["raw_rows", "clean_rows", "unique_passages", "documents", "sites", "events_any", "events_own"]
    th = an.themes
    tot = [sum(an.ladder[t][lv] for t in th) for lv in levels]
    if not any(tot):
        return
    fig, ax = plt.subplots(figsize=(9, 4.5))
    left = [0.0] * len(levels)
    for i, t in enumerate(th):
        vals = [an.ladder[t][lv] / n if n else 0 for lv, n in zip(levels, tot)]
        ax.barh(levels, vals, left=left, color=THEME_COLOUR.get(t, EXTRA_COLOURS[i % 8]), label=t,
                edgecolor=SURFACE, linewidth=1.5)
        for j, (v, lf) in enumerate(zip(vals, left)):
            if v >= 0.08:
                ax.text(lf + v / 2, j, str(an.ladder[t][levels[j]]), ha="center", va="center", fontsize=7,
                        color="white")
        left = [a + b for a, b in zip(left, vals)]
    for j, n in enumerate(tot):
        ax.text(1.01, j, f"n={n}", va="center", fontsize=7, color=INK2)
    ax.invert_yaxis()
    ax.set_xlim(0, 1.12)
    ax.set_xlabel("share of the level's total")
    ax.set_title("Measurement ladder: theme shares from raw rows to events", loc="left")
    ax.legend(title="theme", ncol=len(th), loc="lower center", bbox_to_anchor=(0.5, -0.32))
    _save(plt, fig, path)


def chart_prevalence(an: Analysis, plt, path: Path) -> None:
    """03: own-text prevalence per period, one panel per theme, one marker per stream."""
    rows = [r for r in an.prev_sp if r["measure"] == "own" and r["stream"] != "all"
            and "(all listings)" not in r["stream"] and r["group"] in [p[2] for p in PERIODS] and r["n_events"]]
    if not rows:
        return
    themes = an.themes + ["any"]
    periods = [p[2] for p in PERIODS]
    streams = sorted({r["stream"] for r in rows})
    fig, axes = plt.subplots(1, len(themes), figsize=(3.2 * len(themes), 3.8), sharey=True)
    for ax, t in zip(axes, themes):
        for i, st in enumerate(streams):
            pts = [r for r in rows if r["theme"] == t and r["stream"] == st]
            off = (i - (len(streams) - 1) / 2) * 0.18
            xs = [periods.index(r["group"]) + off for r in pts]
            ys = [r["share"] for r in pts]
            lo = [max(0, r["share"] - r["wilson_lo"]) for r in pts]
            hi = [max(0, r["wilson_hi"] - r["share"]) for r in pts]
            ax.errorbar(xs, ys, yerr=[lo, hi], fmt=STREAM_MARKER.get(st, "D"), color=_stream_colour(st, i),
                        ms=6, capsize=2, lw=1.2, label=st)
            for r in pts:
                # k/n in a band at the top of the panel, one line per stream (in its colour),
                # centred on the period: labels next to the points overlapped
                ax.text(periods.index(r["group"]), 0.97 - 0.055 * i, f"{r['k_events']}/{r['n_events']}",
                        transform=ax.get_xaxis_transform(), ha="center", va="top", fontsize=6,
                        color=_stream_colour(st, i))
        ax.set_xticks(range(len(periods)))
        ax.set_xticklabels(periods, rotation=45)
        ax.set_title(t, loc="left")
    # leave the top band (one text line per stream) free of data
    top = max([r["wilson_hi"] for r in rows if r["wilson_hi"] is not None] + [0.05])
    axes[0].set_ylim(0, top / (1 - 0.055 * len(streams) - 0.05))
    axes[0].set_ylabel("share of events (own listing text)")
    axes[-1].legend(title="stream (k/n at top)", loc="center left", bbox_to_anchor=(1.02, 0.5), fontsize=7,
                    title_fontsize=7)
    fig.suptitle("Theme in the event's own listing text, by period and stream (Wilson 95 % intervals)", x=0.01,
                 y=1.04, ha="left")
    _save(plt, fig, path)


def chart_concentration(an: Analysis, plt, path: Path) -> None:
    """04: the 12 documents with the most clean theme rows, stacked by theme."""
    c = defaultdict(Counter)
    for inst, d in _clean_inst_docs(an):
        c[d][inst["theme"]] += 1
    top = sorted(c, key=lambda d: (-sum(c[d].values()), d))[:12]
    if not top:
        return
    tot = sum(sum(v.values()) for v in c.values())
    fig, ax = plt.subplots(figsize=(10, 0.45 * len(top) + 1.6))
    left = [0] * len(top)
    for i, t in enumerate(an.themes):
        vals = [c[d][t] for d in top]
        ax.barh(range(len(top)), vals, left=left, color=THEME_COLOUR.get(t, EXTRA_COLOURS[i % 8]), label=t,
                edgecolor=SURFACE, linewidth=1.5)
        left = [a + b for a, b in zip(left, vals)]
    for j, n in enumerate(left):
        ax.text(n, j, f" n={n} ({100 * n / tot:.0f} %)", va="center", fontsize=7, color=INK2)
    ax.set_yticks(range(len(top)))
    ax.set_yticklabels([short(pdf_display(d), 70) for d in top], fontsize=7)
    ax.invert_yaxis()
    ax.set_xlabel("clean theme rows")
    ax.set_title(f"Documents with the most clean theme rows (all documents: {len(c)}, rows: {tot})", loc="left")
    ax.legend(title="theme")
    ax.set_xlim(0, max(left) * 1.25)
    _save(plt, fig, path)


def _component_layout(G, seed: int, weight: Optional[str] = "weight", spread: float = 1.6) -> Dict[str, Tuple[float, float]]:
    """Spring layout per connected component, the components packed in rows.
    One spring layout of a disconnected graph pushes the components far apart and
    squeezes each one, so labels in a dense component overlap."""
    import networkx as nx
    comps = sorted((sorted(c) for c in nx.connected_components(G)), key=lambda c: (-len(c), c[0]))
    radius = [max(0.6, len(c) ** 0.5) for c in comps]
    row_width = max(2 * radius[0], sum(2 * r for r in radius) ** 0.5 * 2.2)
    pos, x, y, row_h = {}, 0.0, 0.0, 0.0
    for c, r in zip(comps, radius):
        if x > 0 and x + 2 * r > row_width:
            x, y, row_h = 0.0, y - row_h - 0.8, 0.0
        # a copy with sorted nodes and edges: the layout must not depend on insertion order
        sub = nx.Graph()
        sub.add_nodes_from(c)
        sub.add_edges_from(sorted((min(u, v), max(u, v), d) for u, v, d in G.subgraph(c).edges(data=True)))
        if len(c) == 1:
            p = {c[0]: (0.0, 0.0)}
        elif len(c) > 3 and nx.density(sub) >= 0.5:
            p = nx.circular_layout(sub, scale=r)     # near-cliques: a spring layout piles the labels up
        elif len(c) > 30 and hasattr(nx, "forceatlas2_layout"):
            # ForceAtlas2 (numpy only) with strong gravity keeps dense groups and their bridges
            # readable where a spring layout throws a near-clique far out
            p = nx.forceatlas2_layout(sub, max_iter=400, seed=seed, strong_gravity=True, gravity=0.5,
                                      scaling_ratio=4.0, weight=weight)
            p = nx.rescale_layout_dict(p, scale=r)
        else:
            p = nx.spring_layout(sub, seed=seed, weight=weight, k=spread / len(c) ** 0.5, iterations=200, scale=r)
        for n, (px, py) in p.items():
            pos[n] = (x + r + float(px), y - r + float(py))
        x += 2 * r + 0.8
        row_h = max(row_h, 2 * r)
    return pos


def chart_backbone(an: Analysis, plt, path: Path) -> None:
    """05: the strongest organisations of the disparity backbone: the top --chart-nodes by
    strength among those in two or more series (spring layout per component, fixed
    seed; under 500 nodes networkx lays out with numpy alone, no scipy needed)."""
    import networkx as nx
    B, G = an.B, an.G
    n_top = max(1, int(getattr(an.a, "chart_nodes", 150) or 150))
    multi = [n for n in B if n in G and G.degree(n) >= 2]
    keep = sorted(multi, key=lambda n: (-B.degree(n, weight="weight"), n))[:n_top]
    S = B.subgraph(keep).copy()
    S.remove_nodes_from([n for n in list(S) if S.degree(n) == 0])
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.set_axis_off()
    title = (f"Organisation backbone, scenario {an.default_sc}, alpha {an.a.backbone_alpha}: the "
             f"{count_of(S.number_of_nodes(), 'organisation')} with the highest strength among the "
             f"{len(multi)} in two or more series (top {n_top}; {count_of(S.number_of_edges(), 'tie')} shown; "
             f"whole backbone {B.number_of_nodes()} organisations, {B.number_of_edges()} ties)"
             + ("" if an.curated else "\n(uncurated: colours are the suggested_class guesses of "
                "org_codes_template.xlsx, not codes)"))
    ax.set_title(title, loc="left", fontsize=8, wrap=True)
    if not S.number_of_edges():
        # an empty projection (no two organisations share an event) is not the same as a strict alpha
        msg = ("No backbone ties among organisations in two or more series" if an.P.number_of_edges() else
               "No two organisations share an event series, so there is nothing to project")
        ax.text(0.5, 0.5, msg, ha="center", va="center", transform=ax.transAxes, color=INK2)
        _save(plt, fig, path)
        return
    # unweighted layout: the heaviest ties would otherwise pull whole clusters onto one point
    pos = _component_layout(S, an.a.seed, weight=None, spread=2.4)
    ax.set_aspect("equal", adjustable="datalim")
    ax.margins(0.08)
    cen = {"O:" + r["org_key"]: r for r in an.org_centrality}
    sug = {"O:" + r["org_key"]: r.get("suggested_class") for r in getattr(an, "authority", [])}

    def cls(n: str) -> str:
        if an.curated:
            return cen.get(n, {}).get("actor_class") or "not coded"
        return sug.get(n) or "no suggestion"
    classes = sorted({cls(n) for n in S})
    col = {c: (INK2 if c in ("not coded", "no suggestion") else EXTRA_COLOURS[i % len(EXTRA_COLOURS)])
           for i, c in enumerate(classes)}
    wmax = max(d["weight"] for _, _, d in S.edges(data=True)) or 1.0
    widths = [0.4 + 2.6 * d["weight"] / wmax for _, _, d in S.edges(data=True)]
    nx.draw_networkx_edges(S, pos, ax=ax, width=widths, edge_color=GRID, alpha=0.9)
    for c in classes:
        nodes = sorted(n for n in S if cls(n) == c)
        nx.draw_networkx_nodes(S, pos, nodelist=nodes, ax=ax, node_color=col[c],
                               node_size=[20 + 12 * math.sqrt(G.degree(n)) for n in nodes],
                               edgecolors=SURFACE, linewidths=1.0)
        ax.scatter([], [], s=60, color=col[c], label=c)     # fixed legend marker, whatever the node sizes
    strength = sorted(S, key=lambda n: (-B.degree(n, weight="weight"), n))[:30]
    ys = [y for _, y in pos.values()]
    dy = 0.02 * ((max(ys) - min(ys)) or 1)
    lpos = {n: (x, y + dy) for n, (x, y) in pos.items()}     # labels just above their node
    nx.draw_networkx_labels(S, lpos, labels={n: short(S.nodes[n].get("label", n[2:]), 28) for n in strength},
                            ax=ax, font_size=6.5, font_color=INK)
    ax.legend(title=("actor_class" if an.curated else "suggested_class")
              + "\n(node size = 20 + 12 x sqrt(series);\nlabels: the 30 strongest)", loc="upper left",
              bbox_to_anchor=(1.0, 1.0), fontsize=7, title_fontsize=7)
    _save(plt, fig, path)


def chart_country(an: Analysis, plt, path: Path) -> None:
    """06: prevalence per country (n >= 15 events), own text and any layer side by side."""
    rows = [r for r in an.tables["prev_country"][0] if r["theme"] == "any" and r["n_events"] >= 15
            and not r["group"].startswith(("(", "other"))]
    groups = sorted({r["group"] for r in rows if r["measure"] == "own"},
                    key=lambda g: (-next(r["n_events"] for r in rows if r["group"] == g and r["measure"] == "own"), g))
    if not groups:
        return
    labels = {"own": "own listing text (all events)", "any_web": "any layer (events with web text)"}
    # the any-layer panel only when at least one country has 15 or more events with web text
    measures = ["own"] + (["any_web"] if any(r["measure"] == "any_web" and r["n_events"] >= 15 for r in rows)
                          else [])
    fig, axes = plt.subplots(1, len(measures), figsize=(5 * len(measures), 0.35 * len(groups) + 1.8),
                             sharey=True, squeeze=False)
    axes = list(axes[0])
    if len(measures) == 1:
        fig.text(0.01, -0.02, "Any-layer panel left out: no country has 15 or more events whose linked site "
                 "was read.", fontsize=7, color=INK2)
    for ax, m in zip(axes, measures):
        for j, g in enumerate(groups):
            r = next((x for x in an.tables["prev_country"][0] if x["group"] == g and x["measure"] == m
                      and x["theme"] == "any"), None)
            if not r or r["share"] is None or r["n_events"] < 15:
                ax.text(0, j, f"  n={r['n_events'] if r else 0} (fewer than 15, not shown)", va="center",
                        fontsize=7, color=INK2)
                continue
            ax.errorbar([r["share"]], [j], xerr=[[r["share"] - r["wilson_lo"]], [r["wilson_hi"] - r["share"]]],
                        fmt="o", color=THEME_COLOUR["colonialism"], ms=6, capsize=2)
            ax.annotate(f"{r['k_events']}/{r['n_events']}", (r["wilson_hi"], j), xytext=(4, -3),
                        textcoords="offset points", fontsize=7, color=INK2)
        ax.set_title(labels[m], loc="left")
        top = [x["wilson_hi"] for x in rows if x["measure"] == m and x["wilson_hi"] is not None]
        ax.set_xlim(0, min(1.0, max(top + [0.1]) * 1.25))
        ax.set_xlabel("share of events with any theme")
    axes[0].set_yticks(range(len(groups)))
    axes[0].set_yticklabels([f"{an.p.country_name(g)} ({g})" for g in groups])
    axes[0].invert_yaxis()
    fig.suptitle("Theme prevalence by event country (countries with at least 15 events; Wilson 95 % intervals)",
                 x=0.01, ha="left")
    _save(plt, fig, path)


def write_charts(an: Analysis, out: Path) -> List[str]:
    plt = _plt()
    made = []
    for name, fn in (("01_coverage_frame.png", chart_coverage), ("02_measurement_ladder.png", chart_ladder),
                     ("03_prevalence_stream_period.png", chart_prevalence),
                     ("04_doc_concentration.png", chart_concentration), ("05_backbone.png", chart_backbone),
                     ("06_prevalence_country.png", chart_country)):
        path = out / "charts" / name
        try:
            fn(an, plt, path)
        except Exception as exc:  # noqa: BLE001 - one chart must not stop the others
            an.warn("chart %s not drawn: %s: %s", name, type(exc).__name__, exc)
            plt.close("all")
            continue
        if path.exists():
            made.append(name)
        else:
            an.skipped[f"chart {name}"] = "nothing to draw"
    return made


# ==========================================================================
# Network files (Gephi: GEXF and GraphML, plus node and edge CSVs)
# ==========================================================================
def _node_rows(G) -> Tuple[List[dict], List[str]]:
    fields = ["id"] + sorted({k for _, d in G.nodes(data=True) for k in d})
    return [{"id": n, **d} for n, d in sorted(G.nodes(data=True))], fields


def _edge_rows(G) -> Tuple[List[dict], List[str]]:
    fields = ["source", "target"] + sorted({k for _, _, d in G.edges(data=True) for k in d})
    rows = [{"source": min(a, b), "target": max(a, b), **d} for a, b, d in G.edges(data=True)]
    return sorted(rows, key=lambda r: (r["source"], r["target"])), fields


def write_networks(an: Analysis, out: Path) -> List[str]:
    cen = {"O:" + r["org_key"]: r for r in an.org_centrality}
    made = []
    for sc, (G, B) in an.exports.items():
        for n in G:
            if n in cen:     # centralities as node attributes (only for the scenario they were computed on)
                if sc == an.default_sc:
                    c = cen[n]
                    G.nodes[n].update(n_series=c["n_series"], bip_betweenness=c["bip_betweenness"],
                                      in_backbone=c["in_backbone"], backbone_strength=c["backbone_strength"],
                                      hub_dependence=c["hub_dependence"], core_2_2=c["core_2_2"])
        for n in B:
            B.nodes[n]["strength"] = round(B.degree(n, weight="weight"), 4)
            B.nodes[n]["n_series"] = G.degree(n) if n in G else 0
        base = out / "network"
        for g, stem in ((G, f"bipartite_{sc}"), (B, f"org_backbone_{sc}")):
            made += [p.name for p in write_graph(g, base / stem)]
        if sc == an.default_sc:
            # the backbone's core: organisations in two or more series (small enough for Gephi at once)
            import networkx as nx
            core = B.subgraph([n for n in B if B.nodes[n]["n_series"] >= 2]).copy()
            core.remove_nodes_from([n for n in list(core) if core.degree(n) == 0])
            p = base / f"org_backbone_core_{sc}.gexf"
            tmp = _atomic(p)
            nx.write_gexf(clean_graph(core), str(tmp))
            os.replace(tmp, p)
            made.append(p.name)
            an.facts["backbone_core"] = {"nodes": core.number_of_nodes(), "edges": core.number_of_edges()}
        suffix = "" if sc == an.default_sc else f"_{sc}"
        for g, stem in ((G, "bipartite"), (B, "backbone")):
            rows, fields = _node_rows(clean_graph(g))
            write_csv(base / f"nodes_{stem}{suffix}.csv", rows, fields)
            rows, fields = _edge_rows(clean_graph(g))
            write_csv(base / f"edges_{stem}{suffix}.csv", rows, fields)
            made += [f"nodes_{stem}{suffix}.csv", f"edges_{stem}{suffix}.csv"]
    return made


# ==========================================================================
# Workbooks
# ==========================================================================
SHEET_DOCS = OrderedDict([
    ("README", "This sheet: what each sheet holds."),
    ("diagnostics", "How the inputs were read: layouts, joins of event references, theme filter counts, "
                    "country sources, methods used."),
    ("missingness", "What is missing or uncertain, as k of n (thesis table 2)."),
    ("coverage_by_stream", "Coverage frame per listing stream: links, links read, organisation ties, theme hits, "
                           "countries, day-precise dates; shares with Wilson 95 % intervals (thesis table 1)."),
    ("events_enriched", "One row per event (events.csv) with event_id, series_id, streams, countries, coverage "
                        "and theme flags per layer: own (listing text), landing (event link), linked (other pages "
                        "and PDFs), any, in_title."),
    ("ties", "One row per event x organisation x source, with tie level (L1 listing, L2 landing page, L3 other "
             "linked pages and PDFs), tie_sublevel (L3s on the event's own host, L3x other hosts), doc_kind, "
             "attribution, participation and context_reason, and whether the source passes the theme filter."),
    ("theme_units", "One row per passage x theme x event; passage_id is the stable key for coded workbooks."),
    ("measurement_ladder", "Each theme counted at seven levels, from raw rows to events."),
    ("concentration", "How much of the theme material comes from a few documents (eff_docs = 1 / sum of squared "
                      "shares)."),
    ("attribution", "Clean theme references and documents per layer and theme."),
    ("event_theme_documents", "Per event: theme rows (outside menus and footers), distinct documents and passages, "
                              "by tie sublevel; one_document_dominates when one document gives more than half "
                              "of the rows (and the event has two or more documents)."),
    ("descriptive_themes", "Descriptive themes (global_south, authoritarianism): counted, never used for the "
                           "theme filter or the any-theme measures."),
    ("similar_sources", "Source pairs naming nearly the same organisations (Jaccard >= 0.9, at least 15 each) "
                        "but tied to different series: unmarked copies."),
    ("multi_name_series", "Series (shared event link) whose editions carry three or more unrelated names: "
                          "probably different events; split them with event_codes.csv series_id."),
    ("term_families", "Theme-file lines (term families) with their rows, passages, documents, events and surface "
                      "forms."),
    ("prev_stream_period", "Theme prevalence (focus themes) with the event as unit, per period and stream. own, "
                           "any_all and in_title use all events; landing, linked, any_web and core (listing, "
                           "landing page and the event's own host, no reports) only the events whose link was "
                           "read (column denominator), with share_std standardised over words read. any_all is "
                           "never lower than own; any_web counts fewer events because its denominator is smaller."),
    ("prev_stream_year", "As prev_stream_period, per year."),
    ("prev_words_read", "Web-layer prevalence by words read on the web for the event (exposure bins <1k ... "
                        ">=100k). share_std in the other prev_* tables standardises over these bins."),
    ("prev_country", "As prev_stream_period, per event country (an event in two countries counts in both; "
                     "countries with fewer than 15 events pooled; headline_ok means at least 30 events)."),
    ("prev_lang", "Own-text prevalence per language of the agenda entry."),
    ("entry_rates", "Theme hits per 10,000 words of agenda entry text (Byar 95 % intervals) and the share of "
                    "entries with a hit."),
    ("page_rates", "Theme hits per 10,000 words of the sources that were read, per site role and kind."),
    ("sensitivity", "The event x organisation network under scenarios A-D (and C_curated), for the layer sets "
                    "L1-L3s, L1-L3 and L1+L2, at edition and series level, with and without the theme filter; "
                    "then variants of the default network (column variant): +context rows, reports included, "
                    "cap 50 organisations per source, without common_word? suggestions, +L4 site-level rows."),
    ("tie_levels", "Rows, organisations and sources per tie sublevel and kind of source in the default scenario, "
                   "for every row and for the participation rows."),
    ("org_centrality", "Organisations of the default network: series, years, countries, bip_betweenness "
                       "(shortest-path betweenness on the event-organisation graph, normalised over all its "
                       "nodes by (n-1)(n-2)/2 as for a one-mode graph, not by the bipartite maximum), backbone "
                       "strength, hub dependence, (2,2)-core, country participation."),
    ("org_theme", "Organisation x theme: passage level (organisation named in the theme paragraph, counted in "
                  "documents) and event level (series whose own or landing text carries the theme). share, "
                  "lift and shrunk are blank when the masters were already theme-filtered (see summary.md, C11)."),
    ("communities", "Louvain communities of the default backbone (weight, --seed): size, strongest members, "
                    "dominant countries, share of members seen in one source only, series covered."),
    ("org_bridging", "Organisations robustly tied to two or more themes (robust: two documents from two "
                     "websites, or the event-level rule on unfiltered masters), by series, then documents."),
    ("theme_cooccurrence", "Themes together in the same passage, document or series, with a leave-one-site-out "
                           "range."),
])


def write_analytics_workbook(an: Analysis, out: Path) -> None:
    readme = [{"sheet": k, "description": v} for k, v in SHEET_DOCS.items()
              if k in ("README", "diagnostics") or k in an.tables]
    readme += [{"sheet": k, "description": "(see summary.md)"} for k in an.tables if k not in SHEET_DOCS]
    readme.append({"sheet": "", "description": ""})
    readme.append({"sheet": "theme filter", "description": f"'{an.a.theme_filter}' (applied by make_master.py: "
                                                           f"'{an.master_filter or 'unknown'}')"})
    readme.append({"sheet": "curated", "description": yn(an.curated) + (
        "" if an.curated else ": organisations are not curated yet; read network results as provisional")})
    for k, v in an.skipped.items():
        readme.append({"sheet": "skipped", "description": f"{k}: {v}"})
    sheets = [("README", readme, ["sheet", "description"]),
              ("diagnostics", an.diag, ["item", "value", "note"])]
    sheets += [(k, rows, fields) for k, (rows, fields) in an.tables.items()]
    # Very long tables go to CSV only (writing them to xlsx takes minutes; Excel stops at 1,048,576 rows).
    big = [k for k, rows, _ in sheets if len(rows) > XLSX_MAX_ROWS]
    for k in big:
        readme.append({"sheet": k, "description": f"more than {XLSX_MAX_ROWS:,} rows: only in tables/{k}.csv"})
    write_xlsx(out / "analytics.xlsx", [x for x in sheets if x[0] not in big])
    for title, rows, fields in sheets:
        write_csv(out / "tables" / f"{title}.csv", rows, fields)


XLSX_MAX_ROWS = 100_000

AUTH_FIELDS = ["org_key", "organisation", "priority", "variants", "max_confidence", "found_by", "route", "n_events",
               "n_series", "in_default_network",
               "n_years", "years", "countries", "n_sources", "n_sites", "sites", "tie_levels", "only_in_event_title",
               "only_menus_footers", "theme_passages", "themes_in_paragraph", "example_context_1",
               "example_context_2", "suggested_class", "flags", "variant_group", "merge_into_suggested",
               "include", "merge_into_key", "canonical_name", "actor_class", "country", "is_organisation",
               "is_partial_name", "coder", "notes"]


def write_org_template(an: Analysis, out: Path) -> None:
    readme = [{"text": x} for x in [
        "Organisation authority list. Fill in the curation columns (include ... notes) of the 'organisations' "
        "sheet, save that sheet as org_codes.csv (UTF-8) and run analytics.py again with --org-codes org_codes.csv.",
        "include: 1 keeps the organisation in the network whatever its confidence, 0 drops it, blank leaves it to "
        "--min-confidence.",
        "merge_into_key: the org_key this one is the same organisation as (e.g. 'ende gelaende' -> 'ende gelande').",
        "canonical_name: the name to show. actor_class: one of " + ", ".join(ACTOR_CLASSES) + ". With codes, "
        "place, site_ui, noise and event_name are left out of the curated scenario (C_curated).",
        "is_organisation (yes / no / partial) and is_partial_name: precision coding. The precision_sample sheet "
        "lists a stratified random sample to code first; once coded, analytics writes precision_by_stratum.",
        "suggested_class and flags are guesses from simple rules, not codes. merge_candidates lists pairs that "
        "may be the same organisation (umlaut variants are merged automatically).",
        f"The list covers organisations from sources that pass the theme filter '{an.a.theme_filter}', context "
        "rows included (flag only_in_reference_docs: named only in linked reports and studies).",
        "priority: 1 = in two or more event series of the default network, 2 = in one series there through the "
        "listing or the landing page, 3 = the rest; code in this order. Organisations in no series at all "
        "(priority 4) are in the no_series sheet.",
        "suggested_class 'common_word?' marks one-word names that read as ordinary words (weekdays, or a noun, "
        "adjective or verb in their own contexts by the language's part-of-speech tagger). It is a suggestion: "
        "set include 0 to drop such a name.",
        "sources_to_code: every page and PDF with its organisations and what the analysis took it for "
        "(doc_kind; 'report' sources are left out of the default network). Fill in include (1 = evidence of "
        "participation, 0 = leave out), save the workbook and run analytics.py with --source-codes "
        "org_codes_template.xlsx (or the sheet saved as CSV).",
        "Keys are re-derived from the organisation name when org_codes.csv is read, so a stale org_key only "
        "produces a warning."]]
    main = [r for r in an.authority if r["priority"] != 4]
    no_series = [r for r in an.authority if r["priority"] == 4]
    sheets = [("organisations", main, AUTH_FIELDS),
              ("no_series", no_series, AUTH_FIELDS),
              ("sources_to_code", an.sources_to_code, SOURCE_FIELDS),
              ("merge_candidates", an.merge_cands, ["key_a", "name_a", "key_b", "name_b", "reason", "score", "accept"]),
              ("precision_sample", an.precision_sample,
               ["stratum", "stratum_N", "n", "weight", "org_key", "organisation", "found_by", "n_events",
                "example_context_1", "is_organisation", "is_partial_name", "coder"])]
    if an.precision_rows:
        sheets.append(("precision_by_stratum", an.precision_rows,
                       ["stratum", "N", "n", "k", "precision", "wilson_lo", "wilson_hi", "se", "k_lenient",
                        "precision_lenient", "lenient_lo", "lenient_hi"]))
    sheets.append(("README", readme, ["text"]))
    yn_ = ["yes", "no", "partial"]
    write_xlsx(out / "org_codes_template.xlsx", sheets, validations={
        "organisations": {"include": ["1", "0"], "actor_class": ACTOR_CLASSES, "is_organisation": yn_,
                          "is_partial_name": ["yes", "no"]},
        "no_series": {"include": ["1", "0"], "actor_class": ACTOR_CLASSES, "is_organisation": yn_},
        "sources_to_code": {"include": ["1", "0"]},
        "merge_candidates": {"accept": ["yes", "no"]},
        "precision_sample": {"is_organisation": yn_, "is_partial_name": ["yes", "no"]}})
    if an.precision_rows:
        an.table("precision_by_stratum", an.precision_rows,
                 ["stratum", "N", "n", "k", "precision", "wilson_lo", "wilson_hi", "se", "k_lenient",
                  "precision_lenient", "lenient_lo", "lenient_hi"])


CODEBOOK = [
    ("relevance", "substantive", "The passage is about the theme (colonialism, racism, fascism) in its own right."),
    ("relevance", "name_or_title_only", "The term only appears in a name or title (an organisation, event, book)."),
    ("relevance", "list_or_tag", "The term is one item of a list of causes or a tag, without development."),
    ("relevance", "off_topic_sense", "The word is used in another sense (e.g. 'colony' of animals)."),
    ("relevance", "false_match", "The match is not the term at all (part of another word, a handle, a URL)."),
    ("relevance", "unclear", "Cannot be decided from the passage."),
    ("frame_function", "(open)", "What the theme does in the passage: diagnosis, motivation, identity, alliance, "
                                 "target, history ... to be fixed after the first coding round."),
    ("code_1..3", "(open)", "Inductive codes."),
    ("climate_linkage", "(open)", "How the passage links the theme to climate (none, causal, justice, solidarity ...)."),
    ("memo", "(free text)", "Notes."),
    ("coder", "(initials)", "Who coded the row. Rows flagged double_coded are coded by two people (reliability)."),
    ("sentence", "", "Use the sentence for attribution: in agenda PDFs the context columns can run into the "
                     "neighbouring entry (caveat C6)."),
    ("sample_weight", "", "N_h / n_h of the row's stratum (theme x layer group x period); use it for weighted "
                          "shares of the sample."),
]


def write_kwic(an: Analysis, out: Path, kw: dict) -> None:
    readme = [{"text": x} for x in [
        "Theme passages in keyword-in-context layout for qualitative coding. One row per passage and theme "
        "(furniture and pipeline duplicates removed).",
        f"The per-theme sheets hold at most {an.a.doc_cap} passages per document (--doc-cap), so one long "
        "brochure cannot dominate; all_units has every passage.",
        f"in_sample marks a stratified random sample (theme x layer group x period, proportional with at least 2 "
        f"per stratum, {an.a.sample_size} rows at most, seed {an.a.seed}); double_coded marks 25 % of it for a "
        "second coder. sampling_frame lists the strata with their weights; coding_sample holds the sampled "
        "passages only, sorted by stratum, with sample_weight and double_coded.",
        "event_level: 'no' when the passage is tied to no event (other pages of a website, or not tied); such "
        "passages count for the theme ladder but not for event prevalence.",
        "passage_id is stable across runs: keep it when copying rows, it is the key for merging coded sheets back.",
        "Code relevance first (dropdown); the codebook sheet explains every column.",
        "Caveat C6: in agenda PDFs the context columns can include text of the neighbouring entry; attribute "
        "with the sentence column."]]
    sheets = [("README", readme, ["text"])]
    for t in an.themes:
        rows = sorted([u for u in kw["capped"] if u["theme"] == t], key=lambda u: (not u["in_sample"],
                                                                                   u["stratum"], u["passage_id"]))
        sheets.append((t, rows, KWIC_FIELDS))
    sheets.append(("coding_sample", kw["sample"], ["stratum"] + KWIC_FIELDS + ["document"]))
    sheets.append(("all_units", sorted(kw["all"], key=lambda u: (u["theme"], u["passage_id"])),
                   KWIC_FIELDS + ["document", "stratum"]))
    sheets.append(("sampling_frame", kw["frame"], ["stratum", "theme", "layer_group", "period", "N_h", "n_h",
                                                   "weight"]))
    sheets.append(("codebook", [{"column": a, "value": b, "meaning": c} for a, b, c in CODEBOOK],
                   ["column", "value", "meaning"]))
    lists = [{"relevance": r, "actor_class": ACTOR_CLASSES[i] if i < len(ACTOR_CLASSES) else ""}
             for i, r in enumerate(RELEVANCE)]
    lists += [{"relevance": "", "actor_class": c} for c in ACTOR_CLASSES[len(RELEVANCE):]]
    sheets.append(("lists", lists, ["relevance", "actor_class"]))
    val = {t: {"relevance": RELEVANCE} for t in an.themes}
    val["all_units"] = {"relevance": RELEVANCE}
    val["coding_sample"] = {"relevance": RELEVANCE}
    write_xlsx(out / "kwic_for_coding.xlsx", sheets, validations=val)


# ==========================================================================
# A9 provenance: summary.json and summary.md
# ==========================================================================
def _versions() -> dict:
    v = {"python": platform.python_version()}
    for mod in ("networkx", "matplotlib", "openpyxl"):
        try:
            v[mod] = __import__(mod).__version__
        except Exception:  # noqa: BLE001
            v[mod] = "not installed"
    return v


def _pct(k, n) -> str:
    return f"{k}/{n} ({100 * k / n:.1f} %)" if n else f"{k}/0"


def write_summary(an: Analysis, out: Path, outputs: Dict[str, List[str]], seconds: float) -> dict:
    a = an.a
    params = {k: v for k, v in sorted(vars(a).items()) if k not in ("selftest",)}
    files = {name: {"path": str(p), "sha256": sha256_of(p)} for name, p in sorted(an.inputs.items())}
    for name in ("analytics.py", "gazetteer.py"):
        files[name] = {"path": str(HERE / name), "sha256": sha256_of(HERE / name)}
    th = an.themes
    insts = an.instances
    clean = [i for i in insts if i["clean"]]
    stream_n = Counter(e["stream_primary"] for e in an.events)
    conc = an.concentration.get("ALL", {})
    own = {t: sum(1 for e in an.events if e[f"{t}_own"]) for t in th}
    anyl = {t: sum(1 for e in an.events if e[f"{t}_any"]) for t in th}
    web = [e for e in an.events if e["has_web_text"]]
    conf = Counter(r["max_confidence"] for r in an.authority)
    summary = {
        "tool": "analytics.py", "version": VERSION,
        "made_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "seconds": round(seconds, 1), "results_dir": str(an.results), "runs": [str(r) for r in an.runs],
        "parameters": params, "versions": _versions(), "files": files,
        "layout": {"master_network": "long" if an.long_net else "older", "master_themes": "long" if an.long_th
                   else "older"},
        "theme_filter": {"analytics": a.theme_filter, "make_master": an.master_filter or "unknown",
                         **{k: {"kept": v[0], "total": v[1], "left_out": v[1] - v[0]}
                            for k, v in an.filter_counts.items()},
                         "excluded_sources_in_master": len(an.excluded_sources),
                         "make_master_counts": {k: {"kept_by_theme_filter": v[0], "before_filter": v[1],
                                                    "in_master_file": v[2]}
                                                for k, v in an.master_counts.items()}},
        "curated": yn(an.curated), "has_web_text_method": an.web_method, "country_from": an.country_from,
        "betweenness": getattr(an, "betweenness_method", ""),
        "match": {"network": {k: dict(v) for k, v in an.ref_counts.items()}, "themes": dict(an.theme_ref_counts),
                  "unmatched_names": dict(an.unmatched_names.most_common(20))},
        "numbers": {"events": len(an.events), "events_by_stream": dict(sorted(stream_n.items())),
                    "events_with_country": sum(1 for e in an.events if e["countries"]),
                    "events_with_web_text": len(web), "theme_rows": len(insts), "clean_rows": len(clean),
                    "furniture_rows": sum(1 for i in insts if i["furniture"]),
                    "duplicate_rows": sum(1 for i in insts if i["dup"] and not i["furniture"]),
                    "unique_passages": len({i["passage_id"] for i in clean}),
                    "documents": conc.get("documents", 0), "eff_docs": conc.get("eff_docs"),
                    "top1_doc": conc.get("top1_doc", ""), "top1_share": conc.get("top1_share"),
                    "own_text_events": own, "any_layer_events": anyl,
                    "organisations_by_max_confidence": dict(sorted(conf.items())),
                    **{k: v for k, v in an.facts.items() if k not in ("org_theme",)}},
        "skipped": dict(sorted(an.skipped.items())), "warnings": an.warnings, "outputs": outputs}
    write_text(out / "summary.json", json.dumps(summary, ensure_ascii=False, indent=2, default=str))
    write_text(out / "summary.md", summary_md(an, summary))
    return summary


def summary_md(an: Analysis, sm: dict) -> str:
    a, n = an.a, sm["numbers"]
    ne = n["events"]
    th = an.themes
    fc = an.filter_counts
    streams = ", ".join(f"{k} {v}" for k, v in n["events_by_stream"].items())
    dg = an.facts.get("default_graph", {})
    L = [f"# Analysis summary ({sm['made_at'][:10]})", "",
         f"Results folder: `{an.results}`. analytics.py {VERSION}; parameters and file checksums are in "
         "summary.json.", "", "## Methods (numbers filled in from this run)", ""]
    para = [f"The event list holds {ne} events ({streams}).",
            f"A country was assigned to {_pct(n['events_with_country'], ne)} events (from {an.country_from}).",
            f"For {_pct(n['events_with_web_text'], ne)} events the linked website was read "
            f"(method: {an.web_method}); web-layer shares use only these events as denominator.",
            f"The theme search produced {n['theme_rows']} rows; after removing {n['furniture_rows']} rows in "
            f"menus, footers or boilerplate and {n['duplicate_rows']} repeated rows, {n['clean_rows']} clean rows "
            f"remain, forming {n['unique_passages']} distinct passages in {n['documents']} documents (one agenda "
            "entry counts as one document)."]
    if n.get("eff_docs"):
        para.append(f"The material is concentrated: the effective number of documents is {n['eff_docs']}, and "
                    f"the largest document (`{pdf_display(n['top1_doc'])}`) gives "
                    f"{100 * (n['top1_share'] or 0):.0f} % of the clean rows.")
    para.append("Events with a theme in their own listing text: " + ", ".join(
        f"{t} {n['own_text_events'][t]}" for t in th) + "; on any layer: " + ", ".join(
        f"{t} {n['any_layer_events'][t]}" for t in th) + f" (of all {ne} events; measure any_all in the "
        "prev_* tables, while any_web counts only the events whose link was read).")
    mf_label = f"'{an.master_filter or 'unknown'}'"
    if an.master_filter == "none" and not an.master_about:
        mf_label += " (older masters, built before the filter existed)"
    para.append(f"Inclusion criterion: only sources with a theme hit count for the network (theme filter "
                f"'{a.theme_filter}'; make_master.py applied {mf_label}). This kept "
                f"{fc['rows'][0]} of {fc['rows'][1]} network rows, {fc['organisations'][0]} of "
                f"{fc['organisations'][1]} organisations, {fc['events'][0]} of {fc['events'][1]} events and "
                f"{fc['sources'][0]} of {fc['sources'][1]} sources"
                + (f"; {len(an.excluded_sources)} sources were already left out by make_master.py"
                   if an.excluded_sources else "") + ".")
    # Kept rows that cannot enter the event network: no event (pseudo rows such as
    # "(not tied to an event)") or the organisation was seen only in menus and footers.
    kept = [t for t in an.ties if t["sheet"] != "menus_and_footers" and t["pass"][a.theme_filter]]
    no_ev = sum(1 for t in kept if not t["event_id"])
    menus = sum(1 for t in kept if t["event_id"] and t["in_menus"])
    if no_ev or menus:
        para.append(f"Of the kept rows, {no_ev} are not tied to an event and {menus} name an organisation "
                    "found only in menus or footers; neither enters the event network.")
    ctx = Counter(k for t in kept if t["event_id"] and not t["in_menus"] and not t["participation"]
                  for k in t["ctx_kinds"])
    n_ctx = sum(1 for t in kept if t["event_id"] and not t["in_menus"] and not t["participation"])
    para.append(f"Network scope '{a.network_scope}': {n_ctx} event-tied rows are context, not evidence of taking "
                "part (" + (", ".join(f"{k} {v}" for k, v in sorted(ctx.items())) or "none") + "; doc_kind from "
                f"{an.doc_kind_from})" + (", and the default network leaves them out" if a.network_scope ==
                                          "participation" else ", but --network-scope all keeps them") + ".")
    mc = an.master_counts
    if mc:
        para.append("make_master.py's filter had already kept " + ", ".join(
            f"{v[0]} of {v[1]} {k.replace('_', ' ')}" for k, v in mc.items()) + " before this analysis" + (
            " (master_network holds " + ", ".join(f"{v[2]} {k.replace('_', ' ')}" for k, v in mc.items()
                                                  if v[2] is not None)
            + " after the NER-only names went to ner_only_candidates)."
            if any(v[2] is not None and v[2] != v[0] for v in mc.values()) else "."))
    conf = n["organisations_by_max_confidence"]
    para.append("Organisations in the authority list (org_codes_template.xlsx, after the theme filter, "
                "including those without an event or only in menus) by highest confidence: " + ", ".join(f"{k} {v}" for k, v in conf.items()) + ".")
    if dg:
        top = dg.get("top_event_share_of_proj_edges")
        cm = an.facts.get("communities") or {}
        para.append(f"The default network (scenario {dg['scenario']}, series level, layers {dg.get('layers', '')}, "
                    f"scope {dg.get('network_scope', '')}) has "
                    f"{dg['series']} event series, {count_of(dg['orgs'], 'organisation')} "
                    f"({dg.get('orgs_in_2plus', 0)} in two or more series) and "
                    f"{count_of(dg['edges'], 'tie')}; its projection has "
                    f"{count_of(dg['proj_edges'], 'organisation pair')}"
                    + (f" ({100 * top:.0f} % of them inside the largest event)" if top is not None else "")
                    + f", of which the disparity backbone (alpha {a.backbone_alpha}) keeps {dg['backbone_edges']} "
                    f"between {count_of(dg['backbone_nodes'], 'organisation')}"
                    + (f"; Louvain finds {cm['n']} communities in it (modularity {cm['modularity']}, "
                       f"{cm['n_10plus']} with 10 or more members)" if cm else "")
                    + f". Betweenness: {sm['betweenness']}.")
    prec = an.facts.get("precision")
    if prec:
        para.append(f"Coded organisation precision (stratified estimate): strict {prec['precision']} "
                    f"(95 % interval {prec.get('wilson_lo')} to {prec.get('wilson_hi')}), lenient "
                    f"{prec['precision_lenient']} ({prec.get('lenient_lo')} to {prec.get('lenient_hi')}; "
                    f"n = {prec['n']}).")
    else:
        para.append("Organisation names are not yet coded for precision (see org_codes_template.xlsx).")
    L += [" ".join(para), ""]
    # ---- caveats
    leak = an.facts.get("seed_context_leak", (0, 0))
    top_share = next((r["top_event_share_of_proj_edges"] for r in an.sensitivity
                      if r["scenario"] == "A" and r["layers"] == "L1-L3" and r["node_level"] == "edition"
                      and r["theme_filter"] == a.theme_filter), None)
    web_k = sum(1 for e in an.events if e["has_web_text"] and e["any_any"])
    noweb_k = sum(1 for e in an.events if not e["has_web_text"] and e["any_any"])
    nweb = n["events_with_web_text"]
    racism = conc_r = None
    if "racism" in th:
        racism = n["own_text_events"]["racism"]
        conc_r = an.concentration.get("racism", {}).get("top1_share")
    years = defaultdict(list)
    for e in an.events:
        if e["year"]:
            years[e["stream_primary"]].append(e["year"])
    frame = "; ".join(f"{k} {min(v)}-{max(v)}" for k, v in sorted(years.items()))
    caveats = [
        ("C1 Coverage, not framing", f"Web-layer theme shares apply only to events whose links were read "
         f"({nweb}/{ne}). {web_k} of those {nweb} events have a theme on any layer, against {noweb_k} of the "
         f"{ne - nweb} others. Never compare web-layer prevalence across groups with different crawl coverage."),
        ("C2 Changing sampling frame", f"Years covered per stream: {frame}. Trends must be read within a "
         "stream; the agendas change language and entry length over time."),
        ("C3 Document concentration", f"The top document gives {100 * (n.get('top1_share') or 0):.0f} % of clean "
         "rows" + (f" and {100 * conc_r:.0f} % of racism rows" if conc_r is not None else "") +
         ". Headline figures use events or documents, and the coding workbook caps passages per document."),
        ("C4 Snapshot dating", "Landing and linked pages were read at crawl time and are attached to past "
         "editions; trend tables should use own listing text (L1) or the document year."),
        ("C5 Organisation precision", "Names come from automatic extraction. " + (
            "They are curated (org_codes)." if an.curated else
            "They are not curated: top betweenness and backbone nodes can be places, targets or site menus. "
            "Report network results only after curation, and always with the sensitivity table.")),
        ("C6 Seed context leakage", f"{leak[0]} of {leak[1]} seed theme rows have context that names another "
         "event; coders must use the sentence column for attribution."),
        ("C7 Hub events", "One large event can make most projection edges" + (
            f" (here {100 * top_share:.0f} % in scenario A)" if top_share is not None else "") +
         ". Use the backbone, series nodes and hub_dependence, not eigenvector centrality on the raw projection."),
        ("C8 Country inference", f"Countries come from names, entry text and page locations (method per event "
         f"in events_enriched); {ne - n['events_with_country']} events stay unresolved, and a country named as a "
         "topic can be mistaken for the venue."),
        ("C9 Small numbers", ("Racism appears in own text for " + str(racism) + " events. " if racism is not None
                              else "") + "Groups below the thresholds are pooled, and every share is printed "
         "with k/n and a Wilson interval."),
        ("C10 Term list", "False matches exist (handles, place names, words from other languages). in_handle and "
         "in_event_title are flagged; per-family precision follows from coding."),
    ]
    tl = {(r["rows_in"], r["tie_sublevel"]): 0 for r in an.tables.get("tie_levels", ([], []))[0]}
    for r in an.tables.get("tie_levels", ([], []))[0]:
        tl[(r["rows_in"], r["tie_sublevel"])] += r["rows"]
    lv = {x: tl.get(("every row", x), 0) for x in SUBLEVELS}
    caveats.append(("C12 Tie levels", f"L1 (the listing itself) is structurally small: agenda entries and "
                    f"calendar cells name few organisations ({lv['L1']} rows in scenario {an.default_sc}, against "
                    f"{lv['L2']} on landing pages, {lv['L3s']} on the event's own host (L3s) and {lv['L3x']} on "
                    f"other hosts (L3x)). The default layer set is {a.layers}; the sensitivity table also gives "
                    "L1-L3 and L1+L2."))
    grad = an.facts.get("words_read_gradient") or {}
    if grad:
        lo_b, hi_b = next(iter(grad.items())), list(grad.items())[-1]
        std = [r for r in an.prev_sp if r["stream"] == "all" and r["theme"] == "any" and r["measure"] == "any_web"
               and r["group"] in [p_[2] for p_ in PERIODS]]
        caveats.append(("C13 Prevalence depends on how much was read", f"The share of events with a theme on any "
                        f"layer rises from {100 * (lo_b[1][2] or 0):.0f} % for events with {lo_b[0]} words read to "
                        f"{100 * (hi_b[1][2] or 0):.0f} % for {hi_b[0]} (prev_words_read). share_std standardises "
                        "the web-layer shares over these bins (the bin mix of all events with web text); by period, "
                        "any_web is " + ", ".join(f"{r['group']} {100 * (r['share'] or 0):.0f} % crude, "
                                                   f"{100 * (r['share_std'] or 0):.0f} % standardised" for r in std)
                        + ". The core measure (listing, landing page and the event's own host, no reports) "
                        "depends less on long linked documents."))
    dom = an.facts.get("events_one_document_dominates")
    if dom:
        caveats.append(("C14 One document per event", f"For {dom} events one document gives more than half of "
                        "the event's theme rows (event_theme_documents): raw theme rows are dominated by long "
                        "PDFs. Count documents and passages, not rows."))
    if an.descriptive:
        caveats.append(("C15 Descriptive themes", f"{', '.join(an.descriptive)} are reported in descriptive_themes "
                        "and never drive the theme filter or the any-theme measures."))
    if an.master_filter != "none":
        caveats.append(("C11 Inclusion criterion", "The masters were built with a theme filter, so the 'none' "
                        "rows of the sensitivity table equal the filtered ones, and the event-level share, lift "
                        "and shrunk columns of org_theme are left blank (their base rate would be inflated). "
                        "Rebuild with make_master.py --theme-filter none and run analytics.py with "
                        "--theme-filter source to get both."))
    def _caveat_no(c):
        m = re.match(r"C(\d+)", c[0])
        return int(m.group(1)) if m else 999
    caveats.sort(key=_caveat_no)
    L += ["## Caveats", ""] + [f"- **{h}.** {txt}" for h, txt in caveats] + [""]
    if sm["skipped"]:
        L += ["## Skipped", ""] + [f"- {k}: {v}" for k, v in sm["skipped"].items()] + [""]
    if sm["warnings"]:
        L += ["## Warnings", ""] + [f"- {w}" for w in sm["warnings"]] + [""]
    L += ["## Files", "", "- analytics.xlsx and tables/*.csv: all tables (README sheet first)",
          "- org_codes_template.xlsx: organisations to curate, merge candidates, precision sample",
          "- kwic_for_coding.xlsx: theme passages for coding", "- network/: Gephi files and node/edge CSVs",
          "- charts/: PNG charts", ""]
    return "\n".join(L)


# ==========================================================================
# Command line
# ==========================================================================
def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("results", nargs="?", help="results folder with master_network, master_themes and events.csv")
    ap.add_argument("-o", "--out", help="output folder (must differ from the results folder)")
    ap.add_argument("--run", nargs="+", action="extend", default=[], metavar="DIR",
                    help="raw run folders to look in for seed_links.json, pages.csv, sites.csv, sources.csv")
    ap.add_argument("--themes-file", help="theme file (default: themes_colonialism_racism_fascism.txt here)")
    ap.add_argument("--org-codes", help="curated org_codes.csv or .xlsx (read only)")
    ap.add_argument("--event-codes", help="curated event_codes.csv (read only)")
    ap.add_argument("--theme-filter", choices=THEME_FILTERS, default=None,
                    help="inclusion criterion for the network: direct (as source, but the listing, the event's own "
                         "page and its landing page count when any source of their event has a hit), source "
                         "(only sources with a theme hit), site / event (every source of a site or event with a "
                         "hit; agenda entries still judged one by one), or none (all sources). Default: the filter "
                         "make_master.py recorded in the masters' about sheet, else source")
    ap.add_argument("--network-scope", choices=NETWORK_SCOPES, default="participation",
                    help="participation (default): the network uses only rows that are evidence of taking part "
                         "(no linked reports, other pages of a site, rows about another year or copies); all: "
                         "every row, context_organisations included")
    ap.add_argument("--layers", choices=list(LAYER_SETS), default=DEFAULT_LAYERS,
                    help="tie levels of the default network: L1-L3s (default: listing, landing page and pages "
                         "and PDFs on the event's own host), L1-L3 (all linked pages and PDFs too) or L1+L2")
    ap.add_argument("--source-codes", metavar="FILE",
                    help="sources_to_code sheet of org_codes_template.xlsx (or that sheet as CSV): include 1 keeps "
                         "a source in the network whatever its doc_kind, 0 leaves it out (read only)")
    ap.add_argument("--chart-nodes", type=int, default=150,
                    help="organisations drawn in the backbone chart: the strongest of those in two or more "
                         "series (default 150)")
    ap.add_argument("--min-confidence", choices=list(CONF_RANK), default="medium",
                    help="confidence for the default network scenario C (default medium)")
    ap.add_argument("--backbone-alpha", type=float, default=0.05, help="disparity filter alpha (default 0.05)")
    ap.add_argument("--doc-cap", type=int, default=8, help="most passages per document in the KWIC theme sheets")
    ap.add_argument("--sample-size", type=int, default=120, help="KWIC coding sample size")
    ap.add_argument("--precision-per-stratum", type=int, default=25, help="names per precision stratum")
    ap.add_argument("--seed", type=int, default=20261017, help="random seed (samples, layouts, betweenness)")
    ap.add_argument("--betweenness-k", default="auto",
                    help="'auto' (exact up to 5,000 nodes, else 500 sampled sources), 'exact' or a number")
    ap.add_argument("--no-charts", action="store_true", help="do not draw the PNG charts")
    ap.add_argument("--no-graphs", action="store_true", help="do not write the Gephi network files")
    ap.add_argument("--selftest", action="store_true", help="run the built-in checks on a small synthetic folder")
    ap.add_argument("-q", "--quiet", action="store_true", help="warnings only")
    return ap


def run(args) -> Analysis:
    t0 = time.time()
    results, out = Path(args.results), Path(args.out)
    if not results.is_dir():
        raise InputError(f"results folder {results} not found")
    if out.resolve() == results.resolve():
        raise InputError("-o must not be the results folder (analytics never writes into it)")
    out.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(out / "analytics.log", mode="w", encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    log.addHandler(fh)
    try:
        an = Analysis(args, Project())
        log.info("Reading %s", results)
        an.load()
        an.build_events()
        an.build_refs()
        _derive(an)
        core_tables(an)
        a2_coverage(an)
        a3_ladder(an)
        a3b_documents(an)
        doc_hygiene(an)
        a4_prevalence(an)
        an.common_words = common_word_suggestions(an)
        log.info("Networks (%s)", "betweenness, projections, backbones, communities")
        a5_network(an)
        a7_org_theme(an)
        log.info("Organisation authority list and coding workbook")
        a1_authority(an)
        kw = a8_kwic(an)
        outputs: Dict[str, List[str]] = {}
        write_org_template(an, out)
        write_kwic(an, out, kw)
        write_analytics_workbook(an, out)
        outputs["workbooks"] = ["analytics.xlsx", "org_codes_template.xlsx", "kwic_for_coding.xlsx"]
        outputs["tables"] = sorted(f"tables/{k}.csv" for k in ["README", "diagnostics"] + list(an.tables))
        if args.no_graphs:
            an.skipped["network files"] = "--no-graphs"
        else:
            outputs["network"] = write_networks(an, out)
        if args.no_charts:
            an.skipped["charts"] = "--no-charts"
        else:
            outputs["charts"] = write_charts(an, out)
        an.summary = write_summary(an, out, outputs, time.time() - t0)
        log.info("Done in %.1f s: %s", time.time() - t0, out)
        return an
    finally:
        log.removeHandler(fh)
        fh.close()


def main(argv: Optional[List[str]] = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv[:1] == ["agreement"]:
        print("analytics.py agreement: coder agreement (kappa, alpha) is planned for phase 2 and not available yet.",
              file=sys.stderr)
        return 2
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.WARNING if args.quiet else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    if args.selftest:
        return selftest()
    if not args.results or not args.out:
        build_parser().print_usage(sys.stderr)
        print("analytics.py: RESULTS_DIR and -o OUT are required", file=sys.stderr)
        return 2
    try:
        run(args)
    except InputError as exc:
        log.error("%s", exc)
        return 2
    return 0


# ==========================================================================
# Self-test: a tiny synthetic results folder, in the older and the long layout
# ==========================================================================
_ST_EVENTS = [  # name, date, link, listed_on
    ("Climate camp Berlin [DE]", "2024-07-01", "https://camp.example.org/2024/", "pdf:CJA-campaction-dates.pdf"),
    ("Climate camp Berlin [DE]", "2025-07-01", "https://camp.example.org/2024/", "pdf:CJA-campaction-dates.pdf"),
    ("Antifascist week Paris", "2023-03-10", "https://antifa.example.fr/", "pdf:Agenda_Les_echos_du_vent.pdf"),
    ("Shell must fall", "2024-05-20", "https://shellmustfall.example.org/", "https://themovementhub.org/events/"),
    ("Online workshop", "2025-01-15", "", "pdf:Agenda_Les_echos_du_vent.pdf"),
]
_ST_ECHOS = "pdf:Agenda_Les_echos_du_vent.pdf"
_ST_NET = [  # event index (None = pseudo), organisation, confidence, source_link, site, site_role, title
    (0, "Ende Gelände", "high", "https://camp.example.org/2024/", "camp.example.org", "linked", "no"),
    (0, "Attac", "medium", "https://camp.example.org/2024/", "camp.example.org", "linked", "no"),
    (0, "FAQ", "low", "https://camp.example.org/2024/", "camp.example.org", "linked", "no"),
    (0, "Bakery Ltd", "medium", "https://camp.example.org/about/", "camp.example.org", "linked", "no"),
    (1, "Ende Gelaende", "medium", "https://camp.example.org/2024/", "camp.example.org", "linked", "no"),
    (1, "Attac", "medium", "https://camp.example.org/2024/", "camp.example.org", "linked", "no"),
    (3, "Shell", "high", "https://shellmustfall.example.org/", "shellmustfall.example.org", "linked", "yes"),
    (3, "Ende Gelände", "medium", "https://shellmustfall.example.org/", "shellmustfall.example.org", "linked", "no"),
    (3, "Attac", "low", "https://shellmustfall.example.org/", "shellmustfall.example.org", "linked", "no"),
    (3, "Extinction Rebellion", "medium", "https://shellmustfall.example.org/", "shellmustfall.example.org",
     "linked", "no"),
    (2, "Antifa Ost", "medium", _ST_ECHOS, _ST_ECHOS, "seed", "no"),
    (None, "Radio\x0b Orange", "medium", "https://camp.example.org/imprint/", "camp.example.org", "linked", "no"),
]
_ST_S1 = "We fight racism and colonialism together with Ende Gelände and Attac."
_ST_THEMES = [  # theme, term, event indexes, source_link, site, site_role, sentence, orgs, region, duplicate_of
    ("racism", "racism", [0, 1], "https://camp.example.org/2024/", "camp.example.org", "linked", _ST_S1,
     "Ende Gelände; Attac", "", ""),
    ("colonialism", "colonialism", [0, 1], "https://camp.example.org/2024/", "camp.example.org", "linked", _ST_S1,
     "Ende Gelände; Attac", "", ""),
    ("racism", "racism", [0, 1], "https://camp.example.org/2024/", "camp.example.org", "linked", _ST_S1,
     "Ende Gelände; Attac", "", "https://camp.example.org/2024/"),
    ("fascism", "Antifascist", [2], _ST_ECHOS, _ST_ECHOS, "seed", "Antifascist week against fascism in Paris.",
     "Antifa Ost", "", ""),
    ("fascism", "fascism", [2], _ST_ECHOS, _ST_ECHOS, "seed", "Antifascist week against fascism in Paris.",
     "Antifa Ost", "", ""),
    ("colonialism", "neocolonial", [3], "https://shellmustfall.example.org/", "shellmustfall.example.org", "linked",
     "Shell profits from neocolonial extraction.", "Shell", "", ""),
    ("racism", "racism", [3], "https://shellmustfall.example.org/", "shellmustfall.example.org", "linked",
     "Footer: no racism here.", "", "footer", ""),
]


def _selftest_folder(folder: Path, long_format: bool) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    ev_fields = ["event", "date", "event_link", "listed_on"] + (["event_key", "country_iso", "country_source",
                                                                 "location_text"] if long_format else [])
    key = lambda e: f"{e[2] or 'pdf:x'}#{e[1]}"  # noqa: E731
    rows = []
    for e in _ST_EVENTS:
        r = dict(zip(["event", "date", "event_link", "listed_on"], e))
        if long_format:
            r.update(event_key=key(e), country_iso="FR" if "Paris" in e[0] else "DE", country_source="tag",
                     location_text="x")
        rows.append(r)
    write_csv(folder / "events.csv", rows, ev_fields)
    nf = ["event", "event_date", "event_link", "organisation", "confidence", "source_link", "site", "site_role",
          "named_in_event_title", "found_by", "mentions_on_source", "example_context"]
    if long_format:
        nf += ["event_key", "layer", "source_unit", "source_has_theme"]
    rows = []
    for i, org, conf, src, site, role, title in _ST_NET:
        e = _ST_EVENTS[i] if i is not None else (NO_EVENT, "", "", "")
        r = {"event": e[0], "event_date": e[1], "event_link": e[2], "organisation": org, "confidence": conf,
             "source_link": src, "site": site, "site_role": role, "named_in_event_title": title,
             "found_by": "ner" if conf == "low" else "watchlist", "mentions_on_source": 1,
             "example_context": f"... {org} ..."}
        if long_format:
            seed = role == "seed"
            unit = f"{src}#entry-{i + 1}" if seed else src
            r.update(event_key=key(e) if i is not None else "", source_unit=unit,
                     layer="listing_seed" if seed else ("landing" if norm_link(src) == norm_link(e[2]) else
                                                        "linked_page"),
                     source_has_theme=yn(src in ("https://camp.example.org/2024/",
                                                 "https://shellmustfall.example.org/") or seed))
        rows.append(r)
    write_csv(folder / "master_network.csv", rows, nf)
    tf = ["theme", "term", "event", "event_date", "event_link", "source_link", "sentence", "context_before",
          "context_after", "organisations_in_paragraph", "site", "site_role", "is_pdf", "region", "boilerplate",
          "duplicate_of"] + (["event_key", "layer", "source_unit"] if long_format else [])
    rows = []
    for theme, term, evs, src, site, role, sent, orgs, region, dup in _ST_THEMES:
        base = {"theme": theme, "term": term, "source_link": src, "sentence": sent, "context_before": "Before.",
                "context_after": "After.", "organisations_in_paragraph": orgs, "site": site, "site_role": role,
                "is_pdf": "no", "region": region, "boilerplate": "no", "duplicate_of": dup}
        if long_format:
            for i in evs:
                e = _ST_EVENTS[i]
                seed = role == "seed"
                rows.append({**base, "event": e[0], "event_date": e[1], "event_link": e[2], "event_key": key(e),
                             "layer": "listing_seed" if seed else "landing",
                             "source_unit": f"{src}#entry-{i + 1}" if seed else src})
        else:
            rows.append({**base, "event": " | ".join(_ST_EVENTS[i][0] for i in evs),
                         "event_date": " | ".join(_ST_EVENTS[i][1] for i in evs),
                         "event_link": " | ".join(_ST_EVENTS[i][2] for i in evs)})
    write_csv(folder / "master_themes.csv", rows, tf)


def selftest() -> int:
    import shutil
    import tempfile
    import networkx as nx
    fails: List[str] = []

    def check(label: str, got, want) -> None:
        ok = (abs(got - want) < 0.011) if isinstance(want, float) and got is not None else got == want
        if not ok:
            fails.append(f"{label}: got {got!r}, want {want!r}")
        log.info("selftest %-55s %s", label, "ok" if ok else f"FAIL (got {got!r}, want {want!r})")

    # ---- pure functions
    check("wilson(4, 172) low", wilson(4, 172)[0], 0.0091)
    check("wilson(4, 172) high", wilson(4, 172)[1], 0.0584)
    check("wilson(0, 10)", wilson(0, 10), (0.0, 0.2775))
    check("stratified interval: effective n 20", _stratified_interval(0.5, 0.0125, 40), wilson(10, 20))
    check("stratified interval: no variance uses n", _stratified_interval(1.0, 0.0, 12), wilson(12, 12))
    G = nx.Graph()
    G.add_edges_from([("E:1", "O:a"), ("E:1", "O:b"), ("E:1", "O:c"), ("E:2", "O:a"), ("E:2", "O:b")])
    P = project(G)
    check("projection: Newman weight a-b (1/2 + 1)", P.edge("O:a", "O:b")["weight"], 1.5)
    check("projection: shared events a-b", P.edge("O:a", "O:b")["shared"], 2)
    check("projection: Newman weight a-c", P.edge("O:a", "O:c")["weight"], 0.5)
    check("projection: degree and edge count", (P.degree("O:a"), P.number_of_edges()), (2, 3))
    check("backbone of a projection is a graph", disparity(P, 0.99).number_of_nodes(), 3)
    S = nx.Graph()
    S.add_edge("O:h", "O:x", weight=10.0)
    for i in range(6):
        S.add_edge("O:h", f"O:{i}", weight=0.1)
    check("disparity keeps only the dominant tie", sorted(disparity(S, 0.05).edges), [("O:h", "O:x")])
    check("subtheme('anticolonial')", subtheme("colonialism", "anticolonial"), "anticolonial")
    check("stance 'against' window", stance("racism", "Demo against "), "against")
    check("period_of(2024)", period_of(2024), "2023-24")
    check("doc_year_of", doc_year_of("https://x.org/wp-content/uploads/2022/07/a.pdf"), 2022)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        bad = nx.Graph()
        bad.add_node("O:a\x0bb", label=None, xs=["x", "y"])
        bad.add_edge("O:a\x0bb", "E:1", w=float("nan"))
        write_graph(bad, tmp / "g")
        check("GEXF round trip with None and \\x0b", sorted(nx.read_gexf(str(tmp / "g.gexf")).nodes), ["E:1", "O:ab"])
        check("GraphML round trip", nx.read_graphml(str(tmp / "g.graphml")).number_of_edges(), 1)
        check("no .tmp files left", sorted(p.name for p in tmp.glob("*.tmp")), [])
        # ---- full runs, older and long layout
        for long_format in (False, True):
            lab = "long" if long_format else "older"
            res = tmp / f"res_{lab}"
            _selftest_folder(res, long_format)
            outs = []
            for rep in (1, 2):
                o = tmp / f"out_{lab}_{rep}"
                argv = [str(res), "-o", str(o)] + (["--no-charts"] if rep == 2 or long_format else [])
                an = run(build_parser().parse_args(argv))
                outs.append((o, an))
            o, an = outs[0]
            tf_ = an.a.theme_filter
            check(f"[{lab}] default theme filter without a record in the masters", tf_, "source")
            check(f"[{lab}] events", len(an.events), 5)
            check(f"[{lab}] unique event ids", len({e['event_id'] for e in an.events}), 5)
            check(f"[{lab}] streams", dict(Counter(e["stream_primary"] for e in an.events)),
                  {"cja": 2, "echos": 2, "hub": 1})
            check(f"[{lab}] series: two editions share one series", an.events[0]["series_id"] ==
                  an.events[1]["series_id"], True)
            check(f"[{lab}] theme refs matched", an.theme_ref_counts.get("unmatched", 0), 0)
            check(f"[{lab}] theme instances", len(an.instances), 7)
            check(f"[{lab}] clean instances", sum(i["clean"] for i in an.instances), 5)
            check(f"[{lab}] unique clean passages", len({i["passage_id"] for i in an.instances if i["clean"]}), 3)
            check(f"[{lab}] eff_docs (shares .4/.4/.2)", an.concentration["ALL"]["eff_docs"], 2.78)
            check(f"[{lab}] umlaut variant merged", len({t["org_key"] for t in an.ties
                                                         if "ende" in t["org_key"]}), 1)
            check(f"[{lab}] R1: rows kept / total", an.filter_counts["rows"], (10, 12))
            check(f"[{lab}] R1: Bakery Ltd left out", any(t["organisation"] == "Bakery Ltd" and t["pass"]["source"]
                                                          for t in an.ties), False)
            sens = {(r["theme_filter"], r["scenario"], r["layers"], r["node_level"]): r for r in an.sensitivity
                    if r["variant"] in SCOPE_LABEL.values()}
            check(f"[{lab}] scenario A orgs (filter source)", sens[(tf_, "A", "L1-L3", "edition")]["orgs"], 6)
            check(f"[{lab}] scenario A orgs (filter none)", sens[("none", "A", "L1-L3", "edition")]["orgs"], 7)
            check(f"[{lab}] scenario C orgs (medium+)", sens[(tf_, "C", "L1-L3", "edition")]["orgs"], 5)
            check(f"[{lab}] scenario B drops title-named Shell", sens[(tf_, "B", "L1-L3", "edition")]["orgs"], 5)
            check(f"[{lab}] Ende Gelande in 2+ series (C, series)",
                  sens[(tf_, "C", "L1-L3", "series")]["orgs_in_2plus"], 2)
            cen = {r["org_key"]: r for r in an.org_centrality}
            check(f"[{lab}] FAQ (low) not in default network", "faq" in cen, False)
            check(f"[{lab}] Attac n_series in C", cen.get("attac", {}).get("n_series"), 2)
            own = {t: sum(1 for e in an.events if e[f"{t}_own"]) for t in an.themes}
            check(f"[{lab}] own-text events (fascism)", own.get("fascism"), 1)
            check(f"[{lab}] landing colonialism events", sum(1 for e in an.events if e["colonialism_landing"]), 3)
            check(f"[{lab}] country of Paris event", an.events[2]["countries"][:1], ["FR"])
            check(f"[{lab}] KWIC units (passage x theme)", an.facts["kwic"]["units"], 4)
            for name in ("analytics.xlsx", "kwic_for_coding.xlsx", "org_codes_template.xlsx", "summary.json",
                         "summary.md", "network/bipartite_C.gexf", "network/org_backbone_C.graphml"):
                check(f"[{lab}] wrote {name}", (o / name).exists() and (o / name).stat().st_size > 0, True)
            if not long_format:
                check("[older] six charts or a skip note", len(list((o / "charts").glob("*.png"))) +
                      sum(1 for k in an.skipped if k.startswith("chart ")), 6)
            o2 = outs[1][0]
            diff = [p.name for p in sorted((o / "tables").glob("*.csv"))
                    if p.read_bytes() != (o2 / "tables" / p.name).read_bytes()]
            check(f"[{lab}] re-run gives byte-identical tables", diff, [])
        # ---- long layout: a term found twice in one sentence gives two identical rows
        res2 = tmp / "res_long_twice"
        shutil.copytree(tmp / "res_long", res2)
        fields, th_rows = read_csv(res2 / "master_themes.csv")
        write_csv(res2 / "master_themes.csv", th_rows + [dict(th_rows[-1])], fields)
        an = run(build_parser().parse_args([str(res2), "-o", str(tmp / "out_twice"), "--no-charts",
                                            "--no-graphs"]))
        check("[long] identical row for the same event is a second instance", len(an.instances), 8)
        # ---- a passage repeated in the main content counts; agenda entries are judged one by one under 'site'
        res3 = tmp / "res_long_rep"
        shutil.copytree(tmp / "res_long", res3)
        fields, th_rows = read_csv(res3 / "master_themes.csv")
        rep = dict(th_rows[-2], sentence="A camp against racism.", region="main", boilerplate="yes",
                   menu_or_sidebar="no")
        old_bp = dict(th_rows[-2], sentence="Old layout repeated text on racism.", region="", boilerplate="yes")
        write_csv(res3 / "master_themes.csv", th_rows + [rep], fields + ["menu_or_sidebar"])
        nfields, n_rows = read_csv(res3 / "master_network.csv")
        e4 = _ST_EVENTS[4]
        n_rows.append(dict(n_rows[10], event=e4[0], event_date=e4[1], event_link=e4[2], organisation="Unthemed Collective",
                           event_key=f"pdf:x#{e4[1]}", source_unit=f"{_ST_ECHOS}#entry-5", source_has_theme="no"))
        write_csv(res3 / "master_network.csv", n_rows, nfields)
        an = run(build_parser().parse_args([str(res3), "-o", str(tmp / "out_rep"), "--no-charts", "--no-graphs",
                                            "--theme-filter", "site"]))
        reps = [i for i in an.instances if i["sentence"] == "A camp against racism."]
        check("[long] repeated main-content passage is not furniture", [(i["furniture"], i["repeated"])
                                                                        for i in reps], [(False, True)])
        passes = {t["organisation"]: t["pass"]["site"] for t in an.ties}
        check("[long] 'site': unthemed agenda entry left out", passes.get("Unthemed Collective"), False)
        check("[long] 'site': themed agenda entry kept", passes.get("Antifa Ost"), True)
        res4 = tmp / "res_older_bp"
        shutil.copytree(tmp / "res_older", res4)
        fields, th_rows = read_csv(res4 / "master_themes.csv")
        write_csv(res4 / "master_themes.csv", th_rows + [dict(th_rows[-2], **{k: v for k, v in old_bp.items()
                                                                            if k in fields})], fields)
        an = run(build_parser().parse_args([str(res4), "-o", str(tmp / "out_bp"), "--no-charts", "--no-graphs"]))
        check("[older] boilerplate without menu_or_sidebar is still furniture",
              [i["furniture"] for i in an.instances if i["sentence"].startswith("Old layout")], [True])
        # ---- newer masters: context_organisations, doc_kind/participation and focus_theme columns
        res5 = tmp / "res_new_cols"
        shutil.copytree(tmp / "res_long", res5)
        nfields, n_rows = read_csv(res5 / "master_network.csv")
        extra = ["doc_kind", "doc_kind_reason", "attribution", "source_year", "duplicate_document_of",
                 "translation_of", "participation", "context_reason"]
        for r in n_rows:
            r.update(doc_kind="listing" if r["layer"] == "listing_seed" else "page", attribution="event",
                     participation="yes")
        e3 = _ST_EVENTS[3]
        ctx = [dict(n_rows[6], organisation=f"Report Org {i}", org_key=f"report org {i}", confidence="high",
                    source_link="https://shellmustfall.example.org/report.pdf", is_pdf="yes", layer="linked_pdf",
                    source_unit="https://shellmustfall.example.org/report.pdf", named_in_event_title="no",
                    doc_kind="report", doc_kind_reason="report word", participation="no",
                    context_reason="linked report or study", event=e3[0], event_date=e3[1], event_link=e3[2])
               for i in (1, 2)]
        write_csv(res5 / "master_network.csv", n_rows, nfields + extra)
        write_csv(res5 / "master_network_context.csv", ctx, nfields + extra)
        fields, th_rows = read_csv(res5 / "master_themes.csv")
        for r in th_rows:
            r["focus_theme"] = "yes"
        e4 = _ST_EVENTS[4]
        th_rows.append(dict(th_rows[3], theme="global_south", term="Global South", focus_theme="no",
                            sentence="Solidarity with the Global South.", event=e4[0], event_date=e4[1],
                            event_link=e4[2], event_key=f"pdf:x#{e4[1]}", source_unit=f"{_ST_ECHOS}#entry-5"))
        write_csv(res5 / "master_themes.csv", th_rows, fields + ["focus_theme", "passage_id"])
        an_old = run(build_parser().parse_args([str(tmp / "res_long"), "-o", str(tmp / "out_cols0"), "--no-charts",
                                                "--no-graphs"]))
        an = run(build_parser().parse_args([str(res5), "-o", str(tmp / "out_cols"), "--no-charts", "--no-graphs",
                                            "--theme-filter", "direct"]))
        check("[new columns] --theme-filter direct accepted and recorded", an.a.theme_filter, "direct")
        cen = {r["org_key"]: r for r in an.org_centrality}
        check("[new columns] context rows stay out of the default network", "report org 1" in cen, False)
        sens = {r["variant"]: r for r in an.sensitivity if r["node_level"] == "series" and r["scenario"] == "C"
                and r["layers"] == an.a.layers and r["theme_filter"] == "direct"}
        dflt = sens[SCOPE_LABEL["participation"]]["orgs"]
        plus = next(v["orgs"] for k, v in sens.items() if k.startswith("+context"))
        check("[new columns] '+context rows' counts the two report organisations", plus - dflt, 2)
        check("[new columns] global_south rows leave the any-theme events unchanged",
              sum(e["any_any"] for e in an.events), sum(e["any_any"] for e in an_old.events))
        check("[new columns] focus themes", an.themes, ["colonialism", "racism", "fascism"])
        dt = {r["theme"]: r for r in an.tables["descriptive_themes"][0]}
        check("[new columns] global_south listed in descriptive_themes", dt.get("global_south", {}).get("events_any"), 1)
        an = run(build_parser().parse_args([str(res5), "-o", str(tmp / "out_cols_all"), "--no-charts", "--no-graphs",
                                            "--theme-filter", "direct", "--network-scope", "all"]))
        cen = {r["org_key"]: r for r in an.org_centrality}
        check("[new columns] --network-scope all keeps the context rows", "report org 1" in cen, True)
        # ---- --source-codes: include 1 brings a report back into the default network
        sc_path = tmp / "source_codes.csv"
        write_csv(sc_path, [{"source_unit": "https://shellmustfall.example.org/report.pdf", "include": "1"}],
                  ["source_unit", "include"])
        an = run(build_parser().parse_args([str(res5), "-o", str(tmp / "out_cols_sc"), "--no-charts", "--no-graphs",
                                            "--theme-filter", "direct", "--source-codes", str(sc_path)]))
        cen = {r["org_key"]: r for r in an.org_centrality}
        check("[new columns] --source-codes include 1 keeps a report", "report org 1" in cen, True)
        # ---- tie sublevels and document identity
        check("same_host: subdomain", same_host("events.camp.org", "camp.org"), True)
        check("norm_doc keeps the path's case", norm_doc("HTTPS://WWW.Uemss.org/img/A.pdf#p=2"), "uemss.org/img/A.pdf")
        # ---- curation loop
        res = tmp / "res_older"
        codes = tmp / "org_codes.csv"
        write_csv(codes, [{"org_key": "attac", "organisation": "Attac", "include": "0"},
                          {"org_key": "stale key", "organisation": "Extinction Rebellion", "include": "1",
                           "actor_class": "movement_org", "is_organisation": "yes"},
                          {"org_key": "shell", "organisation": "Shell", "actor_class": "company",
                           "is_organisation": "yes"}],
                  ["org_key", "organisation", "include", "merge_into_key", "actor_class", "is_organisation"])
        an = run(build_parser().parse_args([str(res), "-o", str(tmp / "out_cur"), "--org-codes", str(codes),
                                            "--no-charts", "--no-graphs"]))
        cen = {r["org_key"]: r for r in an.org_centrality}
        check("[codes] include=0 drops Attac", "attac" in cen, False)
        check("[codes] curated scenario is the default", an.default_sc, "C_curated")
        check("[codes] stale org_key gives a warning", any("org_key" in w for w in an.warnings), True)
        check("[codes] precision read back", bool(an.precision_rows), True)
        # ---- missing required input
        empty = tmp / "empty"
        empty.mkdir()
        try:
            run(build_parser().parse_args([str(empty), "-o", str(tmp / "out_empty")]))
            check("missing master gives InputError", False, True)
        except InputError:
            check("missing master gives InputError", True, True)
    print(("SELFTEST FAILED:\n  " + "\n  ".join(fails)) if fails else "selftest: all checks passed")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
